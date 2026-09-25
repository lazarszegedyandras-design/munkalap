from datetime import date, datetime, timedelta
from io import BytesIO
from PIL import Image, ImageDraw
import hashlib
import json
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from fastapi import HTTPException
from fastapi.testclient import TestClient

from app.database import Base, get_db
from app.main import app
from app.models import AuthSession, SecurityRateLimit, Asset, AssetComponent, AssetComponentReplacement, AssetMeterReading, AuditLog, CrmActivity, CrmContract, CrmOpportunity, Customer, CustomerContact, EmployeeProfile, LeaveEntitlement, LeaveRequest, Location, Material, ProcurementAttachment, ProcurementRequest, ROLE_ADMIN, ROLE_OFFICE, ROLE_TECHNICIAN, Role, User, WorkCalendarDay, WorkOrder, WorkOrderArchive, WorkOrderCompletionSnapshot, WorkOrderSignature, MobileDevice, PushDeliveryOutbox
from app.security import decode_token_claims, get_password_hash, verify_password
from app.reset_password import PasswordResetError, reset_user_password
from app.services.work_order_concurrency import commit_with_optimistic_lock
from app.services.auth_security import create_auth_session, encrypt_mfa_secret, generate_recovery_codes, _totp_at
from app.config import get_settings, production_security_errors
from app.services.work_order_archive_storage import archive_path
from app.version import APP_VERSION, semantic_version_parts
from app.access_control import PERMISSION_ADMIN_ACCESS, PERMISSION_ADMIN_BACKUP_MANAGE, PERMISSION_ADMIN_BACKUP_RESTORE, PERMISSION_CRM_ACCESS, PERMISSION_CRM_CONTRACT_MANAGE, PERMISSION_CRM_WRITE, PERMISSION_HR_ACCESS, PERMISSION_HR_ADMIN, PERMISSION_HR_LEAVE_APPROVE, PERMISSION_HR_LEAVE_REQUEST, PERMISSION_HR_LEAVE_SELF, PERMISSION_HR_SUMMARY_VIEW, PERMISSION_PROCUREMENT_APPROVE, PERMISSION_SERVICE_ACCESS
from app.services.permissions import ensure_permission_catalog, grant_service_access_if_missing, set_user_permissions
from app.services.audit import verify_audit_chain

SQLALCHEMY_DATABASE_URL = "sqlite://"
engine = create_engine(SQLALCHEMY_DATABASE_URL, connect_args={"check_same_thread": False}, poolclass=StaticPool)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


def override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()


app.dependency_overrides[get_db] = override_get_db
client = TestClient(app)
FIXTURE_DIR = Path(__file__).resolve().parent / "fixtures"


@pytest.fixture(autouse=True)
def synthetic_private_sources(monkeypatch):
    from app.importers import customer_assets
    from app.services import backup_restore, company_profiles, hr_import

    monkeypatch.setattr(customer_assets, "source_file", lambda: FIXTURE_DIR / "customer_assets_20260616.json")
    monkeypatch.setattr(hr_import, "source_file", lambda: FIXTURE_DIR / "hr_leave_2026_rlb.json")
    monkeypatch.setattr(company_profiles, "STATIC_COMPANY_PROFILES_PATH", FIXTURE_DIR / "company_profiles.json")
    monkeypatch.setattr(backup_restore, "STATIC_COMPANY_PROFILES", FIXTURE_DIR / "company_profiles.json")
    company_profiles.load_company_profiles.cache_clear()
    yield
    company_profiles.load_company_profiles.cache_clear()


def setup_function():
    client.cookies.clear()
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    db = TestingSessionLocal()
    admin_role = Role(name=ROLE_ADMIN)
    office_role = Role(name=ROLE_OFFICE)
    technician_role = Role(name=ROLE_TECHNICIAN)
    db.add_all([admin_role, office_role, technician_role])
    db.flush()
    admin = User(email="admin@example.com", full_name="Admin", password_hash=get_password_hash("admin123"), role_id=admin_role.id)
    customer = Customer(name="Teszt ügyfél")
    material = Material(sku="TESZT-1", name="Teszt anyag", unit="db", unit_price=100)
    db.add_all([admin, customer, material])
    db.flush()
    db.add(Location(customer_id=customer.id, name="Teszt helyszín", address="Teszt cím"))
    db.commit()
    db.close()


def auth_headers_for(email: str, password: str):
    # Most API tests need an authenticated principal, not a second login test.
    # Create a server-side session directly so the browser login endpoint can keep
    # its JWT exclusively in HttpOnly cookies.
    db = TestingSessionLocal()
    user = db.query(User).filter(User.email == email).one()
    assert verify_password(password, user.password_hash)
    _session, token, _csrf = create_auth_session(db, user, get_settings(), mfa_verified=True)
    db.commit()
    db.close()
    return {"Authorization": f"Bearer {token}"}


def auth_headers():
    db = TestingSessionLocal()
    admin = db.query(User).filter(User.email == "admin@example.com").one()
    admin.mfa_enabled = True
    admin.mfa_secret_encrypted = encrypt_mfa_secret("JBSWY3DPEHPK3PXP", get_settings())
    db.commit()
    db.close()
    return auth_headers_for("admin@example.com", "admin123")


def grant_service_access(db, user):
    ensure_permission_catalog(db)
    db.flush()
    grant_service_access_if_missing(db, user)
    db.flush()


def test_login_and_me():
    login_response = client.post("/api/auth/login", data={"username": "admin@example.com", "password": "admin123"})
    assert login_response.status_code == 200, login_response.text
    assert login_response.json()["access_token"] is None
    assert client.cookies.get("workapp_session")
    assert client.cookies.get("workapp_csrf")

    response = client.get("/api/auth/me")
    assert response.status_code == 200
    assert response.json()["email"] == "admin@example.com"


def test_cookie_authenticated_write_requires_csrf_and_logout_revokes_session():
    login_response = client.post("/api/auth/login", data={"username": "admin@example.com", "password": "admin123"})
    assert login_response.status_code == 200
    session_cookie = client.cookies.get("workapp_session")
    csrf = client.cookies.get("workapp_csrf")
    assert session_cookie and csrf

    blocked = client.post("/api/auth/logout")
    assert blocked.status_code == 403
    assert client.get("/api/auth/me").status_code == 200

    logout = client.post("/api/auth/logout", headers={"X-CSRF-Token": csrf})
    assert logout.status_code == 204
    client.cookies.set("workapp_session", session_cookie)
    assert client.get("/api/auth/me").status_code == 401


def test_login_rate_limit_blocks_username_without_locking_shared_ip_at_user_threshold():
    settings = get_settings()
    old_user_limit = settings.login_rate_limit_attempts
    old_ip_limit = settings.login_ip_rate_limit_attempts
    try:
        settings.login_rate_limit_attempts = 3
        settings.login_ip_rate_limit_attempts = 30
        for _ in range(3):
            response = client.post("/api/auth/login", data={"username": "admin@example.com", "password": "rossz"})
            assert response.status_code == 401
        blocked = client.post("/api/auth/login", data={"username": "admin@example.com", "password": "admin123"})
        assert blocked.status_code == 429
        assert "Retry-After" in blocked.headers

        # A different account from the same test client/IP is not locked merely
        # because one username reached its tighter threshold.
        db = TestingSessionLocal()
        role = db.query(Role).filter(Role.name == ROLE_OFFICE).one()
        db.add(User(email="second@example.com", full_name="Second", password_hash=get_password_hash("Masik!Jelszo9"), role_id=role.id))
        db.commit()
        db.close()
        other = client.post("/api/auth/login", data={"username": "second@example.com", "password": "Masik!Jelszo9"})
        assert other.status_code == 200, other.text
    finally:
        settings.login_rate_limit_attempts = old_user_limit
        settings.login_ip_rate_limit_attempts = old_ip_limit


def test_admin_mfa_setup_login_and_totp_replay_protection():
    settings = get_settings()
    old_required = settings.admin_mfa_required
    old_key = settings.mfa_encryption_key
    try:
        settings.admin_mfa_required = True
        settings.mfa_encryption_key = "test-mfa-encryption-key-that-is-long-enough-123"

        login_response = client.post("/api/auth/login", data={"username": "admin@example.com", "password": "admin123"})
        assert login_response.status_code == 200
        body = login_response.json()
        assert body["mfa_setup_required"] is True
        assert body["access_token"] is None

        setup = client.post("/api/auth/mfa/setup", json={"challenge_token": body["challenge_token"]})
        assert setup.status_code == 200, setup.text
        secret = setup.json()["secret"]
        from datetime import datetime as _dt
        code = _totp_at(secret, int(_dt.utcnow().timestamp()) // 30)
        confirm = client.post("/api/auth/mfa/confirm", json={"challenge_token": body["challenge_token"], "code": code})
        assert confirm.status_code == 200, confirm.text
        assert confirm.json()["access_token"] is None
        assert len(confirm.json()["recovery_codes"]) == 10
        assert client.get("/api/auth/me").status_code == 200

        csrf = client.cookies.get("workapp_csrf")
        assert client.post("/api/auth/logout", headers={"X-CSRF-Token": csrf}).status_code == 204
        second_login = client.post("/api/auth/login", data={"username": "admin@example.com", "password": "admin123"})
        assert second_login.status_code == 200
        assert second_login.json()["mfa_required"] is True
        replay = client.post("/api/auth/mfa/verify", json={"challenge_token": second_login.json()["challenge_token"], "code": code})
        assert replay.status_code == 401
        recovery = confirm.json()["recovery_codes"][0]
        recovered = client.post("/api/auth/mfa/verify", json={"challenge_token": second_login.json()["challenge_token"], "code": recovery})
        assert recovered.status_code == 200, recovered.text
    finally:
        settings.admin_mfa_required = old_required
        settings.mfa_encryption_key = old_key


def test_security_headers_and_health_readiness_contract():
    response = client.get("/api/health/live")
    assert response.status_code == 200
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["x-frame-options"] == "DENY"
    assert "frame-ancestors 'none'" in response.headers["content-security-policy"]
    ready = client.get("/api/health/ready")
    assert ready.status_code == 200
    assert ready.json()["status"] == "ready"


def test_forced_first_password_change_with_empty_admin_password():
    db = TestingSessionLocal()
    admin = db.query(User).filter(User.email == "admin@example.com").one()
    admin.password_hash = get_password_hash("")
    admin.must_change_password = True
    db.commit()
    db.close()

    login_response = client.post("/api/auth/login", data={"username": "admin@example.com", "password": ""})
    assert login_response.status_code == 200, login_response.text
    assert login_response.json()["access_token"] is None
    headers = {"X-CSRF-Token": client.cookies.get("workapp_csrf")}

    me_response = client.get("/api/auth/me")
    assert me_response.status_code == 200
    assert me_response.json()["must_change_password"] is True

    blocked_response = client.get("/api/customers", headers=headers)
    assert blocked_response.status_code == 403
    assert blocked_response.json()["detail"]["code"] == "password_change_required"

    wrong_response = client.post(
        "/api/auth/change-password",
        headers=headers,
        json={"current_password": "rossz", "new_password": "UjAdmin!123"},
    )
    assert wrong_response.status_code == 400

    change_response = client.post(
        "/api/auth/change-password",
        headers=headers,
        json={"current_password": "", "new_password": "UjAdmin!123"},
    )
    assert change_response.status_code == 200, change_response.text
    assert change_response.json()["must_change_password"] is False

    allowed_response = client.get("/api/customers", headers=headers)
    assert allowed_response.status_code == 200
    assert client.post("/api/auth/login", data={"username": "admin@example.com", "password": "UjAdmin!123"}).status_code == 200


def test_password_policy_and_own_password_change():
    headers = auth_headers()
    db = TestingSessionLocal()
    role_id = db.query(Role).filter(Role.name == ROLE_OFFICE).one().id
    db.close()

    weak_response = client.post(
        "/api/users",
        headers=headers,
        json={
            "email": "weak@example.com",
            "full_name": "Gyenge Jelszó",
            "password": "gyengejelszo",
            "role_id": role_id,
            "is_active": True,
        },
    )
    assert weak_response.status_code == 422

    create_response = client.post(
        "/api/users",
        headers=headers,
        json={
            "email": "office-change@example.com",
            "full_name": "Jelszó Csere",
            "password": "Ideiglenes!A",
            "role_id": role_id,
            "is_active": True,
        },
    )
    assert create_response.status_code == 201, create_response.text
    assert create_response.json()["must_change_password"] is True

    user_headers = auth_headers_for("office-change@example.com", "Ideiglenes!A")
    change_response = client.post(
        "/api/auth/change-password",
        headers=user_headers,
        json={"current_password": "Ideiglenes!A", "new_password": "Vegleges!B"},
    )
    assert change_response.status_code == 200, change_response.text
    assert change_response.json()["must_change_password"] is False



def test_local_password_reset_forces_change_and_writes_audit_log():
    db = TestingSessionLocal()
    user = reset_user_password(
        db,
        email="ADMIN@EXAMPLE.COM",
        new_password="Ideiglenes!9",
        activate=True,
    )
    assert user.email == "admin@example.com"
    assert user.is_active is True
    assert user.must_change_password is True
    assert verify_password("Ideiglenes!9", user.password_hash)

    audit = (
        db.query(AuditLog)
        .filter(AuditLog.entity_type == "user", AuditLog.entity_id == str(user.id))
        .order_by(AuditLog.id.desc())
        .first()
    )
    assert audit is not None
    assert audit.user_id is None
    assert audit.action == "jelszó-visszaállítás"
    assert "helyi helyreállító" in audit.description
    db.close()


def test_local_password_reset_rejects_weak_password_and_unknown_user():
    db = TestingSessionLocal()
    try:
        reset_user_password(db, email="admin@example.com", new_password="gyenge")
    except ValueError:
        pass
    else:
        raise AssertionError("A gyenge jelszót el kellett volna utasítani")

    try:
        reset_user_password(db, email="nincs@example.com", new_password="ErosJelszo!9")
    except PasswordResetError:
        pass
    else:
        raise AssertionError("Az ismeretlen felhasználót el kellett volna utasítani")
    db.close()

def test_new_work_order_types_are_supported():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer_id = db.query(Customer).first().id
    location_id = db.query(Location).first().id
    db.close()

    for work_type in ["kiszállítás", "üzembe helyezés"]:
        response = client.post(
            "/api/work-orders",
            headers=headers,
            json={
                "customer_id": customer_id,
                "location_id": location_id,
                "description": f"{work_type} teszt",
                "priority": "normál",
                "status": "új",
                "work_type": work_type,
            },
        )
        assert response.status_code == 201, response.text
        assert response.json()["work_type"] == work_type


def test_version_and_health_endpoints_report_current_release():
    version_response = client.get("/api/version")
    assert version_response.status_code == 200
    assert version_response.json()["version"] == APP_VERSION
    assert version_response.json()["semantic_version"] == semantic_version_parts(APP_VERSION)

    health_response = client.get("/api/health")
    assert health_response.status_code == 200
    assert health_response.json() == {"status": "ok"}


def test_create_asset_and_history():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    customer_id = customer.id
    location_id = location.id
    db.close()
    response = client.post(
        "/api/assets",
        headers=headers,
        json={
            "internal_id": "A-001",
            "serial_number": "SN-001",
            "type": "Teszt gép",
            "status": "aktív",
            "customer_id": customer_id,
            "current_location_id": location_id,
        },
    )
    assert response.status_code == 201, response.text
    asset_id = response.json()["id"]
    history = client.get(f"/api/assets/{asset_id}/history", headers=headers)
    assert history.status_code == 200
    assert history.json()[0]["action"] == "létrehozás"


def test_create_work_order():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    customer_id = customer.id
    location_id = location.id
    db.close()
    response = client.post(
        "/api/work-orders",
        headers=headers,
        json={"customer_id": customer_id, "location_id": location_id, "description": "Teszt munka", "priority": "normál", "status": "új", "work_type": "javítás", "contract_type": "Karbantartási szerződés"},
    )
    assert response.status_code == 201, response.text
    assert response.json()["number"].startswith("ML-")
    assert response.json()["work_type"] == "javítás"
    assert response.json()["contract_type"] == "Karbantartási szerződés"



def test_work_order_rejects_assets_from_other_customer_or_location():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer_one = db.query(Customer).first()
    location_one = db.query(Location).first()
    customer_two = Customer(name="Másik ügyfél", address="Másik cím")
    db.add(customer_two)
    db.flush()
    location_two = Location(customer_id=customer_two.id, name="Másik helyszín", address="Másik telephely")
    db.add(location_two)
    db.commit()
    customer_one_id = customer_one.id
    location_one_id = location_one.id
    customer_two_id = customer_two.id
    location_two_id = location_two.id
    db.close()

    other_customer_asset = client.post(
        "/api/assets",
        headers=headers,
        json={
            "internal_id": "A-OTHER-CUSTOMER",
            "serial_number": "SN-OTHER",
            "type": "Másik ügyfél gépe",
            "status": "aktív",
            "customer_id": customer_two_id,
            "current_location_id": location_two_id,
        },
    )
    assert other_customer_asset.status_code == 201, other_customer_asset.text

    wrong_customer_order = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_one_id,
            "location_id": location_one_id,
            "asset_ids": [other_customer_asset.json()["id"]],
            "description": "Nem kapcsolható másik ügyfél eszköze",
            "priority": "normál",
            "status": "új",
        },
    )
    assert wrong_customer_order.status_code == 400
    assert "nem a kiválasztott ügyfélhez tartozik" in wrong_customer_order.json()["detail"]

    same_customer_other_location_asset = client.post(
        "/api/assets",
        headers=headers,
        json={
            "internal_id": "A-OTHER-LOCATION",
            "serial_number": "SN-LOC",
            "type": "Másik helyszín gépe",
            "status": "aktív",
            "customer_id": customer_two_id,
            "current_location_id": location_two_id,
        },
    )
    assert same_customer_other_location_asset.status_code == 201, same_customer_other_location_asset.text

    wrong_location_order = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_two_id,
            "location_id": location_one_id,
            "asset_ids": [same_customer_other_location_asset.json()["id"]],
            "description": "Nem kapcsolható másik helyszín eszköze",
            "priority": "normál",
            "status": "új",
        },
    )
    assert wrong_location_order.status_code == 400
    assert "helyszín nem a megadott ügyfélhez tartozik" in wrong_location_order.json()["detail"]

def test_customer_can_have_classified_contacts():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    customer_id = customer.id
    location_id = location.id
    db.close()

    response = client.post(
        f"/api/customers/{customer_id}/contacts",
        headers=headers,
        json={
            "category": "Szerviz",
            "name": "Szerviz Kapcsolat",
            "phone": "+36 30 123 4567",
            "email": "szerviz@example.com",
            "location_id": location_id,
            "is_primary": True,
        },
    )
    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["category"] == "Szerviz"
    assert payload["location_name"] == "Teszt helyszín"

    sales_response = client.post(
        f"/api/customers/{customer_id}/contacts",
        headers=headers,
        json={"category": "Sales", "name": "Sales Kapcsolat", "email": "sales@example.com"},
    )
    assert sales_response.status_code == 201, sales_response.text

    customer_response = client.get(f"/api/customers/{customer_id}", headers=headers)
    assert customer_response.status_code == 200, customer_response.text
    categories = {contact["category"] for contact in customer_response.json()["contacts"]}
    assert {"Szerviz", "Sales"}.issubset(categories)


def test_work_order_uses_selected_customer_contact():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    contact = CustomerContact(customer_id=customer.id, location_id=location.id, category="Szerviz", name="Munkalap Kontakt", phone="+36 1 555 0000", email="kontakt@example.com")
    db.add(contact)
    db.commit()
    customer_id = customer.id
    location_id = location.id
    contact_id = contact.id
    db.close()

    response = client.post(
        "/api/work-orders",
        headers=headers,
        json={"customer_id": customer_id, "location_id": location_id, "contact_id": contact_id, "description": "Kontaktos munkalap", "priority": "normál", "status": "új"},
    )
    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["contact_id"] == contact_id
    assert payload["contact"]["category"] == "Szerviz"
    assert payload["customer_contact_person"] == "Munkalap Kontakt"
    assert payload["customer_phone"] == "+36 1 555 0000"


def test_dashboard_summary():
    headers = auth_headers()
    response = client.get("/api/dashboard/summary", headers=headers)
    assert response.status_code == 200
    assert set(response.json().keys()) == {"open_work_orders", "urgent_work_orders", "due_today_work_orders", "due_maintenance_assets", "due_soon_maintenance_assets", "missing_maintenance_assets", "faulty_assets"}


def test_asset_history_shows_linked_work_order_full_content():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    material = db.query(Material).first()
    customer_id = customer.id
    location_id = location.id
    material_id = material.id
    db.close()

    asset_response = client.post(
        "/api/assets",
        headers=headers,
        json={
            "internal_id": "A-WO-001",
            "serial_number": "SN-WO-001",
            "type": "Teszt berendezés",
            "status": "aktív",
            "customer_id": customer_id,
            "current_location_id": location_id,
        },
    )
    assert asset_response.status_code == 201, asset_response.text
    asset_id = asset_response.json()["id"]

    order_response = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "asset_ids": [asset_id],
            "description": "Teljes tartalmú korábbi munkalap teszt",
            "priority": "sürgős",
            "status": "új",
            "work_done": "Diagnosztika és próbaüzem",
            "labor_hours": "1.50",
            "internal_note": "Belső teszt megjegyzés",
            "customer_note": "Ügyfélnek látható teszt megjegyzés",
            "materials": [{"material_id": material_id, "quantity": "2", "unit_price": "100", "note": "Teszt anyag sor"}],
        },
    )
    assert order_response.status_code == 201, order_response.text

    history_orders = client.get(f"/api/assets/{asset_id}/work-orders", headers=headers)
    assert history_orders.status_code == 200, history_orders.text
    data = history_orders.json()
    assert len(data) == 1
    assert data[0]["description"] == "Teljes tartalmú korábbi munkalap teszt"
    assert data[0]["work_done"] == "Diagnosztika és próbaüzem"
    assert data[0]["internal_note"] == "Belső teszt megjegyzés"
    assert data[0]["customer_note"] == "Ügyfélnek látható teszt megjegyzés"
    assert data[0]["assets"][0]["id"] == asset_id
    assert data[0]["materials"][0]["name"] == "Teszt anyag"


def test_customer_asset_import_preserves_excel_fields():
    from app.importers.customer_assets import import_customer_asset_rows
    from app.models import Asset

    db = TestingSessionLocal()
    admin = db.query(User).filter(User.email == "admin@example.com").first()
    result = import_customer_asset_rows(
        db,
        [
            {
                "company_code": "DRh",
                "customer_name": "Import Teszt Kft.",
                "customer_address": "1111 Budapest, Teszt utca 1.",
                "location_name": "Iroda",
                "location_address": "1111 Budapest, Teszt utca 2.",
                "contact_person": "Teszt Elek",
                "phone": "06-30-123-4567",
                "email": "teszt@example.com",
                "machine_type": "Pitney Bowes Relay 3000",
                "serial_number": "TEST-SERIAL-128",
                "source_row": 2,
            }
        ],
        admin,
    )
    db.commit()
    asset = db.query(Asset).filter(Asset.internal_id == "IMP-20260616-0002").first()
    customer = asset.customer
    location = asset.current_location
    contact = db.query(CustomerContact).filter(CustomerContact.customer_id == customer.id, CustomerContact.name == "Teszt Elek").first()

    assert result.rows == 1
    assert result.assets_created == 1
    assert customer.company_code == "DRh"
    assert customer.contact_person == "Teszt Elek"
    assert contact is not None
    assert contact.category == "Szerviz"
    assert contact.location_id == location.id
    assert location.name == "Iroda"
    assert location.address == "1111 Budapest, Teszt utca 2."
    assert asset.serial_number == "TEST-SERIAL-128"
    assert asset.company_code == "DRh"
    assert asset.import_row == 2
    assert asset.note is None

    # A korábbi verziók automatikus import-megjegyzését és az IMP sort
    # egy ismételt frissítés eltávolítja, de a kézi megjegyzést megőrzi.
    asset.note = (
        "Import forrás: Ügyfél_lista_20260616 másolata.xlsx\n"
        "Excel sor: 2\n"
        "Cég: DRh\n"
        "IMP-20260616-0002\n"
        "Kézi megjegyzés"
    )
    db.commit()
    second_result = import_customer_asset_rows(
        db,
        [
            {
                "company_code": "DRh",
                "customer_name": "Import Teszt Kft.",
                "customer_address": "1111 Budapest, Teszt utca 1.",
                "location_name": "Iroda",
                "location_address": "1111 Budapest, Teszt utca 2.",
                "contact_person": "Teszt Elek",
                "phone": "06-30-123-4567",
                "email": "teszt@example.com",
                "machine_type": "Pitney Bowes Relay 3000",
                "serial_number": "TEST-SERIAL-128",
                "source_row": 2,
            }
        ],
        admin,
    )
    db.commit()
    db.refresh(asset)
    assert second_result.asset_notes_cleared == 1
    assert asset.note == "Kézi megjegyzés"
    db.close()


def test_admin_can_run_customer_asset_import_and_read_status():
    headers = auth_headers()
    status_before = client.get('/api/imports/customer-assets/status', headers=headers)
    assert status_before.status_code == 200, status_before.text
    assert status_before.json()['source_rows'] >= 1

    run_response = client.post('/api/imports/customer-assets/run', headers=headers)
    assert run_response.status_code == 200, run_response.text
    payload = run_response.json()
    assert payload['result']['rows'] >= 1
    assert payload['status']['imported_assets_count'] >= 1
    assert payload['status']['is_imported'] is True


def test_admin_imports_report_missing_private_sources(monkeypatch, tmp_path):
    from app.importers import customer_assets
    from app.services import hr_import

    missing = tmp_path / "missing-private-source.json"
    monkeypatch.setattr(customer_assets, "source_file", lambda: missing)
    monkeypatch.setattr(hr_import, "source_file", lambda: missing)
    headers = auth_headers()

    status = client.get('/api/imports/customer-assets/status', headers=headers)
    assert status.status_code == 200
    assert status.json()['source_rows'] == 0
    assert client.post('/api/imports/customer-assets/run', headers=headers).status_code == 409
    assert client.post('/api/hr/admin/imported-leave/auto-link', headers=headers).status_code == 409


def test_company_profiles_endpoint_contains_known_companies():
    headers = auth_headers()
    response = client.get('/api/settings/companies', headers=headers)
    assert response.status_code == 200, response.text
    data = response.json()
    codes = {row['code'] for row in data}
    assert {'DRH', 'SP', 'GXR'}.issubset(codes)
    drh = next(row for row in data if row['code'] == 'DRH')
    assert drh['tax_number'] == '00000000-0-00'
    assert drh['company_registration_number'] == '00-00-000000'
    assert drh['display_name'] == 'Minta Rendszerház Kft.'
    assert drh['address_for_worksheet'] == '1111 Mintaváros, Példa utca 1.'
    assert drh['phone_for_worksheet'] == '+36 1 000 0000'
    assert drh['email_for_worksheet'] == 'helpdesk@example.com'
    sp = next(row for row in data if row['code'] == 'SP')
    assert sp['phone_for_worksheet'] == '+36 1 000 0001'
    assert sp['email_for_worksheet'] == 'helpdesk@example.com'
    gxr = next(row for row in data if row['code'] == 'GXR')
    assert gxr['phone_for_worksheet'] == '+36 1 000 0002'
    assert gxr['email_for_worksheet'] == 'helpdesk@example.com'


def test_work_order_excel_template_download():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    customer.company_code = "DRh"
    location = db.query(Location).first()
    material = db.query(Material).first()
    customer_id = customer.id
    location_id = location.id
    material_id = material.id
    technician_role = db.query(Role).filter(Role.name == ROLE_TECHNICIAN).one()
    technician = User(
        email="xlsx-technikus@example.com",
        full_name="Excel Teszt Technikus",
        password_hash=get_password_hash("technikus123"),
        role_id=technician_role.id,
    )
    db.add(technician)
    db.commit()
    technician_id = technician.id
    db.close()

    asset_response = client.post(
        "/api/assets",
        headers=headers,
        json={
            "internal_id": "A-XLSX-001",
            "serial_number": "SN-XLSX-001",
            "type": "Teszt gép",
            "status": "aktív",
            "customer_id": customer_id,
            "current_location_id": location_id,
        },
    )
    assert asset_response.status_code == 201, asset_response.text
    asset_id = asset_response.json()["id"]

    order_response = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "asset_ids": [asset_id],
            "description": "Excel sablon export teszt",
            "priority": "normál",
            "status": "új",
            "work_type": "karbantartás",
            "technician_id": technician_id,
            "work_done": "A berendezés tisztítása és beállítása megtörtént.",
            "customer_note": "Ez a korábbi ügyfél-megjegyzés nem jelenhet meg.",
            "materials": [{"material_id": material_id, "quantity": "1", "unit_price": "100"}],
        },
    )
    assert order_response.status_code == 201, order_response.text
    order_payload = order_response.json()
    order_id = order_payload["id"]
    assert order_payload["supplier"]["code"] == "DRH"
    assert order_payload["supplier"]["tax_number"] == "00000000-0-00"
    assert order_payload["customer_company_code"] == "DRh"
    assert order_payload["work_type"] == "karbantartás"

    detail_response = client.get(f"/api/work-orders/{order_id}", headers=headers)
    assert detail_response.status_code == 200, detail_response.text
    detail_payload = detail_response.json()
    assert detail_payload["supplier"]["short_name"] == "Minta Rendszerház Kft."
    assert detail_payload["supplier"]["address_for_worksheet"] == "1111 Mintaváros, Példa utca 1."
    assert detail_payload["assets"][0]["type"] == "Teszt gép"
    assert detail_payload["assets"][0]["serial_number"] == "SN-XLSX-001"
    assert detail_payload["customer_address"] is not None or detail_payload["location_address"] is not None

    export_response = client.get(f"/api/work-orders/{order_id}/worksheet.xlsx", headers=headers)
    assert export_response.status_code == 200, export_response.text
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in export_response.headers["content-type"]
    assert export_response.content.startswith(b"PK")
    assert len(export_response.content) > 5000

    workbook = load_workbook(BytesIO(export_response.content))
    worksheet = workbook.active
    assert worksheet["A4"].value == "Minta Rendszerház Kft."
    assert "Adószám: 00000000-0-00" in worksheet["A5"].value
    assert "Cégjegyzékszám: 00-00-000000" in worksheet["A5"].value
    assert "Cím:" not in worksheet["A5"].value
    assert "Web:" not in worksheet["A5"].value
    assert "helpdesk@example.com" in worksheet["A5"].value
    assert "+36 1 000 0000" in worksheet["A5"].value
    assert worksheet["A10"].value in (None, "")
    assert "Kapcsolattartó:" not in str(worksheet["A10"].value or "")
    assert worksheet["A12"].value == "Telefon: +36 1 000 0000"
    assert worksheet["A14"].value == "E-Mail: helpdesk@example.com"
    assert worksheet["G17"].value == "karbantartás"
    assert worksheet["Q17"].value in (None, "")
    assert worksheet["L38"].value in (None, "")
    assert worksheet["P38"].value in (None, "")
    assert worksheet["A31"].value == "Elvégzett munka:"
    assert worksheet["A33"].value == "A berendezés tisztítása és beállítása megtörtént."
    assert worksheet.row_dimensions[40].hidden is True
    assert worksheet.row_dimensions[43].hidden is True
    workbook_text = " ".join(str(cell.value or "") for row in worksheet.iter_rows() for cell in row)
    assert "Ügyfélnek látható megjegyzés" not in workbook_text
    assert "Ez a korábbi ügyfél-megjegyzés nem jelenhet meg." not in workbook_text
    assert "Felhasznált anyagok" not in workbook_text
    assert "Excel Teszt Technikus" in str(worksheet["A56"].value)


def test_admin_can_manage_technicians_from_settings():
    headers = auth_headers()
    create_response = client.post(
        "/api/settings/technicians",
        headers=headers,
        json={
            "email": "technikus@example.com",
            "full_name": "Teszt Technikus",
            "password": "Technikus!123",
            "is_active": True,
        },
    )
    assert create_response.status_code == 201, create_response.text
    technician = create_response.json()
    assert technician["role"]["name"] == "Technikus / szerelő"

    list_response = client.get("/api/settings/technicians", headers=headers)
    assert list_response.status_code == 200, list_response.text
    assert any(row["email"] == "technikus@example.com" for row in list_response.json())

    update_response = client.patch(
        f"/api/settings/technicians/{technician['id']}",
        headers=headers,
        json={"full_name": "Módosított Technikus", "is_active": False},
    )
    assert update_response.status_code == 200, update_response.text
    assert update_response.json()["full_name"] == "Módosított Technikus"
    assert update_response.json()["is_active"] is False


def test_full_backup_and_restore_replaces_database_content():
    headers = auth_headers()
    backup_response = client.get("/api/settings/backup", headers=headers)
    assert backup_response.status_code == 200, backup_response.text
    assert backup_response.headers["content-type"].startswith("application/zip")
    assert f"v{APP_VERSION}" in backup_response.headers["content-disposition"]
    with zipfile.ZipFile(BytesIO(backup_response.content), "r") as archive:
        manifest = json.loads(archive.read("manifest.json"))
    assert manifest["application_version"] == APP_VERSION

    db = TestingSessionLocal()
    db.add(Customer(name="Mentés után létrehozott ügyfél"))
    db.commit()
    assert db.query(Customer).filter(Customer.name == "Mentés után létrehozott ügyfél").count() == 1
    db.close()

    restore_response = client.post(
        "/api/settings/restore",
        headers=headers,
        data={"confirmation": "VISSZAÁLLÍTÁS"},
        files={"backup_file": ("backup.zip", backup_response.content, "application/zip")},
    )
    assert restore_response.status_code == 200, restore_response.text
    assert restore_response.json()["restart_required"] is True
    assert restore_response.json()["backup_application_version"] == APP_VERSION

    db = TestingSessionLocal()
    assert db.query(Customer).filter(Customer.name == "Mentés után létrehozott ügyfél").count() == 0
    assert db.query(User).filter(User.email == "admin@example.com").count() == 1
    db.close()




def test_restore_accepts_backup_without_secondary_technician_column():
    headers = auth_headers()
    backup_response = client.get("/api/settings/backup", headers=headers)
    assert backup_response.status_code == 200, backup_response.text

    with zipfile.ZipFile(BytesIO(backup_response.content), "r") as source:
        files = {name: source.read(name) for name in source.namelist()}

    database_payload = json.loads(files["database.json"])
    work_orders = database_payload["tables"]["work_orders"]
    work_orders["columns"].remove("secondary_technician_id")
    for row in work_orders["rows"]:
        row.pop("secondary_technician_id", None)
    files["database.json"] = json.dumps(database_payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")

    manifest = json.loads(files["manifest.json"])
    manifest["checksums"]["database.json"] = hashlib.sha256(files["database.json"]).hexdigest()
    files["manifest.json"] = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")

    legacy_backup = BytesIO()
    with zipfile.ZipFile(legacy_backup, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for name, content in files.items():
            target.writestr(name, content)

    restore_response = client.post(
        "/api/settings/restore",
        headers=headers,
        data={"confirmation": "VISSZAÁLLÍTÁS"},
        files={"backup_file": ("legacy-backup.zip", legacy_backup.getvalue(), "application/zip")},
    )
    assert restore_response.status_code == 200, restore_response.text


def test_deleting_technician_clears_secondary_work_order_assignment():
    headers = auth_headers()
    technician = client.post(
        "/api/settings/technicians",
        headers=headers,
        json={
            "email": "secondary-delete@example.com",
            "full_name": "Törlendő Második Technikus",
            "password": "DeleteMe!123",
            "is_active": True,
        },
    ).json()

    db = TestingSessionLocal()
    customer_id = db.query(Customer).first().id
    location_id = db.query(Location).first().id
    db.close()
    created = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "description": "Második technikus törlési teszt",
            "secondary_technician_id": technician["id"],
        },
    )
    assert created.status_code == 201, created.text

    deleted = client.delete(f"/api/settings/technicians/{technician['id']}", headers=headers)
    assert deleted.status_code == 204, deleted.text

    loaded = client.get(f"/api/work-orders/{created.json()['id']}", headers=headers)
    assert loaded.status_code == 200, loaded.text
    assert loaded.json()["secondary_technician_id"] is None


def test_admin_can_edit_company_profile(monkeypatch, tmp_path):
    from app.config import get_settings
    from app.services.company_profiles import load_company_profiles

    monkeypatch.setenv("RUNTIME_DIR", str(tmp_path))
    get_settings.cache_clear()
    load_company_profiles.cache_clear()
    try:
        headers = auth_headers()
        response = client.put(
            "/api/settings/companies/DRH",
            headers=headers,
            json={
                "short_name": "Delfin Rendszerház Teszt Kft.",
                "phone": "+36 1 111 2222",
                "worksheet_phone": "+36 1 333 4444 / +36 20 555 6666",
                "worksheet_email": "szerviz@example.com",
                "aliases": ["DRH", "DRh"],
            },
        )
        assert response.status_code == 200, response.text
        assert response.json()["short_name"] == "Delfin Rendszerház Teszt Kft."
        assert response.json()["phone"] == "+36 1 111 2222"
        assert response.json()["phone_for_worksheet"] == "+36 1 333 4444 / +36 20 555 6666"
        assert response.json()["email_for_worksheet"] == "szerviz@example.com"

        list_response = client.get("/api/settings/companies", headers=headers)
        assert list_response.status_code == 200
        drh = next(row for row in list_response.json() if row["code"] == "DRH")
        assert drh["short_name"] == "Delfin Rendszerház Teszt Kft."
        assert drh["worksheet_phone"] == "+36 1 333 4444 / +36 20 555 6666"
        assert drh["worksheet_email"] == "szerviz@example.com"
        assert (tmp_path / "company_profiles.json").exists()
    finally:
        load_company_profiles.cache_clear()
        get_settings.cache_clear()


def test_legacy_runtime_company_profile_uses_general_contacts(monkeypatch, tmp_path):
    from app.config import get_settings
    from app.services.company_profiles import load_company_profiles

    legacy_rows = json.loads((FIXTURE_DIR / 'company_profiles.json').read_text(encoding='utf-8'))
    for row in legacy_rows:
        row.pop('worksheet_phone', None)
        row.pop('worksheet_email', None)
    (tmp_path / 'company_profiles.json').write_text(json.dumps(legacy_rows, ensure_ascii=False), encoding='utf-8')

    monkeypatch.setenv('RUNTIME_DIR', str(tmp_path))
    get_settings.cache_clear()
    load_company_profiles.cache_clear()
    try:
        profiles = load_company_profiles()
        assert profiles['DRH'].phone_for_worksheet == '+36 1 000 0000'
        assert profiles['DRH'].email_for_worksheet == 'iroda@example.com'
        assert profiles['GXR'].phone_for_worksheet == '+36 1 000 0002'
        assert profiles['GXR'].email_for_worksheet == 'szerviz@example.com'
    finally:
        load_company_profiles.cache_clear()
        get_settings.cache_clear()


def test_dashboard_weekly_calendar_uses_planned_and_actual_intervals():
    headers = auth_headers()
    technician_response = client.post(
        "/api/settings/technicians",
        headers=headers,
        json={
            "email": "calendar-tech@example.com",
            "full_name": "Naptár Technikus",
            "password": "Calendar!123",
            "is_active": True,
        },
    )
    assert technician_response.status_code == 201, technician_response.text
    technician_id = technician_response.json()["id"]

    db = TestingSessionLocal()
    customer_id = db.query(Customer).first().id
    location_id = db.query(Location).first().id
    db.close()

    planned_response = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "technician_id": technician_id,
            "description": "Tervezett naptárbejegyzés",
            "status": "ütemezve",
            "priority": "normál",
            "planned_start_at": "2026-07-13T08:30:00",
            "planned_end_at": "2026-07-13T10:00:00",
        },
    )
    assert planned_response.status_code == 201, planned_response.text
    assert planned_response.json()["planned_date"] == "2026-07-13"

    actual_response = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "technician_id": technician_id,
            "description": "Lezárt naptárbejegyzés",
            "status": "lezárva",
            "priority": "sürgős",
            "planned_date": "2026-07-14",
            "started_at": "2026-07-14T11:00:00",
            "completed_at": "2026-07-14T12:30:00",
        },
    )
    assert actual_response.status_code == 201, actual_response.text

    calendar_response = client.get("/api/dashboard/calendar?week_start=2026-07-13", headers=headers)
    assert calendar_response.status_code == 200, calendar_response.text
    payload = calendar_response.json()
    assert payload["week_start"] == "2026-07-13"
    assert len(payload["days"]) == 7
    by_description = {event["description"]: event for event in payload["events"]}
    assert by_description["Tervezett naptárbejegyzés"]["duration_source"] == "planned"
    assert by_description["Tervezett naptárbejegyzés"]["start"] == "2026-07-13T08:30:00"
    assert by_description["Lezárt naptárbejegyzés"]["duration_source"] == "actual"
    assert by_description["Lezárt naptárbejegyzés"]["is_closed"] is True
    assert any(item["id"] == technician_id for item in payload["technicians"])


def test_dashboard_calendar_includes_only_approved_technician_leave():
    headers = auth_headers()
    technician_response = client.post(
        "/api/settings/technicians",
        headers=headers,
        json={
            "email": "leave-calendar-tech@example.com",
            "full_name": "Szabadságos Technikus",
            "password": "Calendar!123",
            "is_active": True,
        },
    )
    assert technician_response.status_code == 201, technician_response.text
    technician_id = technician_response.json()["id"]

    db = TestingSessionLocal()
    office_role = db.query(Role).filter(Role.name == ROLE_OFFICE).one()
    office_user = User(
        email="leave-calendar-office@example.com",
        full_name="Irodai Szabadság",
        password_hash=get_password_hash("Office!123"),
        role_id=office_role.id,
    )
    db.add(office_user)
    db.flush()
    technician_profile = EmployeeProfile(user_id=technician_id, active=True)
    office_profile = EmployeeProfile(user_id=office_user.id, active=True)
    db.add_all([technician_profile, office_profile])
    db.flush()
    db.add_all([
        LeaveRequest(
            employee_id=technician_profile.id,
            start_date=date(2026, 7, 12),
            end_date=date(2026, 7, 14),
            requested_days=2,
            status="approved",
            leave_type="annual_leave",
        ),
        LeaveRequest(
            employee_id=technician_profile.id,
            start_date=date(2026, 7, 15),
            end_date=date(2026, 7, 15),
            requested_days=1,
            status="pending",
            leave_type="annual_leave",
        ),
        LeaveRequest(
            employee_id=office_profile.id,
            start_date=date(2026, 7, 16),
            end_date=date(2026, 7, 16),
            requested_days=1,
            status="approved",
            leave_type="annual_leave",
        ),
    ])
    db.commit()
    db.close()

    response = client.get("/api/dashboard/calendar?week_start=2026-07-13", headers=headers)
    assert response.status_code == 200, response.text
    leave_events = response.json()["leave_events"]
    assert leave_events == [{
        "id": leave_events[0]["id"],
        "technician_id": technician_id,
        "technician_name": "Szabadságos Technikus",
        "leave_type": "annual_leave",
        "start_date": "2026-07-12",
        "end_date": "2026-07-14",
    }]


def test_customer_location_and_contact_changes_are_admin_only():
    admin_headers = auth_headers()
    db = TestingSessionLocal()
    office_role = db.query(Role).filter(Role.name == ROLE_OFFICE).one()
    office_user = User(
        email="customer-office@example.com",
        full_name="Ügyfél Irodai",
        password_hash=get_password_hash("Office!123"),
        role_id=office_role.id,
    )
    db.add(office_user)
    grant_service_access(db, office_user)
    customer = db.query(Customer).first()
    location = db.query(Location).filter(Location.customer_id == customer.id).first()
    customer_id, location_id = customer.id, location.id
    db.commit()
    db.close()
    office_headers = auth_headers_for("customer-office@example.com", "Office!123")

    contact_response = client.post(
        f"/api/customers/{customer_id}/contacts",
        headers=admin_headers,
        json={"name": "Admin Kontakt", "category": "Szerviz", "location_id": location_id},
    )
    assert contact_response.status_code == 201, contact_response.text
    contact_id = contact_response.json()["id"]

    assert client.put(f"/api/customers/{customer_id}", headers=office_headers, json={"name": "Tiltott"}).status_code == 403
    assert client.put(f"/api/locations/{location_id}", headers=office_headers, json={"name": "Tiltott"}).status_code == 403
    assert client.put(
        f"/api/customers/{customer_id}/contacts/{contact_id}",
        headers=office_headers,
        json={"name": "Tiltott"},
    ).status_code == 403

    customer_update = client.put(
        f"/api/customers/{customer_id}", headers=admin_headers,
        json={"name": "Szerkesztett ügyfél", "address": "Új ügyfélcím", "note": "Admin módosítás"},
    )
    assert customer_update.status_code == 200, customer_update.text
    assert customer_update.json()["address"] == "Új ügyfélcím"

    location_update = client.put(
        f"/api/locations/{location_id}", headers=admin_headers,
        json={"name": "Szerkesztett helyszín", "address": "Új helyszíncím", "note": "Admin módosítás"},
    )
    assert location_update.status_code == 200, location_update.text
    assert location_update.json()["address"] == "Új helyszíncím"

    contact_update = client.put(
        f"/api/customers/{customer_id}/contacts/{contact_id}", headers=admin_headers,
        json={"name": "Szerkesztett kontakt", "phone": "+36 30 123 4567"},
    )
    assert contact_update.status_code == 200, contact_update.text
    assert contact_update.json()["phone"] == "+36 30 123 4567"

    assert client.delete(f"/api/locations/{location_id}", headers=admin_headers).status_code == 400
    assert client.delete(f"/api/customers/{customer_id}/contacts/{contact_id}", headers=office_headers).status_code == 403
    assert client.delete(f"/api/customers/{customer_id}/contacts/{contact_id}", headers=admin_headers).status_code == 204
    assert client.delete(f"/api/locations/{location_id}", headers=admin_headers).status_code == 204

    empty_customer = client.post("/api/customers", headers=admin_headers, json={"name": "Törölhető ügyfél"})
    assert empty_customer.status_code == 201, empty_customer.text
    empty_customer_id = empty_customer.json()["id"]
    assert client.delete(f"/api/customers/{empty_customer_id}", headers=office_headers).status_code == 403
    assert client.delete(f"/api/customers/{empty_customer_id}", headers=admin_headers).status_code == 204


def test_inactive_location_keeps_history_but_blocks_new_assignments():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer_id = db.query(Customer).first().id
    db.close()

    location_response = client.post(
        f"/api/customers/{customer_id}/locations",
        headers=headers,
        json={"name": "Megszűnő telephely", "address": "1111 Budapest, Régi út 1."},
    )
    assert location_response.status_code == 201, location_response.text
    location_id = location_response.json()["id"]
    assert location_response.json()["is_active"] is True

    asset = client.post(
        "/api/assets",
        headers=headers,
        json={"internal_id": "INACTIVE-LOC-ASSET", "customer_id": customer_id, "current_location_id": location_id},
    )
    assert asset.status_code == 201, asset.text
    contact = client.post(
        f"/api/customers/{customer_id}/contacts",
        headers=headers,
        json={"name": "Régi telephelyi kontakt", "category": "Szerviz", "location_id": location_id},
    )
    assert contact.status_code == 201, contact.text
    order = client.post(
        "/api/work-orders",
        headers=headers,
        json={"customer_id": customer_id, "location_id": location_id, "description": "Történeti munkalap"},
    )
    assert order.status_code == 201, order.text
    contract = client.post(
        "/api/crm/contracts",
        headers=headers,
        json={
            "customer_id": customer_id,
            "contract_number": "INACTIVE-LOC-CONTRACT",
            "contract_type": "Teszt",
            "location_ids": [location_id],
        },
    )
    assert contract.status_code == 201, contract.text

    deactivated = client.put(f"/api/locations/{location_id}", headers=headers, json={"is_active": False})
    assert deactivated.status_code == 200, deactivated.text
    assert deactivated.json()["is_active"] is False
    customer = client.get(f"/api/customers/{customer_id}", headers=headers).json()
    assert next(row for row in customer["locations"] if row["id"] == location_id)["is_active"] is False
    assert client.delete(f"/api/locations/{location_id}", headers=headers).status_code == 400

    # A meglévő hivatkozások szerkeszthetők, és a helyszín változatlanul megmarad rajtuk.
    kept_asset = client.put(f"/api/assets/{asset.json()['id']}", headers=headers, json={"note": "Megőrzött előzmény"})
    assert kept_asset.status_code == 200, kept_asset.text
    assert kept_asset.json()["current_location_id"] == location_id
    kept_contact = client.put(
        f"/api/customers/{customer_id}/contacts/{contact.json()['id']}",
        headers=headers,
        json={"name": "Régi telephelyi kontakt módosítva"},
    )
    assert kept_contact.status_code == 200, kept_contact.text
    assert kept_contact.json()["location_id"] == location_id
    kept_order = client.put(
        f"/api/work-orders/{order.json()['id']}",
        headers=headers,
        json={"version": order.json()["version"], "description": "Történeti munkalap módosítva"},
    )
    assert kept_order.status_code == 200, kept_order.text
    assert kept_order.json()["location_id"] == location_id
    kept_contract = client.put(
        f"/api/crm/contracts/{contract.json()['id']}",
        headers=headers,
        json={"version": contract.json()["version"], "note": "Megőrzött szerződés"},
    )
    assert kept_contract.status_code == 200, kept_contract.text
    assert kept_contract.json()["location_ids"] == [location_id]

    assert client.post(
        "/api/assets",
        headers=headers,
        json={"internal_id": "INACTIVE-LOC-NEW-ASSET", "customer_id": customer_id, "current_location_id": location_id},
    ).status_code == 400
    assert client.post(
        f"/api/customers/{customer_id}/contacts",
        headers=headers,
        json={"name": "Tiltott új kontakt", "category": "Szerviz", "location_id": location_id},
    ).status_code == 400
    assert client.post(
        "/api/work-orders",
        headers=headers,
        json={"customer_id": customer_id, "location_id": location_id, "description": "Tiltott új munkalap"},
    ).status_code == 400
    assert client.post(
        "/api/crm/contracts",
        headers=headers,
        json={
            "customer_id": customer_id,
            "contract_number": "INACTIVE-LOC-NEW-CONTRACT",
            "contract_type": "Teszt",
            "location_ids": [location_id],
        },
    ).status_code == 400

    reactivated = client.put(f"/api/locations/{location_id}", headers=headers, json={"is_active": True})
    assert reactivated.status_code == 200, reactivated.text
    assert reactivated.json()["is_active"] is True


def test_import_identifiers_are_hidden_from_user_facing_outputs():
    from app.importers.customer_assets import import_customer_asset_rows
    from app.models import Asset

    headers = auth_headers()
    db = TestingSessionLocal()
    admin = db.query(User).filter(User.email == "admin@example.com").first()
    result = import_customer_asset_rows(
        db,
        [
            {
                "company_code": "DRh",
                "customer_name": "Rejtett importazonosító Kft.",
                "customer_address": "1111 Budapest, Rejtett utca 1.",
                "location_name": "Iroda",
                "location_address": "1111 Budapest, Rejtett utca 2.",
                "contact_person": "Teszt Elek",
                "phone": "06-30-123-4567",
                "email": "teszt@example.com",
                "machine_type": "Pitney Bowes Relay 3000",
                "serial_number": "SN-HIDDEN-001",
                "source_row": 211,
            }
        ],
        admin,
    )
    db.commit()
    assert result.assets_created == 1
    asset = db.query(Asset).filter(Asset.import_row == 211).first()
    assert asset.internal_id == "IMP-20260616-0211"
    asset_id = asset.id
    customer_id = asset.customer_id
    location_id = asset.current_location_id
    db.close()

    assets_response = client.get("/api/assets", headers=headers)
    assert assets_response.status_code == 200, assets_response.text
    imported_asset = next(row for row in assets_response.json() if row["id"] == asset_id)
    assert imported_asset["internal_id_hidden"] is True
    assert imported_asset["display_internal_id"] is None
    assert "IMP-" not in imported_asset["display_name"]

    order_response = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "asset_ids": [asset_id],
            "description": "Importazonosító elrejtési teszt",
            "priority": "normál",
            "status": "új",
        },
    )
    assert order_response.status_code == 201, order_response.text
    order = order_response.json()
    assert order["assets"][0]["internal_id"] is None
    assert "IMP-" not in order["assets"][0]["display_name"]

    csv_response = client.get("/api/export/assets.csv", headers=headers)
    assert csv_response.status_code == 200, csv_response.text
    assert "IMP-20260616-0211" not in csv_response.text

    worksheet_response = client.get(f"/api/work-orders/{order['id']}/worksheet.xlsx", headers=headers)
    assert worksheet_response.status_code == 200, worksheet_response.text
    workbook = load_workbook(BytesIO(worksheet_response.content))
    worksheet = workbook.active
    assert "IMP-20260616-0211" not in " ".join(str(cell.value or "") for row in worksheet.iter_rows() for cell in row)


def test_work_order_optimistic_lock_prevents_stale_overwrite():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer_id = db.query(Customer).first().id
    location_id = db.query(Location).first().id
    db.close()

    created = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "description": "Eredeti leírás",
            "priority": "normál",
            "status": "új",
        },
    )
    assert created.status_code == 201, created.text
    original = created.json()
    assert original["version"] == 1

    first_update = client.put(
        f"/api/work-orders/{original['id']}",
        headers=headers,
        json={"version": original["version"], "description": "Első felhasználó módosítása"},
    )
    assert first_update.status_code == 200, first_update.text
    current = first_update.json()
    assert current["version"] == 2

    stale_update = client.put(
        f"/api/work-orders/{original['id']}",
        headers=headers,
        json={"version": original["version"], "description": "Elavult felülírás"},
    )
    assert stale_update.status_code == 409, stale_update.text
    assert "másik felhasználó" in stale_update.json()["detail"]

    detail = client.get(f"/api/work-orders/{original['id']}", headers=headers)
    assert detail.status_code == 200
    assert detail.json()["description"] == "Első felhasználó módosítása"
    assert detail.json()["version"] == 2


def test_work_order_numbers_are_unique_for_multiple_creates():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer_id = db.query(Customer).first().id
    location_id = db.query(Location).first().id
    db.close()

    numbers = []
    for index in range(3):
        response = client.post(
            "/api/work-orders",
            headers=headers,
            json={
                "customer_id": customer_id,
                "location_id": location_id,
                "description": f"Sorszámteszt {index}",
                "priority": "normál",
                "status": "új",
            },
        )
        assert response.status_code == 201, response.text
        numbers.append(response.json()["number"])

    assert len(numbers) == len(set(numbers))
    assert all(number.startswith("ML-") for number in numbers)


def test_database_level_optimistic_lock_catches_near_simultaneous_updates():
    db = TestingSessionLocal()
    customer_id = db.query(Customer).first().id
    order = WorkOrder(
        number="ML-2026-LOCK-000001",
        customer_id=customer_id,
        description="Eredeti adatbázisérték",
    )
    db.add(order)
    db.commit()
    order_id = order.id
    db.close()

    first_session = TestingSessionLocal()
    second_session = TestingSessionLocal()
    try:
        first_copy = first_session.get(WorkOrder, order_id)
        second_copy = second_session.get(WorkOrder, order_id)

        first_copy.description = "Első párhuzamos mentés"
        commit_with_optimistic_lock(first_session)

        second_copy.description = "Második párhuzamos mentés"
        try:
            commit_with_optimistic_lock(second_session)
            assert False, "Az elavult adatbázis-frissítésnek 409 ütközést kellett volna okoznia"
        except HTTPException as exc:
            assert exc.status_code == 409
    finally:
        first_session.close()
        second_session.close()

    verify = TestingSessionLocal()
    assert verify.get(WorkOrder, order_id).description == "Első párhuzamos mentés"
    verify.close()


def test_work_order_edit_presence_warns_both_users():
    admin_headers = auth_headers()
    db = TestingSessionLocal()
    office_role = db.query(Role).filter(Role.name == ROLE_OFFICE).first()
    office_user = User(
        email="office@example.com",
        full_name="Irodai Felhasználó",
        password_hash=get_password_hash("office123"),
        role_id=office_role.id,
    )
    db.add(office_user)
    grant_service_access(db, office_user)
    customer_id = db.query(Customer).first().id
    location_id = db.query(Location).first().id
    db.commit()
    db.close()

    created = client.post(
        "/api/work-orders",
        headers=admin_headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "description": "Szerkesztési jelenlét teszt",
            "priority": "normál",
            "status": "új",
        },
    )
    assert created.status_code == 201, created.text
    order_id = created.json()["id"]

    admin_start = client.post(f"/api/work-orders/{order_id}/editing/start", headers=admin_headers)
    assert admin_start.status_code == 200, admin_start.text
    assert admin_start.json()["active_editors"] == []
    admin_token = admin_start.json()["session_token"]

    office_headers = auth_headers_for("office@example.com", "office123")
    office_start = client.post(f"/api/work-orders/{order_id}/editing/start", headers=office_headers)
    assert office_start.status_code == 200, office_start.text
    assert office_start.json()["has_other_editors"] is True
    assert office_start.json()["active_editors"][0]["email"] == "admin@example.com"
    office_token = office_start.json()["session_token"]

    admin_heartbeat = client.post(
        f"/api/work-orders/{order_id}/editing/heartbeat",
        headers=admin_headers,
        json={"session_token": admin_token},
    )
    assert admin_heartbeat.status_code == 200, admin_heartbeat.text
    assert admin_heartbeat.json()["has_other_editors"] is True
    assert admin_heartbeat.json()["active_editors"][0]["email"] == "office@example.com"

    office_stop = client.delete(
        f"/api/work-orders/{order_id}/editing/{office_token}",
        headers=office_headers,
    )
    assert office_stop.status_code == 204, office_stop.text

    admin_after_stop = client.post(
        f"/api/work-orders/{order_id}/editing/heartbeat",
        headers=admin_headers,
        json={"session_token": admin_token},
    )
    assert admin_after_stop.status_code == 200, admin_after_stop.text
    assert admin_after_stop.json()["active_editors"] == []



def _create_printer_asset(headers, internal_id="PRN-001"):
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    customer_id, location_id = customer.id, location.id
    db.close()
    response = client.post(
        "/api/assets",
        headers=headers,
        json={
            "internal_id": internal_id,
            "serial_number": f"SN-{internal_id}",
            "type": "Nyomtató",
            "manufacturer": "Teszt",
            "model": "Print 1000",
            "category": "Nyomtató",
            "status": "aktív",
            "customer_id": customer_id,
            "current_location_id": location_id,
        },
    )
    assert response.status_code == 201, response.text
    return response.json(), customer_id, location_id


def test_printer_meter_component_and_forecast():
    headers = auth_headers()
    asset, _, _ = _create_printer_asset(headers)

    component_response = client.post(
        f"/api/assets/{asset['id']}/components",
        headers=headers,
        json={
            "name": "Dobegység",
            "part_number": "DR-TEST",
            "expected_life_pages": 7000,
            "last_replacement_at": "2026-01-01T08:00:00",
            "last_replacement_meter": 0,
        },
    )
    assert component_response.status_code == 201, component_response.text

    first = client.post(
        f"/api/assets/{asset['id']}/meter-readings",
        headers=headers,
        json={"value": 1000, "recorded_at": "2026-01-01T08:00:00"},
    )
    assert first.status_code == 201, first.text
    second = client.post(
        f"/api/assets/{asset['id']}/meter-readings",
        headers=headers,
        json={"value": 4000, "recorded_at": "2026-01-11T08:00:00"},
    )
    assert second.status_code == 201, second.text
    payload = second.json()
    assert payload["latest_meter"]["value"] == 4000
    assert payload["average_monthly_usage"] == 9000.0
    component = payload["components"][0]
    assert component["used_pages"] == 4000
    assert component["remaining_pages"] == 3000
    assert component["forecast_status"] == "due_30_days"
    assert component["forecast_at"].startswith("2026-01-21")


def test_detailed_meter_reading_is_exposed_and_added_to_work_order_with_date():
    headers = auth_headers()
    asset, customer_id, location_id = _create_printer_asset(headers, "PRN-DETAIL")

    first = client.post(
        f"/api/assets/{asset['id']}/meter-readings",
        headers=headers,
        json={
            "value": 1200,
            "black_white_value": 1000,
            "color_value": 200,
            "scan_value": 75,
            "recorded_at": "2026-04-05T09:30:00",
        },
    )
    assert first.status_code == 201, first.text
    latest = first.json()["latest_meter"]
    assert latest["value"] == 1200
    assert latest["black_white_value"] == 1000
    assert latest["color_value"] == 200
    assert latest["scan_value"] == 75

    # A részletes csatornák vezérlő- vagy eszközcsere után újraindulhatnak.
    reset = client.post(
        f"/api/assets/{asset['id']}/meter-readings",
        headers=headers,
        json={
            "value": 1500,
            "black_white_value": 1250,
            "color_value": 250,
            "scan_value": 5,
            "recorded_at": "2026-04-06T09:30:00",
        },
    )
    assert reset.status_code == 201, reset.text

    order_response = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "asset_ids": [asset["id"]],
            "description": "Részletes számláló teszt",
            "priority": "normál",
            "status": "új",
            "work_type": "karbantartás",
        },
    )
    assert order_response.status_code == 201, order_response.text
    order = order_response.json()
    linked_asset = order["assets"][0]
    assert linked_asset["latest_meter_value"] == 1500
    assert linked_asset["latest_meter_black_white_value"] == 1250
    assert linked_asset["latest_meter_color_value"] == 250
    assert linked_asset["latest_meter_scan_value"] == 5
    assert linked_asset["latest_meter_at"].startswith("2026-04-06T09:30:00")

    worksheet_response = client.get(f"/api/work-orders/{order['id']}/worksheet.xlsx", headers=headers)
    assert worksheet_response.status_code == 200, worksheet_response.text
    workbook = load_workbook(BytesIO(worksheet_response.content), data_only=True)
    worksheet = workbook.active
    assert "2026.04.06. 09:30" in str(worksheet["B25"].value)
    assert "Összes / FF / Színes / Scan" in str(worksheet["B25"].value)
    assert "1500 / 1250 / 250 / 5" in str(worksheet["F25"].value)
    workbook.close()


def test_same_timestamp_total_reading_can_be_enriched_with_detailed_channels():
    headers = auth_headers()
    asset, _, _ = _create_printer_asset(headers, "PRN-ENRICH")
    timestamp = "2026-05-01T08:00:00"
    base = client.post(
        f"/api/assets/{asset['id']}/meter-readings",
        headers=headers,
        json={"value": 5000, "recorded_at": timestamp},
    )
    assert base.status_code == 201, base.text
    enriched = client.post(
        f"/api/assets/{asset['id']}/meter-readings",
        headers=headers,
        json={
            "value": 5000,
            "black_white_value": 4200,
            "color_value": 800,
            "scan_value": 120,
            "recorded_at": timestamp,
        },
    )
    assert enriched.status_code == 201, enriched.text
    assert enriched.json()["latest_meter"]["color_value"] == 800
    db = TestingSessionLocal()
    rows = db.query(AssetMeterReading).filter(AssetMeterReading.asset_id == asset["id"]).all()
    assert len(rows) == 1
    assert (rows[0].black_white_value, rows[0].color_value, rows[0].scan_value) == (4200, 800, 120)
    db.close()


def test_component_replacement_can_be_recorded_from_work_order():
    headers = auth_headers()
    asset, customer_id, location_id = _create_printer_asset(headers, "PRN-002")
    component_response = client.post(
        f"/api/assets/{asset['id']}/components",
        headers=headers,
        json={"name": "Fekete toner", "part_number": "TN-TEST", "expected_life_pages": 5000},
    )
    component_id = component_response.json()["components"][0]["id"]

    order_response = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "asset_ids": [asset["id"]],
            "description": "Toner csere",
            "priority": "normál",
            "status": "folyamatban",
            "work_type": "karbantartás",
        },
    )
    assert order_response.status_code == 201, order_response.text
    order_id = order_response.json()["id"]
    db = TestingSessionLocal()
    material_id = db.query(Material).first().id
    db.close()

    replacement = client.post(
        f"/api/assets/{asset['id']}/component-replacements",
        headers=headers,
        json={
            "asset_component_id": component_id,
            "meter_value": 12345,
            "work_order_id": order_id,
            "material_id": material_id,
            "note": "Munkalapról rögzítve",
        },
    )
    assert replacement.status_code == 201, replacement.text
    payload = replacement.json()
    assert payload["latest_meter"]["value"] == 12345
    assert payload["replacements"][0]["work_order_id"] == order_id
    assert payload["replacements"][0]["material_id"] == material_id
    assert payload["components"][0]["last_replacement_meter"] == 12345

    db = TestingSessionLocal()
    assert db.query(AssetMeterReading).filter(AssetMeterReading.asset_id == asset["id"]).count() == 1
    assert db.query(AssetComponentReplacement).filter(AssetComponentReplacement.asset_id == asset["id"]).count() == 1
    db.close()


def test_printer_dashboard_reports_due_and_stale_assets():
    headers = auth_headers()
    asset, _, _ = _create_printer_asset(headers, "PRN-003")
    component_response = client.post(
        f"/api/assets/{asset['id']}/components",
        headers=headers,
        json={"name": "Fixáló", "expected_life_pages": 1000, "last_replacement_meter": 0},
    )
    assert component_response.status_code == 201
    client.post(f"/api/assets/{asset['id']}/meter-readings", headers=headers, json={"value": 900, "recorded_at": "2025-01-01T08:00:00"})
    client.post(f"/api/assets/{asset['id']}/meter-readings", headers=headers, json={"value": 1100, "recorded_at": "2025-01-11T08:00:00"})

    response = client.get("/api/dashboard/printer-maintenance", headers=headers)
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["overdue_count"] >= 1
    assert payload["stale_meter_assets_count"] >= 1
    assert any(item["asset_id"] == asset["id"] for item in payload["due_items"])


def test_asset_list_round_two_filters_and_tracking_summary():
    headers = auth_headers()
    asset, _, _ = _create_printer_asset(headers, "PRN-LIST-001")
    component = client.post(
        f"/api/assets/{asset['id']}/components",
        headers=headers,
        json={"name": "Dobegység", "expected_life_pages": 7000, "last_replacement_meter": 0},
    )
    assert component.status_code == 201, component.text
    client.post(
        f"/api/assets/{asset['id']}/meter-readings",
        headers=headers,
        json={"value": 1000, "recorded_at": "2026-01-01T08:00:00"},
    )
    client.post(
        f"/api/assets/{asset['id']}/meter-readings",
        headers=headers,
        json={"value": 4000, "recorded_at": "2026-01-11T08:00:00"},
    )

    response = client.get("/api/assets?meter_state=stale&forecast_state=due", headers=headers)
    assert response.status_code == 200, response.text
    row = next(item for item in response.json() if item["id"] == asset["id"])
    assert row["latest_meter_value"] == 4000
    assert row["meter_state"] == "stale"
    assert row["component_forecast_status"] == "due_30_days"
    assert row["next_component_name"] == "Dobegység"

    q_response = client.get("/api/assets?q=Teszt%20ügyfél", headers=headers)
    assert q_response.status_code == 200, q_response.text
    assert any(item["id"] == asset["id"] for item in q_response.json())

    options = client.get("/api/assets/filter-options", headers=headers)
    assert options.status_code == 200, options.text
    assert "Teszt" in options.json()["manufacturers"]
    assert "Nyomtató" in options.json()["types"]


def test_asset_list_round_three_server_pagination_and_sorting():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    for index in range(23):
        db.add(
            Asset(
                internal_id=f"PAGE-{index:03d}",
                serial_number=f"SER-{index:03d}",
                type="Lapozott nyomtató",
                manufacturer="Teszt gyártó",
                model=f"Modell {index:03d}",
                category="Lapozási teszt",
                status="aktív" if index % 2 == 0 else "raktáron",
                customer_id=customer.id,
                current_location_id=location.id,
                company_code="DRH" if index < 12 else "SP",
            )
        )
    db.commit()
    db.close()

    response = client.get(
        "/api/assets/paged?category=Lapoz%C3%A1si%20teszt&page=2&page_size=10&sort=serial_number&order=desc",
        headers=headers,
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["total"] == 23
    assert payload["pages"] == 3
    assert payload["page"] == 2
    assert payload["page_size"] == 10
    assert len(payload["items"]) == 10
    serials = [item["serial_number"] for item in payload["items"]]
    assert serials == sorted(serials, reverse=True)
    assert payload["summary"]["total"] == 23
    assert payload["summary"]["active"] == 12

    clamped = client.get(
        "/api/assets/paged?category=Lapoz%C3%A1si%20teszt&page=99&page_size=10",
        headers=headers,
    )
    assert clamped.status_code == 200, clamped.text
    assert clamped.json()["page"] == 3
    assert len(clamped.json()["items"]) == 3


def _meter_import_workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Számláló import"
    sheet.append([
        "import_batch",
        "import_row",
        "gyári szám",
        "cégkód",
        "ügyfél",
        "eszköz",
        "mérés időpontja",
        "számlálóállás",
        "forrás munkalap",
        "forrás sor",
        "FF számlálóállás",
        "színes számlálóállás",
        "scan számlálóállás",
    ])
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_meter_reading_excel_import_preview_run_and_idempotency():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    asset = Asset(
        internal_id="IMP-20260616-0003",
        serial_number="TEST-SERIAL-113",
        type="Pitney Bowes DI 380",
        status="aktív",
        customer_id=customer.id,
        current_location_id=location.id,
        company_code="DRh",
        import_batch="customer_assets_20260616",
        import_row=3,
    )
    db.add(asset)
    db.commit()
    asset_id = asset.id
    db.close()

    workbook_bytes = _meter_import_workbook_bytes([
        ["customer_assets_20260616", 3, "TEST-SERIAL-113", "DRh", "Teszt ügyfél", "Pitney Bowes DI 380", "2024-08-01 12:00", 251214, "2024-08", 4, 200000, 51214, 1000],
        ["customer_assets_20260616", 3, "TEST-SERIAL-113", "DRh", "Teszt ügyfél", "Pitney Bowes DI 380", "2024-09-01 12:00", 252397, "2024-09", 4, 201000, 51397, 50],
    ])
    files = {"import_file": ("szamlalo_adatok_import.xlsx", workbook_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    preview = client.post("/api/imports/meter-readings/preview", headers=headers, files=files)
    assert preview.status_code == 200, preview.text
    preview_payload = preview.json()
    assert preview_payload["would_import_count"] == 2
    assert preview_payload["affected_assets_count"] == 1
    assert preview_payload["conflict_count"] == 0

    files = {"import_file": ("szamlalo_adatok_import.xlsx", workbook_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    imported = client.post("/api/imports/meter-readings/run", headers=headers, files=files)
    assert imported.status_code == 200, imported.text
    assert imported.json()["imported_count"] == 2

    db = TestingSessionLocal()
    readings = db.query(AssetMeterReading).filter(AssetMeterReading.asset_id == asset_id).order_by(AssetMeterReading.recorded_at).all()
    assert [row.value for row in readings] == [251214, 252397]
    assert [(row.black_white_value, row.color_value, row.scan_value) for row in readings] == [(200000, 51214, 1000), (201000, 51397, 50)]
    db.close()

    files = {"import_file": ("szamlalo_adatok_import.xlsx", workbook_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    repeated = client.post("/api/imports/meter-readings/run", headers=headers, files=files)
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()["imported_count"] == 0
    assert repeated.json()["unchanged_count"] == 2


def test_meter_reading_excel_import_reports_conflict_and_missing_asset():
    headers = auth_headers()
    asset, _, _ = _create_printer_asset(headers, "PRN-IMPORT-CONFLICT")
    existing = client.post(
        f"/api/assets/{asset['id']}/meter-readings",
        headers=headers,
        json={"value": 2000, "recorded_at": "2026-02-01T12:00:00"},
    )
    assert existing.status_code == 201, existing.text

    workbook_bytes = _meter_import_workbook_bytes([
        [None, None, asset["serial_number"], None, None, None, "2026-01-01 12:00", 3000, "2026-01", 2],
        [None, None, "NINCS-ILYEN", "DRH", None, None, "2026-01-01 12:00", 100, "2026-01", 3],
    ])
    files = {"import_file": ("szamlalo_adatok_import.xlsx", workbook_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    preview = client.post("/api/imports/meter-readings/preview", headers=headers, files=files)
    assert preview.status_code == 200, preview.text
    payload = preview.json()
    assert payload["conflict_count"] == 1
    assert payload["missing_asset_count"] == 1
    assert payload["would_import_count"] == 0


def test_bulk_customer_maintenance_work_orders_create_one_per_asset_and_skip_duplicates():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    customer_id = customer.id
    location_id = location.id
    db.close()

    for index in range(2):
        response = client.post(
            "/api/assets",
            headers=headers,
            json={
                "internal_id": f"BULK-{index + 1}",
                "serial_number": f"BULK-SN-{index + 1}",
                "type": "Teszt nyomtató",
                "status": "aktív",
                "customer_id": customer_id,
                "current_location_id": location_id,
            },
        )
        assert response.status_code == 201, response.text

    first = client.post(f"/api/customers/{customer_id}/maintenance-work-orders", headers=headers)
    assert first.status_code == 200, first.text
    payload = first.json()
    assert payload["total_assets"] == 2
    assert payload["created_count"] == 2
    assert payload["skipped_count"] == 0
    assert len(payload["created_work_order_ids"]) == 2

    db = TestingSessionLocal()
    orders = db.query(WorkOrder).order_by(WorkOrder.id).all()
    assert len(orders) == 2
    assert all(order.work_type == "karbantartás" for order in orders)
    assert all(order.description == "Tervezett karbantartás" for order in orders)
    assert all(order.status == "új" for order in orders)
    assert all(len(order.asset_links) == 1 for order in orders)
    db.close()

    second = client.post(f"/api/customers/{customer_id}/maintenance-work-orders", headers=headers)
    assert second.status_code == 200, second.text
    payload = second.json()
    assert payload["created_count"] == 0
    assert payload["skipped_count"] == 2


def test_bulk_work_order_update_changes_status_planned_date_and_technician():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    technician_role = db.query(Role).filter(Role.name == ROLE_TECHNICIAN).first()
    technician = User(
        email="bulk-tech@example.com",
        full_name="Tömeges Technikus",
        password_hash=get_password_hash("secret123"),
        role_id=technician_role.id,
        is_active=True,
    )
    db.add(technician)
    db.commit()
    customer_id = customer.id
    location_id = location.id
    technician_id = technician.id
    db.close()

    asset_ids = []
    for index in range(2):
        asset_response = client.post(
            "/api/assets",
            headers=headers,
            json={
                "internal_id": f"BULK-EDIT-ASSET-{index + 1}",
                "serial_number": f"BULK-EDIT-SN-{index + 1}",
                "type": "Teszt nyomtató",
                "status": "aktív",
                "customer_id": customer_id,
                "current_location_id": location_id,
            },
        )
        assert asset_response.status_code == 201, asset_response.text
        asset_ids.append(asset_response.json()["id"])

    created_orders = []
    for index, hour in enumerate((8, 13)):
        response = client.post(
            "/api/work-orders",
            headers=headers,
            json={
                "customer_id": customer_id,
                "location_id": location_id,
                "asset_ids": [asset_ids[index]],
                "description": f"Tömegesen szerkesztendő {index + 1}",
                "priority": "normál",
                "status": "ütemezve",
                "planned_date": "2026-07-20",
                "planned_start_at": f"2026-07-20T{hour:02d}:15:00",
                "planned_end_at": f"2026-07-20T{hour + 1:02d}:45:00",
            },
        )
        assert response.status_code == 201, response.text
        created_orders.append(response.json())

    bulk_response = client.post(
        "/api/work-orders/bulk-update",
        headers=headers,
        json={
            "items": [{"id": order["id"], "version": order["version"]} for order in created_orders],
            "apply_status": True,
            "status": "folyamatban",
            "apply_planned_date": True,
            "planned_date": "2026-08-03",
            "apply_technician": True,
            "technician_id": technician_id,
        },
    )
    assert bulk_response.status_code == 200, bulk_response.text
    assert bulk_response.json()["updated_count"] == 2

    for index, original in enumerate(created_orders):
        refreshed = client.get(f"/api/work-orders/{original['id']}", headers=headers)
        assert refreshed.status_code == 200, refreshed.text
        payload = refreshed.json()
        assert payload["status"] == "folyamatban"
        assert payload["planned_date"] == "2026-08-03"
        assert payload["planned_start_at"].startswith("2026-08-03T")
        assert payload["planned_end_at"].startswith("2026-08-03T")
        assert payload["planned_start_at"][11:16] == original["planned_start_at"][11:16]
        assert payload["planned_end_at"][11:16] == original["planned_end_at"][11:16]
        assert payload["technician_id"] == technician_id
        assert payload["customer_id"] == customer_id
        assert payload["location_id"] == location_id
        assert payload["asset_ids"] == [asset_ids[index]]
        assert payload["version"] == original["version"] + 1


def test_bulk_work_order_status_to_closed_applies_close_side_effects():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    customer_id = customer.id
    location_id = location.id
    db.close()

    asset_response = client.post(
        "/api/assets",
        headers=headers,
        json={
            "internal_id": "BULK-STATUS-CLOSE-ASSET",
            "serial_number": "BULK-STATUS-CLOSE-SN",
            "type": "Teszt nyomtató",
            "status": "aktív",
            "customer_id": customer_id,
            "current_location_id": location_id,
            "last_maintenance_date": "2026-01-10",
        },
    )
    assert asset_response.status_code == 201, asset_response.text
    asset_id = asset_response.json()["id"]

    order_response = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "asset_ids": [asset_id],
            "description": "Tömeges státusz lezárási teszt",
            "priority": "normál",
            "status": "folyamatban",
            "work_type": "karbantartás",
        },
    )
    assert order_response.status_code == 201, order_response.text
    order = order_response.json()

    bulk_response = client.post(
        "/api/work-orders/bulk-update",
        headers=headers,
        json={
            "items": [{"id": order["id"], "version": order["version"]}],
            "apply_status": True,
            "status": "lezárva",
        },
    )
    assert bulk_response.status_code == 200, bulk_response.text
    assert bulk_response.json()["updated_count"] == 1

    refreshed = client.get(f"/api/work-orders/{order['id']}", headers=headers)
    assert refreshed.status_code == 200, refreshed.text
    payload = refreshed.json()
    assert payload["status"] == "lezárva"
    assert payload["completed_at"] is not None
    assert payload["closed_at"] is not None
    assert payload["version"] == order["version"] + 1

    completed_date = datetime.fromisoformat(payload["completed_at"]).date()
    db = TestingSessionLocal()
    asset = db.get(Asset, asset_id)
    assert asset.last_maintenance_date == completed_date
    maintenance_entries = [entry for entry in asset.history if entry.action == "karbantartás elvégezve"]
    assert len(maintenance_entries) == 1
    assert maintenance_entries[0].work_order_id == order["id"]
    db.close()


def test_bulk_work_order_update_is_atomic_on_version_conflict():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    customer_id = customer.id
    location_id = location.id
    db.close()

    orders = []
    for index in range(2):
        response = client.post(
            "/api/work-orders",
            headers=headers,
            json={
                "customer_id": customer_id,
                "location_id": location_id,
                "description": f"Konfliktus teszt {index + 1}",
                "priority": "normál",
                "status": "új",
                "planned_date": "2026-07-01",
            },
        )
        assert response.status_code == 201, response.text
        orders.append(response.json())

    single_update = client.put(
        f"/api/work-orders/{orders[1]['id']}",
        headers=headers,
        json={"version": orders[1]["version"], "priority": "sürgős"},
    )
    assert single_update.status_code == 200, single_update.text

    bulk_response = client.post(
        "/api/work-orders/bulk-update",
        headers=headers,
        json={
            "items": [{"id": order["id"], "version": order["version"]} for order in orders],
            "apply_planned_date": True,
            "planned_date": "2026-09-09",
            "apply_technician": False,
        },
    )
    assert bulk_response.status_code == 409, bulk_response.text

    first_refreshed = client.get(f"/api/work-orders/{orders[0]['id']}", headers=headers).json()
    second_refreshed = client.get(f"/api/work-orders/{orders[1]['id']}", headers=headers).json()
    assert first_refreshed["planned_date"] == "2026-07-01"
    assert second_refreshed["planned_date"] == "2026-07-01"
    assert second_refreshed["priority"] == "sürgős"


def test_closing_maintenance_work_order_updates_linked_assets_last_maintenance_date():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    customer_id = customer.id
    location_id = location.id
    db.close()

    asset_ids = []
    for index in range(2):
        asset_response = client.post(
            "/api/assets",
            headers=headers,
            json={
                "internal_id": f"MAINT-CLOSE-{index + 1}",
                "serial_number": f"MAINT-CLOSE-SN-{index + 1}",
                "type": "Teszt nyomtató",
                "status": "aktív",
                "customer_id": customer_id,
                "current_location_id": location_id,
                "last_maintenance_date": "2026-01-10",
            },
        )
        assert asset_response.status_code == 201, asset_response.text
        asset_ids.append(asset_response.json()["id"])

    order_response = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "asset_ids": asset_ids,
            "description": "Időszakos karbantartás",
            "priority": "normál",
            "status": "folyamatban",
            "work_type": "karbantartás",
        },
    )
    assert order_response.status_code == 201, order_response.text
    order = order_response.json()

    close_response = client.post(
        f"/api/work-orders/{order['id']}/close",
        headers=headers,
        json={
            "version": order["version"],
            "work_done": "Karbantartás elvégezve",
            "final_status": "Üzemképes",
            "labor_hours": 1.5,
            "completed_at": "2026-07-15T14:30:00",
        },
    )
    assert close_response.status_code == 200, close_response.text

    db = TestingSessionLocal()
    assets = db.query(Asset).filter(Asset.id.in_(asset_ids)).all()
    assert all(asset.last_maintenance_date == date(2026, 7, 15) for asset in assets)
    for asset in assets:
        maintenance_entries = [entry for entry in asset.history if entry.action == "karbantartás elvégezve"]
        assert len(maintenance_entries) == 1
        assert maintenance_entries[0].work_order_id == order["id"]
        assert order["number"] in maintenance_entries[0].description
    db.close()


def test_closing_older_maintenance_does_not_move_last_maintenance_date_backwards():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    customer_id = customer.id
    location_id = location.id
    db.close()

    asset_response = client.post(
        "/api/assets",
        headers=headers,
        json={
            "internal_id": "MAINT-NO-ROLLBACK",
            "serial_number": "MAINT-NO-ROLLBACK-SN",
            "type": "Teszt nyomtató",
            "status": "aktív",
            "customer_id": customer_id,
            "current_location_id": location_id,
            "last_maintenance_date": "2026-08-20",
        },
    )
    assert asset_response.status_code == 201, asset_response.text
    asset_id = asset_response.json()["id"]

    order_response = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "asset_ids": [asset_id],
            "description": "Korábbi karbantartás utólagos lezárása",
            "priority": "normál",
            "status": "folyamatban",
            "work_type": "karbantartás",
        },
    )
    order = order_response.json()
    close_response = client.post(
        f"/api/work-orders/{order['id']}/close",
        headers=headers,
        json={
            "version": order["version"],
            "work_done": "Korábbi munka lezárva",
            "final_status": "Üzemképes",
            "labor_hours": 1,
            "completed_at": "2026-07-01T10:00:00",
        },
    )
    assert close_response.status_code == 200, close_response.text

    db = TestingSessionLocal()
    asset = db.get(Asset, asset_id)
    assert asset.last_maintenance_date == date(2026, 8, 20)
    assert any(entry.action == "karbantartás elvégezve" and entry.work_order_id == order["id"] for entry in asset.history)
    db.close()


def _maintenance_import_workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Karbantartás import"
    sheet.append([
        "import_batch",
        "import_row",
        "gyári szám",
        "cégkód",
        "ügyfél",
        "eszköz",
        "utolsó karbantartás dátuma",
        "forrás munkalap",
        "forrás Excel sor",
        "munkalap száma",
        "technikus",
    ])
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_maintenance_date_excel_import_preview_run_and_idempotency():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    asset = Asset(
        internal_id="IMP-20260616-0003",
        serial_number="TEST-SERIAL-113",
        type="Pitney Bowes DI 380",
        status="aktív",
        customer_id=customer.id,
        current_location_id=location.id,
        company_code="DRh",
        import_batch="customer_assets_20260616",
        import_row=3,
    )
    db.add(asset)
    db.commit()
    asset_id = asset.id
    db.close()

    workbook_bytes = _maintenance_import_workbook_bytes([
        ["customer_assets_20260616", 3, "TEST-SERIAL-113", "DRh", "Teszt ügyfél", "Pitney Bowes DI 380", "2026-04-28", "Munka1", 660, "MKL0016362", "Kedves István"],
    ])
    files = {"import_file": ("utolso_karbantartas_import.xlsx", workbook_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    preview = client.post("/api/imports/maintenance-dates/preview", headers=headers, files=files)
    assert preview.status_code == 200, preview.text
    payload = preview.json()
    assert payload["would_update_count"] == 1
    assert payload["affected_assets_count"] == 1
    assert payload["conflict_count"] == 0

    files = {"import_file": ("utolso_karbantartas_import.xlsx", workbook_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    imported = client.post("/api/imports/maintenance-dates/run", headers=headers, files=files)
    assert imported.status_code == 200, imported.text
    assert imported.json()["updated_count"] == 1

    db = TestingSessionLocal()
    refreshed = db.get(Asset, asset_id)
    assert refreshed.last_maintenance_date == date(2026, 4, 28)
    assert any(entry.action == "karbantartási dátum import" for entry in refreshed.history)
    db.close()

    files = {"import_file": ("utolso_karbantartas_import.xlsx", workbook_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    repeated = client.post("/api/imports/maintenance-dates/run", headers=headers, files=files)
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()["updated_count"] == 0
    assert repeated.json()["unchanged_count"] == 1


def test_maintenance_date_import_does_not_overwrite_newer_database_value():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    asset = Asset(
        internal_id="MAINT-CONFLICT-1",
        serial_number="MAINT-SN-1",
        type="Teszt gép",
        status="aktív",
        customer_id=customer.id,
        current_location_id=location.id,
        company_code="DRH",
        last_maintenance_date=date(2026, 6, 1),
    )
    db.add(asset)
    db.commit()
    asset_id = asset.id
    db.close()

    workbook_bytes = _maintenance_import_workbook_bytes([
        [None, None, "MAINT-SN-1", "DRH", "Teszt ügyfél", "Teszt gép", "2026-05-01", "Munka1", 10, "MKL-OLD", "Teszt technikus"],
    ])
    files = {"import_file": ("utolso_karbantartas_import.xlsx", workbook_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    preview = client.post("/api/imports/maintenance-dates/preview", headers=headers, files=files)
    assert preview.status_code == 200, preview.text
    payload = preview.json()
    assert payload["conflict_count"] == 1
    assert payload["would_update_count"] == 0

    files = {"import_file": ("utolso_karbantartas_import.xlsx", workbook_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    imported = client.post("/api/imports/maintenance-dates/run", headers=headers, files=files)
    assert imported.status_code == 200, imported.text
    assert imported.json()["updated_count"] == 0
    assert imported.json()["conflict_count"] == 1

    db = TestingSessionLocal()
    refreshed = db.get(Asset, asset_id)
    assert refreshed.last_maintenance_date == date(2026, 6, 1)
    db.close()


def _maintenance_cycle_raw_workbook_bytes(rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Munka1"
    worksheet.append(["Cég", "Ügyfél", "Telephely cím", "Gép típusa", "Gyáriszám", "Bérelt/üf.saját", "Karb. Ciklus"])
    for row in rows:
        worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_calendar_month_cycle_uses_calendar_boundaries():
    from app.services.maintenance_schedule import add_calendar_months

    assert add_calendar_months(date(2026, 1, 31), 1) == date(2026, 2, 28)
    assert add_calendar_months(date(2024, 1, 31), 1) == date(2024, 2, 29)
    assert add_calendar_months(date(2026, 8, 31), 6) == date(2027, 2, 28)


def test_maintenance_cycle_excel_g_column_import_preview_run_and_idempotency():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    asset = Asset(
        internal_id="IMP-20260616-0211",
        serial_number="CYCLE-SN-001",
        type="Teszt nyomtató",
        status="aktív",
        company_code="DRh",
        import_batch="customer_assets_20260616",
        import_row=211,
        customer_id=customer.id,
        current_location_id=location.id,
        last_maintenance_date=date(2026, 1, 31),
    )
    db.add(asset)
    db.commit()
    asset_id = asset.id
    db.close()

    workbook_bytes = _maintenance_cycle_raw_workbook_bytes([
        ["DRh", "Teszt ügyfél", "Teszt cím", "Teszt nyomtató", "CYCLE-SN-001", "Bérelt", "6 hó"],
    ])
    files = {"import_file": ("karb ciklus.xlsx", workbook_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    preview = client.post("/api/imports/maintenance-cycles/preview", headers=headers, files=files)
    assert preview.status_code == 200, preview.text
    assert preview.json()["would_update_count"] == 1
    assert preview.json()["missing_asset_count"] == 0

    imported = client.post("/api/imports/maintenance-cycles/run", headers=headers, files=files)
    assert imported.status_code == 200, imported.text
    assert imported.json()["updated_count"] == 1

    db = TestingSessionLocal()
    refreshed = db.get(Asset, asset_id)
    assert refreshed.maintenance_cycle_months == 6
    assert refreshed.maintenance_cycle_note is None
    db.close()

    repeated = client.post("/api/imports/maintenance-cycles/run", headers=headers, files=files)
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()["updated_count"] == 0
    assert repeated.json()["unchanged_count"] == 1


def test_contract_cycle_drives_asset_and_dashboard_maintenance_planning():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    asset = Asset(
        internal_id="CYCLE-DASH-001",
        serial_number="CYCLE-DASH-SN",
        type="Ciklusos nyomtató",
        status="aktív",
        customer_id=customer.id,
        current_location_id=location.id,
        last_maintenance_date=date(2020, 1, 31),
        maintenance_cycle_months=6,
    )
    db.add(asset)
    db.commit()
    asset_id = asset.id
    db.close()

    detail = client.get(f"/api/assets/{asset_id}", headers=headers)
    assert detail.status_code == 200, detail.text
    payload = detail.json()
    assert payload["maintenance_cycle_label"] == "6 hónap"
    assert payload["calculated_next_maintenance_date"] == "2020-07-31"
    assert payload["effective_next_maintenance_date"] == "2020-07-31"
    assert payload["maintenance_due_source"] == "contract_cycle"
    assert payload["maintenance_state"] == "overdue"

    planning = client.get("/api/dashboard/maintenance-planning", headers=headers)
    assert planning.status_code == 200, planning.text
    item = next(row for row in planning.json()["items"] if row["asset_id"] == asset_id)
    assert item["maintenance_cycle_label"] == "6 hónap"
    assert item["maintenance_state"] == "overdue"


def _global_import_workbook_bytes():
    workbook = Workbook()
    assets = workbook.active
    assets.title = "Eszközök"
    assets.append([
        "Cégkód", "Ügyfél neve", "Ügyfél címe", "Helyszín neve", "Helyszín címe",
        "Kapcsolattartó neve", "Kapcsolattartó típusa", "Kapcsolattartó beosztása",
        "Kapcsolattartó telefon", "Kapcsolattartó email", "Kapcsolattartó elsődleges",
        "Belső azonosító", "Gyári szám", "Eszköz típusa", "Gyártó", "Modell", "Kategória",
        "Állapot", "Belső használat", "Vásárlás dátuma", "Garancia lejárata",
        "Utolsó karbantartás", "Karbantartási ciklus hónap",
        "Karbantartási ciklus megjegyzés", "Következő karbantartás", "Eszköz megjegyzés",
    ])
    assets.append([
        "DRH", "Globál Import Kft.", "1111 Budapest, Import utca 1.", "Központ",
        "1111 Budapest, Import utca 1.", "Szerviz Elek", "Szerviz", "üzemeltető",
        "+36 30 111 2222", "szerviz@global.test", "Igen", "GLOBAL-001", "GLOBAL-SN-001",
        "Multifunkciós nyomtató", "Ricoh", "IM C3000", "Nyomtató / MFP", "aktív", "Nem",
        "2024-01-15", "2027-01-15", "2026-01-15", 3, None, None, "Globális import teszt",
    ])

    contacts = workbook.create_sheet("Kapcsolattartók")
    contacts.append([
        "Cégkód", "Ügyfél neve", "Helyszín neve", "Helyszín címe", "Kapcsolattartó típusa",
        "Név", "Beosztás", "Telefon", "Email", "Elsődleges", "Megjegyzés",
    ])
    contacts.append([
        "DRH", "Globál Import Kft.", "Központ", "1111 Budapest, Import utca 1.", "Sales",
        "Értékesítő Anna", "kapcsolattartó", "+36 30 333 4444", "sales@global.test", "Igen", None,
    ])

    meters = workbook.create_sheet("Számlálók")
    meters.append([
        "Cégkód", "Ügyfél neve", "Helyszín neve", "Belső azonosító", "Gyári szám",
        "Mérés időpontja", "Számlálóállás", "FF számlálóállás", "Színes számlálóállás",
        "Scan számlálóállás", "Munkalapszám", "Megjegyzés",
    ])
    meters.append(["DRH", "Globál Import Kft.", "Központ", "GLOBAL-001", "GLOBAL-SN-001", "2026-02-01 12:00", 1000, 800, 200, 25, None, None])
    meters.append(["DRH", "Globál Import Kft.", "Központ", "GLOBAL-001", "GLOBAL-SN-001", "2026-03-01 12:00", 1200, 950, 250, 5, None, None])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_global_import_template_download_and_single_workbook_import():
    headers = auth_headers()
    template = client.get("/api/imports/global/template", headers=headers)
    assert template.status_code == 200, template.text
    assert "spreadsheetml" in template.headers["content-type"]
    template_workbook = load_workbook(BytesIO(template.content), read_only=True)
    assert {"Útmutató", "Eszközök", "Kapcsolattartók", "Számlálók", "Példák"}.issubset(template_workbook.sheetnames)
    assert [cell.value for cell in template_workbook["Számlálók"][1]][6:10] == [
        "Számlálóállás", "FF számlálóállás", "Színes számlálóállás", "Scan számlálóállás"
    ]
    template_workbook.close()

    workbook_bytes = _global_import_workbook_bytes()
    files = {"import_file": ("globalis_import.xlsx", workbook_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    preview = client.post("/api/imports/global/preview", headers=headers, files=files)
    assert preview.status_code == 200, preview.text
    payload = preview.json()
    assert payload["asset_rows"] == 1
    assert payload["contact_rows"] == 1
    assert payload["meter_rows"] == 2
    assert payload["assets_created"] == 1
    assert payload["contacts_created"] == 2
    assert payload["maintenance_dates_updated"] == 1
    assert payload["maintenance_cycles_updated"] == 1
    assert payload["meter_readings_would_import"] == 2
    assert payload["skipped_count"] == 0

    files = {"import_file": ("globalis_import.xlsx", workbook_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    imported = client.post("/api/imports/global/run", headers=headers, files=files)
    assert imported.status_code == 200, imported.text
    result = imported.json()
    assert result["assets_created"] == 1
    assert result["meter_readings_imported"] == 2

    db = TestingSessionLocal()
    asset = db.query(Asset).filter(Asset.internal_id == "GLOBAL-001").one()
    assert asset.serial_number == "GLOBAL-SN-001"
    assert asset.last_maintenance_date == date(2026, 1, 15)
    assert asset.maintenance_cycle_months == 3
    assert asset.customer.name == "Globál Import Kft."
    assert asset.current_location.name == "Központ"
    assert db.query(CustomerContact).filter(CustomerContact.customer_id == asset.customer_id).count() == 2
    imported_readings = db.query(AssetMeterReading).filter(AssetMeterReading.asset_id == asset.id).order_by(AssetMeterReading.recorded_at).all()
    assert [row.value for row in imported_readings] == [1000, 1200]
    assert [(row.black_white_value, row.color_value, row.scan_value) for row in imported_readings] == [(800, 200, 25), (950, 250, 5)]
    db.close()

    files = {"import_file": ("globalis_import.xlsx", workbook_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    repeated = client.post("/api/imports/global/run", headers=headers, files=files)
    assert repeated.status_code == 200, repeated.text
    repeated_payload = repeated.json()
    assert repeated_payload["assets_created"] == 0
    assert repeated_payload["assets_unchanged"] == 1
    assert repeated_payload["meter_readings_imported"] == 0
    assert repeated_payload["meter_readings_unchanged"] == 2


def test_global_database_export_matches_import_template_and_roundtrips():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    customer.name = "Export Teszt Kft."
    customer.company_code = "DRH"
    customer.address = "1111 Budapest, Export utca 1."
    location = db.query(Location).first()
    location.name = "Export telephely"
    location.address = "1111 Budapest, Export utca 2."
    contact = CustomerContact(
        customer_id=customer.id,
        location_id=location.id,
        category="Szerviz",
        name="Export Elek",
        title="üzemeltető",
        phone="+36 30 123 4567",
        email="export@example.com",
        is_primary=True,
        note="Kapcsolattartó megjegyzés",
    )
    asset = Asset(
        internal_id="EXPORT-001",
        serial_number="EXPORT-SN-001",
        type="Multifunkciós nyomtató",
        manufacturer="Ricoh",
        model="IM C3000",
        category="Nyomtató / MFP",
        status="aktív",
        customer_id=customer.id,
        current_location_id=location.id,
        company_code="DRH",
        purchase_date=date(2024, 1, 15),
        warranty_expiry=date(2027, 1, 15),
        last_maintenance_date=date(2026, 1, 20),
        maintenance_cycle_months=3,
        maintenance_cycle_note="Szerződés szerint",
        next_maintenance_date=date(2026, 4, 20),
        note="Eszköz megjegyzés",
    )
    db.add_all([contact, asset])
    db.flush()
    reading = AssetMeterReading(
        asset_id=asset.id,
        value=123456,
        black_white_value=100000,
        color_value=23456,
        scan_value=7890,
        recorded_at=datetime(2026, 2, 1, 12, 30),
        recorded_by_user_id=db.query(User).filter(User.email == "admin@example.com").one().id,
        note="Havi számláló",
    )
    db.add(reading)
    db.commit()
    db.close()

    response = client.get("/api/imports/global/export", headers=headers)
    assert response.status_code == 200, response.text
    assert "spreadsheetml" in response.headers["content-type"]
    assert "globalis_adat_export_" in response.headers["content-disposition"]

    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert {"Útmutató", "Eszközök", "Kapcsolattartók", "Számlálók"}.issubset(workbook.sheetnames)
    assert "Példák" not in workbook.sheetnames

    asset_sheet = workbook["Eszközök"]
    asset_headers = [cell.value for cell in asset_sheet[1]]
    asset_row = {asset_headers[index]: value for index, value in enumerate(next(asset_sheet.iter_rows(min_row=2, values_only=True)))}
    assert asset_row["Cégkód"] == "DRH"
    assert asset_row["Ügyfél neve"] == "Export Teszt Kft."
    assert asset_row["Belső azonosító"] == "EXPORT-001"
    assert asset_row["Utolsó karbantartás"] == datetime(2026, 1, 20, 0, 0)
    assert asset_row["Karbantartási ciklus hónap"] == 3

    contact_sheet = workbook["Kapcsolattartók"]
    contact_values = list(contact_sheet.iter_rows(min_row=2, values_only=True))
    assert any(row[5] == "Export Elek" and row[8] == "export@example.com" for row in contact_values)

    meter_sheet = workbook["Számlálók"]
    meter_headers = [cell.value for cell in meter_sheet[1]]
    assert meter_headers[6:10] == ["Számlálóállás", "FF számlálóállás", "Színes számlálóállás", "Scan számlálóállás"]
    meter_values = list(meter_sheet.iter_rows(min_row=2, values_only=True))
    assert any(row[3] == "EXPORT-001" and row[6:10] == (123456, 100000, 23456, 7890) for row in meter_values)
    workbook.close()

    files = {
        "import_file": (
            "globalis_adat_export.xlsx",
            response.content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    preview = client.post("/api/imports/global/preview", headers=headers, files=files)
    assert preview.status_code == 200, preview.text
    payload = preview.json()
    assert payload["skipped_count"] == 0
    assert payload["assets_created"] == 0
    assert payload["meter_readings_unchanged"] >= 1


def test_work_order_history_combines_audit_and_asset_events():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    customer_id = customer.id
    location_id = location.id
    db.close()

    asset_response = client.post(
        "/api/assets",
        headers=headers,
        json={
            "internal_id": "HISTORY-ASSET-1",
            "serial_number": "HISTORY-SN-1",
            "type": "Teszt nyomtató",
            "status": "aktív",
            "customer_id": customer_id,
            "current_location_id": location_id,
        },
    )
    assert asset_response.status_code == 201, asset_response.text

    order_response = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "asset_ids": [asset_response.json()["id"]],
            "description": "Történet ellenőrzése",
            "priority": "normál",
            "status": "új",
            "work_type": "javítás",
        },
    )
    assert order_response.status_code == 201, order_response.text
    order = order_response.json()

    history_response = client.get(f"/api/work-orders/{order['id']}/history", headers=headers)
    assert history_response.status_code == 200, history_response.text
    history = history_response.json()
    assert any(entry["source"] == "audit" and entry["action"] == "létrehozás" for entry in history)
    assert any(entry["source"] == "asset_history" and entry["asset_id"] == asset_response.json()["id"] for entry in history)


def test_close_work_order_stores_actual_start_and_end_time():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    customer_id = customer.id
    location_id = location.id
    db.close()

    order_response = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "description": "Külön lezárási oldal teszt",
            "priority": "normál",
            "status": "folyamatban",
            "work_type": "javítás",
        },
    )
    order = order_response.json()
    response = client.post(
        f"/api/work-orders/{order['id']}/close",
        headers=headers,
        json={
            "version": order["version"],
            "work_done": "Javítás elkészült",
            "final_status": "Üzemképes",
            "labor_hours": 1.5,
            "started_at": "2026-07-23T08:00:00",
            "completed_at": "2026-07-23T09:30:00",
        },
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["started_at"].startswith("2026-07-23T08:00:00")
    assert payload["completed_at"].startswith("2026-07-23T09:30:00")


def test_work_order_paged_list_supports_server_sort_filter_and_pagination():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    customer_id = customer.id
    location_id = location.id
    db.close()

    created = []
    for index in range(30):
        response = client.post(
            "/api/work-orders",
            headers=headers,
            json={
                "customer_id": customer_id,
                "location_id": location_id,
                "description": f"Lapozási feladat {index:02d}",
                "priority": "sürgős" if index % 2 == 0 else "normál",
                "status": "lezárva" if index < 2 else "új",
                "work_type": "karbantartás" if index % 3 == 0 else "javítás",
                "planned_date": f"2026-08-{index + 1:02d}",
            },
        )
        assert response.status_code == 201, response.text
        created.append(response.json())

    response = client.get(
        "/api/work-orders/paged?page=2&page_size=25&sort=planned_date&order=asc&q=Lapoz%C3%A1si",
        headers=headers,
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["total"] == 30
    assert payload["page"] == 2
    assert payload["pages"] == 2
    assert len(payload["items"]) == 5
    assert [item["planned_date"] for item in payload["items"]] == [
        "2026-08-26", "2026-08-27", "2026-08-28", "2026-08-29", "2026-08-30"
    ]

    open_response = client.get(
        "/api/work-orders/paged?open_state=open&work_type=karbantart%C3%A1s&page_size=25",
        headers=headers,
    )
    assert open_response.status_code == 200
    open_items = open_response.json()["items"]
    assert open_items
    assert all(item["status"] not in {"lezárva", "törölve / sztornózva"} for item in open_items)
    assert all(item["work_type"] == "karbantartás" for item in open_items)


def test_work_order_saved_views_are_user_specific_and_validate_state():
    admin_headers = auth_headers()
    db = TestingSessionLocal()
    office_role = db.query(Role).filter(Role.name == ROLE_OFFICE).one()
    office = User(
        email="office.views@example.com",
        full_name="Office Views",
        password_hash=get_password_hash("office123"),
        role_id=office_role.id,
    )
    db.add(office)
    grant_service_access(db, office)
    db.commit()
    db.close()
    office_headers = auth_headers_for("office.views@example.com", "office123")

    create_response = client.post(
        "/api/work-orders/saved-views",
        headers=admin_headers,
        json={
            "name": "Nyitott karbantartások",
            "is_default": True,
            "state": {
                "open_state": "open",
                "work_type": "karbantartás",
                "page_size": 100,
                "sort": "planned_date",
                "order": "asc",
                "columns": ["status", "number", "customer", "unknown-column"],
                "ignored": "value",
            },
        },
    )
    assert create_response.status_code == 201, create_response.text
    saved = create_response.json()
    assert saved["is_default"] is True
    assert saved["state"]["columns"] == ["status", "number", "customer"]
    assert "ignored" not in saved["state"]

    admin_list = client.get("/api/work-orders/saved-views", headers=admin_headers)
    office_list = client.get("/api/work-orders/saved-views", headers=office_headers)
    assert admin_list.status_code == 200
    assert len(admin_list.json()) == 1
    assert office_list.status_code == 200
    assert office_list.json() == []

    duplicate = client.post(
        "/api/work-orders/saved-views",
        headers=admin_headers,
        json={"name": "Nyitott karbantartások", "state": {}},
    )
    assert duplicate.status_code == 409

    update = client.put(
        f"/api/work-orders/saved-views/{saved['id']}",
        headers=admin_headers,
        json={"name": "Karbantartási terv", "state": {"page_size": 999, "sort": "invalid", "order": "invalid"}},
    )
    assert update.status_code == 200, update.text
    updated = update.json()
    assert updated["name"] == "Karbantartási terv"
    assert updated["state"]["page_size"] == 50
    assert updated["state"]["sort"] == "created_at"
    assert updated["state"]["order"] == "desc"

    delete = client.delete(f"/api/work-orders/saved-views/{saved['id']}", headers=admin_headers)
    assert delete.status_code == 204
    assert client.get("/api/work-orders/saved-views", headers=admin_headers).json() == []


def _create_archive_test_asset(headers, suffix="ARCH"):
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    customer_id = customer.id
    location_id = location.id
    db.close()
    response = client.post(
        "/api/assets",
        headers=headers,
        json={
            "internal_id": f"A-{suffix}",
            "serial_number": f"SN-{suffix}",
            "type": "Archiválandó eszköz",
            "status": "aktív",
            "customer_id": customer_id,
            "current_location_id": location_id,
        },
    )
    assert response.status_code == 201, response.text
    return response.json(), customer_id, location_id


def test_work_order_archive_upload_links_work_order_assets_and_downloads():
    headers = auth_headers()
    asset, customer_id, location_id = _create_archive_test_asset(headers, "WO-ARCH")
    work_order = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "asset_ids": [asset["id"]],
            "description": "Archiválási teszt",
            "priority": "normál",
            "status": "új",
        },
    )
    assert work_order.status_code == 201, work_order.text
    order = work_order.json()
    pdf_bytes = b"%PDF-1.4\n% archive test\n%%EOF\n"
    upload = client.post(
        f"/api/work-orders/{order['id']}/archives",
        headers=headers,
        data={"source_number": "PAPIR-2020-17", "work_date": "2020-05-04", "note": "Régi aláírt példány"},
        files={"archive_file": ("munkalap.pdf", pdf_bytes, "application/pdf")},
    )
    assert upload.status_code == 201, upload.text
    archive = upload.json()
    assert archive["archive_number"].startswith("AR-")
    assert archive["work_order_id"] == order["id"]
    assert archive["asset_ids"] == [asset["id"]]
    assert archive["source_number"] == "PAPIR-2020-17"

    order_archives = client.get(f"/api/work-orders/{order['id']}/archives", headers=headers)
    assert order_archives.status_code == 200
    assert order_archives.json()[0]["id"] == archive["id"]

    asset_archives = client.get(f"/api/assets/{asset['id']}/work-order-archives", headers=headers)
    assert asset_archives.status_code == 200
    assert asset_archives.json()[0]["archive_number"] == archive["archive_number"]

    global_archives = client.get("/api/work-order-archives/paged?q=PAPIR-2020-17", headers=headers)
    assert global_archives.status_code == 200
    assert global_archives.json()["total"] == 1

    preview = client.get(f"/api/work-order-archives/{archive['id']}/preview", headers=headers)
    assert preview.status_code == 200
    assert preview.content == pdf_bytes
    assert preview.headers["content-disposition"].startswith("inline")

    download = client.get(f"/api/work-order-archives/{archive['id']}/download", headers=headers)
    assert download.status_code == 200
    assert download.content == pdf_bytes
    assert download.headers["content-type"].startswith("application/pdf")


def test_global_archive_upload_requires_asset_and_can_link_matching_work_order():
    headers = auth_headers()
    first_asset, customer_id, location_id = _create_archive_test_asset(headers, "GLOBAL-ONE")
    second_asset, _, _ = _create_archive_test_asset(headers, "GLOBAL-TWO")
    work_order_response = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "asset_ids": [first_asset["id"], second_asset["id"]],
            "description": "Globális archív feltöltés",
            "priority": "normál",
            "status": "új",
        },
    )
    assert work_order_response.status_code == 201, work_order_response.text
    work_order = work_order_response.json()

    legacy_upload = client.post(
        "/api/work-order-archives",
        headers=headers,
        data={
            "asset_id": str(first_asset["id"]),
            "source_number": "REGI-2021-44",
            "work_date": "2021-09-15",
        },
        files={"archive_file": ("regi.pdf", b"%PDF-1.4\n%%EOF\n", "application/pdf")},
    )
    assert legacy_upload.status_code == 201, legacy_upload.text
    legacy_archive = legacy_upload.json()
    assert legacy_archive["work_order_id"] is None
    assert legacy_archive["asset_ids"] == [first_asset["id"]]
    assert legacy_archive["source_number"] == "REGI-2021-44"

    upload = client.post(
        "/api/work-order-archives",
        headers=headers,
        data={
            "asset_id": str(first_asset["id"]),
            "work_order_id": str(work_order["id"]),
            "note": "Archívum aloldalról feltöltve",
        },
        files={"archive_file": ("globalis.pdf", b"%PDF-1.4\n%%EOF\n", "application/pdf")},
    )
    assert upload.status_code == 201, upload.text
    archive = upload.json()
    assert archive["work_order_id"] == work_order["id"]
    assert archive["source_number"] == work_order["number"]
    assert sorted(archive["asset_ids"]) == sorted([first_asset["id"], second_asset["id"]])

    unrelated_asset, _, _ = _create_archive_test_asset(headers, "GLOBAL-OTHER")
    mismatch = client.post(
        "/api/work-order-archives",
        headers=headers,
        data={"asset_id": str(unrelated_asset["id"]), "work_order_id": str(work_order["id"])},
        files={"archive_file": ("rossz.pdf", b"%PDF-1.4\n%%EOF\n", "application/pdf")},
    )
    assert mismatch.status_code == 400
    assert "nem kapcsolódik" in mismatch.json()["detail"]


def test_asset_archive_upload_supports_legacy_migration_without_work_order():
    headers = auth_headers()
    asset, _, _ = _create_archive_test_asset(headers, "LEGACY-ARCH")
    png_bytes = b"\x89PNG\r\n\x1a\n" + b"legacy scan"
    upload = client.post(
        f"/api/assets/{asset['id']}/work-order-archives",
        headers=headers,
        data={"source_number": "4587/2019", "work_date": "2019-04-12"},
        files={"archive_file": ("4587-2019.png", png_bytes, "image/png")},
    )
    assert upload.status_code == 201, upload.text
    archive = upload.json()
    assert archive["work_order_id"] is None
    assert archive["work_order_number"] is None
    assert archive["asset_ids"] == [asset["id"]]
    assert archive["source_number"] == "4587/2019"


def test_archive_delete_requires_admin_and_exact_confirmation():
    admin_headers = auth_headers()
    asset, _, _ = _create_archive_test_asset(admin_headers, "ARCH-DELETE")
    upload = client.post(
        f"/api/assets/{asset['id']}/work-order-archives",
        headers=admin_headers,
        files={"archive_file": ("torlendo.pdf", b"%PDF-1.4\n%%EOF\n", "application/pdf")},
    )
    assert upload.status_code == 201, upload.text
    archive = upload.json()

    db = TestingSessionLocal()
    archive_row = db.query(WorkOrderArchive).filter(WorkOrderArchive.id == archive["id"]).one()
    stored_path = archive_path(get_settings(), archive_row.storage_key)
    assert stored_path.is_file()
    office_role = db.query(Role).filter(Role.name == ROLE_OFFICE).one()
    db.add(User(
        email="archive-office@example.com",
        full_name="Archív irodai",
        password_hash=get_password_hash("Office!123"),
        role_id=office_role.id,
    ))
    db.commit()
    db.close()
    office_headers = auth_headers_for("archive-office@example.com", "Office!123")

    forbidden = client.request(
        "DELETE",
        f"/api/work-order-archives/{archive['id']}",
        headers=office_headers,
        json={"confirmation": "TÖRÖL"},
    )
    assert forbidden.status_code == 403

    wrong_confirmation = client.request(
        "DELETE",
        f"/api/work-order-archives/{archive['id']}",
        headers=admin_headers,
        json={"confirmation": "torol"},
    )
    assert wrong_confirmation.status_code == 400

    deleted = client.request(
        "DELETE",
        f"/api/work-order-archives/{archive['id']}",
        headers=admin_headers,
        json={"confirmation": "TÖRÖL"},
    )
    assert deleted.status_code == 204, deleted.text
    assert not stored_path.exists()
    assert client.get(f"/api/work-order-archives/{archive['id']}", headers=admin_headers).status_code == 404
    assert client.get(f"/api/work-order-archives/{archive['id']}/download", headers=admin_headers).status_code == 404
    asset_archives = client.get(f"/api/assets/{asset['id']}/work-order-archives", headers=admin_headers)
    assert asset_archives.status_code == 200
    assert asset_archives.json() == []


def test_archive_upload_rejects_unsupported_file_content():
    headers = auth_headers()
    asset, _, _ = _create_archive_test_asset(headers, "BAD-ARCH")
    response = client.post(
        f"/api/assets/{asset['id']}/work-order-archives",
        headers=headers,
        files={"archive_file": ("not-a-pdf.pdf", b"plain text", "application/pdf")},
    )
    assert response.status_code == 415


def test_full_backup_contains_and_restores_archive_files():
    headers = auth_headers()
    asset, _, _ = _create_archive_test_asset(headers, "BACKUP-ARCH")
    pdf_bytes = b"%PDF-1.4\n% backup archive\n%%EOF\n"
    upload = client.post(
        f"/api/assets/{asset['id']}/work-order-archives",
        headers=headers,
        data={"source_number": "BACKUP-OLD-1"},
        files={"archive_file": ("backup.pdf", pdf_bytes, "application/pdf")},
    )
    assert upload.status_code == 201, upload.text
    archive_id = upload.json()["id"]

    db = TestingSessionLocal()
    customer_id = db.query(Customer).first().id
    db.close()
    opportunity = client.post(
        "/api/crm/opportunities",
        headers=headers,
        json={"customer_id": customer_id, "title": "Backup CRM lehetőség", "stage": "ajánlat", "estimated_value": 990000},
    )
    assert opportunity.status_code == 201, opportunity.text
    opportunity_id = opportunity.json()["id"]
    activity = client.post(
        "/api/crm/activities",
        headers=headers,
        json={"customer_id": customer_id, "opportunity_id": opportunity_id, "activity_type": "jegyzet", "subject": "Backup CRM aktivitás"},
    )
    assert activity.status_code == 201, activity.text
    activity_id = activity.json()["id"]
    contract = client.post(
        "/api/crm/contracts", headers=headers,
        json={"customer_id": customer_id, "contract_number":"BACKUP-SZ-1", "contract_type":"Backup szerződés", "asset_ids":[asset["id"]]},
    )
    assert contract.status_code == 201, contract.text
    contract_id = contract.json()["id"]
    contract_pdf = b"%PDF-1.4\n% contract backup archive\n%%EOF\n"
    contract_archive = client.post(
        f"/api/crm/contracts/{contract_id}/archives", headers=headers,
        files={"archive_file": ("contract-backup.pdf", contract_pdf, "application/pdf")},
    )
    assert contract_archive.status_code == 201, contract_archive.text
    contract_archive_id = contract_archive.json()["id"]

    backup = client.get("/api/settings/backup", headers=headers)
    assert backup.status_code == 200, backup.text
    with zipfile.ZipFile(BytesIO(backup.content), "r") as archive_zip:
        manifest = json.loads(archive_zip.read("manifest.json"))
        assert manifest["version"] == 16
        assert any(name.startswith("archives/") for name in archive_zip.namelist())
        assert any(name.startswith("contract_archives/") for name in archive_zip.namelist())

    restore = client.post(
        "/api/settings/restore",
        headers=headers,
        data={"confirmation": "VISSZAÁLLÍTÁS"},
        files={"backup_file": ("backup.zip", backup.content, "application/zip")},
    )
    assert restore.status_code == 200, restore.text
    # Security hardening: a restore must not resurrect or preserve previously
    # issued sessions. The caller must authenticate again after restore.
    stale = client.get(f"/api/work-order-archives/{archive_id}/download", headers=headers)
    assert stale.status_code == 401, stale.text
    headers = auth_headers()
    download = client.get(f"/api/work-order-archives/{archive_id}/download", headers=headers)
    assert download.status_code == 200, download.text
    assert download.content == pdf_bytes
    contract_download = client.get(f"/api/crm/contract-archives/{contract_archive_id}/download", headers=headers)
    assert contract_download.status_code == 200, contract_download.text
    assert contract_download.content == contract_pdf
    db = TestingSessionLocal()
    restored_opportunity = db.get(CrmOpportunity, opportunity_id)
    restored_activity = db.get(CrmActivity, activity_id)
    restored_contract = db.get(CrmContract, contract_id)
    assert restored_opportunity is not None and restored_opportunity.title == "Backup CRM lehetőség"
    assert restored_activity is not None and restored_activity.subject == "Backup CRM aktivitás"
    assert restored_activity.opportunity_id == opportunity_id
    assert restored_contract is not None and restored_contract.contract_number == "BACKUP-SZ-1"
    assert restored_contract.asset_ids == [asset["id"]]
    db.close()


def test_asset_with_archive_cannot_be_deleted():
    headers = auth_headers()
    asset, _, _ = _create_archive_test_asset(headers, "DELETE-ARCH")
    upload = client.post(
        f"/api/assets/{asset['id']}/work-order-archives",
        headers=headers,
        files={"archive_file": ("archive.pdf", b"%PDF-1.4\n%%EOF\n", "application/pdf")},
    )
    assert upload.status_code == 201, upload.text
    delete_response = client.delete(f"/api/assets/{asset['id']}", headers=headers)
    assert delete_response.status_code == 400
    assert "archivált munkalap" in delete_response.json()["detail"]


def test_work_order_supports_two_technicians_and_secondary_filter_and_close():
    headers = auth_headers()
    db = TestingSessionLocal()
    technician_role = db.query(Role).filter(Role.name == ROLE_TECHNICIAN).one()
    first = User(email="tech-one@example.com", full_name="Első Technikus", password_hash=get_password_hash("TechOne!1"), role_id=technician_role.id)
    second = User(email="tech-two@example.com", full_name="Második Technikus", password_hash=get_password_hash("TechTwo!1"), role_id=technician_role.id)
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    db.add_all([first, second])
    grant_service_access(db, first)
    grant_service_access(db, second)
    db.commit()
    first_id, second_id = first.id, second.id
    customer_id, location_id = customer.id, location.id
    db.close()

    created = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "description": "Két technikus teszt",
            "status": "folyamatban",
            "priority": "normál",
            "technician_id": first_id,
            "secondary_technician_id": second_id,
        },
    )
    assert created.status_code == 201, created.text
    payload = created.json()
    assert payload["technician_ids"] == [first_id, second_id]
    assert payload["technician_names"] == ["Első Technikus", "Második Technikus"]

    filtered = client.get(f"/api/work-orders/paged?technician_id={second_id}", headers=headers)
    assert filtered.status_code == 200, filtered.text
    assert [row["id"] for row in filtered.json()["items"]] == [payload["id"]]

    duplicate = client.put(
        f"/api/work-orders/{payload['id']}",
        headers=headers,
        json={"version": payload["version"], "secondary_technician_id": first_id},
    )
    assert duplicate.status_code == 400

    second_headers = auth_headers_for("tech-two@example.com", "TechTwo!1")
    closed = client.post(
        f"/api/work-orders/{payload['id']}/close",
        headers=second_headers,
        json={
            "version": payload["version"],
            "work_done": "A munka elkészült",
            "final_status": "Üzemképes",
            "labor_hours": 1,
            "started_at": "2026-07-31T08:00:00",
            "completed_at": "2026-07-31T09:00:00",
        },
    )
    assert closed.status_code == 200, closed.text
    assert closed.json()["technician_ids"] == [first_id, second_id]

    calendar = client.get("/api/dashboard/calendar?week_start=2026-07-27", headers=headers)
    assert calendar.status_code == 200, calendar.text
    calendar_event = next(event for event in calendar.json()["events"] if event["id"] == payload["id"])
    assert calendar_event["technician_ids"] == [first_id, second_id]
    assert calendar_event["technician_name"] == "Első Technikus, Második Technikus"


def test_work_order_create_and_update_preserves_multiple_material_lines():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    first_material = db.query(Material).first()
    second_material = Material(sku="TESZT-2", name="Második teszt anyag", unit="db", unit_price=250)
    db.add(second_material)
    db.commit()
    customer_id, location_id = customer.id, location.id
    first_material_id, second_material_id = first_material.id, second_material.id
    db.close()

    created = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "description": "Több anyagsor teszt",
            "status": "új",
            "priority": "normál",
            "materials": [
                {"material_id": first_material_id, "quantity": 2, "unit_price": 100, "note": "Első sor"},
                {"material_id": second_material_id, "quantity": 3, "unit_price": 250, "note": "Második sor"},
            ],
        },
    )
    assert created.status_code == 201, created.text
    payload = created.json()
    assert [line["material_id"] for line in payload["materials"]] == [first_material_id, second_material_id]

    updated = client.put(
        f"/api/work-orders/{payload['id']}",
        headers=headers,
        json={
            "version": payload["version"],
            "materials": [
                {"material_id": first_material_id, "quantity": 1, "unit_price": 100, "note": "Módosított első"},
                {"material_id": second_material_id, "quantity": 4, "unit_price": 250, "note": "Módosított második"},
            ],
        },
    )
    assert updated.status_code == 200, updated.text
    assert len(updated.json()["materials"]) == 2
    assert updated.json()["materials"][1]["quantity"] == 4


def test_work_order_linked_asset_printer_tracking_loads_and_accepts_meter_reading():
    headers = auth_headers()
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).first()
    customer_id, location_id = customer.id, location.id
    db.close()

    asset_response = client.post(
        "/api/assets",
        headers=headers,
        json={
            "internal_id": "WO-METER-001",
            "serial_number": "WO-METER-SN-001",
            "type": "Teszt nyomtató",
            "status": "aktív",
            "customer_id": customer_id,
            "current_location_id": location_id,
        },
    )
    asset_id = asset_response.json()["id"]
    order_response = client.post(
        "/api/work-orders",
        headers=headers,
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "asset_ids": [asset_id],
            "description": "Munkalapos számláló teszt",
            "status": "folyamatban",
            "priority": "normál",
        },
    )
    order_id = order_response.json()["id"]

    tracking = client.get(f"/api/assets/{asset_id}/printer-tracking", headers=headers)
    assert tracking.status_code == 200, tracking.text
    assert tracking.json()["asset"]["id"] == asset_id
    assert tracking.json()["latest_meter"] is None

    recorded = client.post(
        f"/api/assets/{asset_id}/meter-readings",
        headers=headers,
        json={
            "value": 1200,
            "black_white_value": 1000,
            "color_value": 200,
            "scan_value": 350,
            "recorded_at": "2026-07-31T10:00:00",
            "work_order_id": order_id,
        },
    )
    assert recorded.status_code == 201, recorded.text
    assert recorded.json()["latest_meter"]["value"] == 1200
    assert recorded.json()["latest_meter"]["work_order_id"] == order_id


def test_module_permissions_are_database_backed_and_apply_without_new_token():
    admin_headers = auth_headers()
    db = TestingSessionLocal()
    office_role_id = db.query(Role).filter(Role.name == ROLE_OFFICE).one().id
    db.close()

    catalog = client.get('/api/users/permissions', headers=admin_headers)
    assert catalog.status_code == 200, catalog.text
    catalog_codes = {item['code'] for item in catalog.json()}
    assert {PERMISSION_SERVICE_ACCESS, PERMISSION_HR_ACCESS, PERMISSION_HR_SUMMARY_VIEW, PERMISSION_CRM_ACCESS, PERMISSION_PROCUREMENT_APPROVE} <= catalog_codes

    created = client.post(
        '/api/users',
        headers=admin_headers,
        json={
            'email': 'module-user@example.com',
            'full_name': 'Modul Teszt Felhasználó',
            'password': 'Ideiglenes!A1',
            'role_id': office_role_id,
            'is_active': True,
            'permission_codes': [PERMISSION_HR_ACCESS, PERMISSION_HR_SUMMARY_VIEW],
        },
    )
    assert created.status_code == 201, created.text
    user_id = created.json()['id']
    assert created.json()['modules'] == ['hr', 'procurement']
    assert set(created.json()['permissions']) == {PERMISSION_HR_ACCESS, PERMISSION_HR_SUMMARY_VIEW}

    user_headers = auth_headers_for('module-user@example.com', 'Ideiglenes!A1')
    changed = client.post(
        '/api/auth/change-password',
        headers=user_headers,
        json={'current_password': 'Ideiglenes!A1', 'new_password': 'Vegleges!B2'},
    )
    assert changed.status_code == 200, changed.text

    blocked = client.get('/api/customers', headers=user_headers)
    assert blocked.status_code == 403
    assert blocked.json()['detail']['code'] == 'permission_required'
    assert blocked.json()['detail']['permission'] == PERMISSION_SERVICE_ACCESS

    granted = client.put(
        f'/api/users/{user_id}/permissions',
        headers=admin_headers,
        json={'permission_codes': [PERMISSION_SERVICE_ACCESS, PERMISSION_HR_ACCESS, PERMISSION_HR_SUMMARY_VIEW, PERMISSION_CRM_ACCESS]},
    )
    assert granted.status_code == 200, granted.text
    assert granted.json()['modules'] == ['service', 'hr', 'crm', 'procurement']

    # Ugyanaz a korábban kiadott JWT azonnal az adatbázisban lévő új jogokat használja.
    allowed = client.get('/api/customers', headers=user_headers)
    assert allowed.status_code == 200, allowed.text

    revoked = client.put(
        f'/api/users/{user_id}/permissions',
        headers=admin_headers,
        json={'permission_codes': [PERMISSION_HR_ACCESS, PERMISSION_HR_SUMMARY_VIEW]},
    )
    assert revoked.status_code == 200, revoked.text
    blocked_again = client.get('/api/customers', headers=user_headers)
    assert blocked_again.status_code == 403
    assert blocked_again.json()['detail']['permission'] == PERMISSION_SERVICE_ACCESS


def test_permission_update_requires_parent_module_access():
    admin_headers = auth_headers()
    db = TestingSessionLocal()
    office_role_id = db.query(Role).filter(Role.name == ROLE_OFFICE).one().id
    db.close()
    created = client.post(
        '/api/users',
        headers=admin_headers,
        json={
            'email': 'invalid-permission@example.com',
            'full_name': 'Jogosultság Teszt',
            'password': 'Ideiglenes!A1',
            'role_id': office_role_id,
            'is_active': True,
        },
    )
    assert created.status_code == 201, created.text
    response = client.put(
        f"/api/users/{created.json()['id']}/permissions",
        headers=admin_headers,
        json={'permission_codes': [PERMISSION_HR_SUMMARY_VIEW]},
    )
    assert response.status_code == 400
    assert PERMISSION_HR_ACCESS in response.json()['detail']


def test_admin_me_exposes_all_module_entries():
    response = client.get('/api/auth/me', headers=auth_headers())
    assert response.status_code == 200
    assert response.json()['modules'] == ['service', 'hr', 'crm', 'procurement', 'admin']



def test_structured_audit_login_contains_request_actor_session_and_hash_chain():
    response = client.post(
        "/api/auth/login",
        data={"username": "admin@example.com", "password": "admin123"},
        headers={"X-Correlation-ID": "audit-test-correlation", "User-Agent": "AuditTest/1.0"},
    )
    assert response.status_code == 200, response.text
    assert response.headers.get("x-request-id")
    assert response.headers.get("x-correlation-id") == "audit-test-correlation"

    db = TestingSessionLocal()
    row = db.query(AuditLog).filter(AuditLog.action == "LOGIN_SUCCESS").order_by(AuditLog.id.desc()).first()
    assert row is not None
    assert row.actor_name == "Admin"
    assert row.actor_email == "admin@example.com"
    assert row.actor_role == ROLE_ADMIN
    assert row.result == "success"
    assert row.severity == "info"
    assert row.request_id
    assert row.correlation_id == "audit-test-correlation"
    assert row.session_id
    assert row.user_agent == "AuditTest/1.0"
    assert len(row.entry_hash) == 64
    assert verify_audit_chain(db)["valid"] is True
    db.close()


def test_failed_login_is_audited_without_password_or_token_data():
    response = client.post(
        "/api/auth/login",
        data={"username": "admin@example.com", "password": "NagyonTitkos!123"},
    )
    assert response.status_code == 401

    db = TestingSessionLocal()
    row = db.query(AuditLog).filter(AuditLog.action == "LOGIN_FAILED").order_by(AuditLog.id.desc()).first()
    assert row is not None
    assert row.result == "failure"
    assert row.severity == "warning"
    serialized = json.dumps(
        {"before": row.before_data, "after": row.after_data, "changes": row.changes, "description": row.description},
        ensure_ascii=False,
    )
    assert "NagyonTitkos!123" not in serialized
    assert "password_hash" not in serialized
    assert "access_token" not in serialized
    assert verify_audit_chain(db)["valid"] is True
    db.close()


def test_user_permission_change_audit_has_structured_before_after_without_password_hash():
    headers = auth_headers()
    db = TestingSessionLocal()
    office_role = db.query(Role).filter(Role.name == ROLE_OFFICE).one()
    db.close()

    create = client.post(
        "/api/users",
        headers=headers,
        json={
            "email": "audit-office@example.com",
            "full_name": "Audit Office",
            "password": "Ideiglenes!9",
            "role_id": office_role.id,
            "is_active": True,
            "permission_codes": [PERMISSION_SERVICE_ACCESS],
        },
    )
    assert create.status_code == 201, create.text
    user_id = create.json()["id"]
    update = client.put(
        f"/api/users/{user_id}/permissions",
        headers=headers,
        json={"permission_codes": [PERMISSION_SERVICE_ACCESS, PERMISSION_HR_ACCESS, PERMISSION_HR_SUMMARY_VIEW]},
    )
    assert update.status_code == 200, update.text

    db = TestingSessionLocal()
    row = (
        db.query(AuditLog)
        .filter(AuditLog.action == "jogosultság-módosítás", AuditLog.entity_id == str(user_id))
        .order_by(AuditLog.id.desc())
        .first()
    )
    assert row is not None
    assert row.before_data["permission_codes"] == [PERMISSION_SERVICE_ACCESS]
    assert set(row.after_data["permission_codes"]) == {PERMISSION_SERVICE_ACCESS, PERMISSION_HR_ACCESS, PERMISSION_HR_SUMMARY_VIEW}
    assert "permission_codes" in row.changes
    serialized = json.dumps({"before": row.before_data, "after": row.after_data, "changes": row.changes}, ensure_ascii=False)
    assert "password_hash" not in serialized
    assert "Ideiglenes!9" not in serialized
    assert verify_audit_chain(db)["valid"] is True
    db.close()


def test_audit_log_api_is_admin_only_and_integrity_detects_tampering():
    admin_headers = auth_headers()
    # Create an auditable business/security change explicitly. Authentication test
    # helpers create sessions directly and intentionally do not fabricate audit rows.
    db = TestingSessionLocal()
    office_role_id = db.query(Role).filter(Role.name == ROLE_OFFICE).one().id
    db.close()
    created = client.post(
        "/api/users",
        headers=admin_headers,
        json={
            "email": "audit-seed@example.com",
            "full_name": "Audit Seed",
            "password": "AuditSeed!9A",
            "role_id": office_role_id,
            "is_active": True,
            "permission_codes": [PERMISSION_SERVICE_ACCESS],
        },
    )
    assert created.status_code == 201, created.text
    response = client.get("/api/audit-logs", headers=admin_headers)
    assert response.status_code == 200, response.text
    assert response.json()["total"] >= 1
    assert response.json()["items"][0]["entry_hash"]

    db = TestingSessionLocal()
    office_role = db.query(Role).filter(Role.name == ROLE_OFFICE).one()
    office = User(
        email="audit-reader@example.com",
        full_name="Audit Reader",
        password_hash=get_password_hash("AuditReader!9"),
        role_id=office_role.id,
        is_active=True,
    )
    db.add(office)
    db.flush()
    grant_service_access(db, office)
    db.commit()
    first_audit_id = db.query(AuditLog.id).order_by(AuditLog.id.asc()).scalar()
    db.close()

    office_headers = auth_headers_for("audit-reader@example.com", "AuditReader!9")
    denied = client.get("/api/audit-logs", headers=office_headers)
    assert denied.status_code == 403

    integrity = client.get("/api/audit-logs/integrity", headers=admin_headers)
    assert integrity.status_code == 200
    assert integrity.json()["valid"] is True

    db = TestingSessionLocal()
    row = db.get(AuditLog, first_audit_id)
    row.description = row.description + " [TAMPER]"
    db.commit()
    assert verify_audit_chain(db)["valid"] is False
    db.close()


def _create_hr_user(admin_headers, email, name, permission_codes):
    db = TestingSessionLocal()
    role_id = db.query(Role).filter(Role.name == ROLE_OFFICE).one().id
    db.close()
    response = client.post(
        '/api/users', headers=admin_headers,
        json={
            'email': email, 'full_name': name, 'password': 'Ideiglenes!A1',
            'role_id': role_id, 'is_active': True, 'permission_codes': permission_codes,
        },
    )
    assert response.status_code == 201, response.text
    headers = auth_headers_for(email, 'Ideiglenes!A1')
    changed = client.post('/api/auth/change-password', headers=headers, json={'current_password':'Ideiglenes!A1','new_password':'Vegleges!B2'})
    assert changed.status_code == 200, changed.text
    return response.json(), headers


def test_hr_leave_self_isolation_approval_and_summary_permissions():
    admin_headers = auth_headers()
    approver, approver_headers = _create_hr_user(
        admin_headers, 'leader@example.com', 'HR Vezető',
        [PERMISSION_HR_ACCESS, PERMISSION_HR_LEAVE_APPROVE],
    )
    employee1, employee1_headers = _create_hr_user(
        admin_headers, 'employee1@example.com', 'Első Dolgozó',
        [PERMISSION_HR_ACCESS, PERMISSION_HR_LEAVE_SELF, PERMISSION_HR_LEAVE_REQUEST],
    )
    employee2, employee2_headers = _create_hr_user(
        admin_headers, 'employee2@example.com', 'Második Dolgozó',
        [PERMISSION_HR_ACCESS, PERMISSION_HR_LEAVE_SELF, PERMISSION_HR_LEAVE_REQUEST],
    )
    summary_user, summary_headers = _create_hr_user(
        admin_headers, 'summary@example.com', 'HR Összesítő',
        [PERMISSION_HR_ACCESS, PERMISSION_HR_SUMMARY_VIEW],
    )

    for employee in (employee1, employee2):
        profile = client.put(
            f"/api/hr/admin/employees/{employee['id']}/profile", headers=admin_headers,
            json={'leave_approver_user_id': approver['id'], 'active': True, 'employee_number': f"T-{employee['id']}"},
        )
        assert profile.status_code == 200, profile.text
        entitlement = client.put(
            f"/api/hr/admin/employees/{employee['id']}/entitlement", headers=admin_headers,
            json={'year': 2026, 'base_days': 25, 'carried_days': 0, 'adjustment_days': 0},
        )
        assert entitlement.status_code == 200, entitlement.text

    created = client.post(
        '/api/hr/me/leave-requests', headers=employee1_headers,
        json={'start_date':'2026-09-07','end_date':'2026-09-11','employee_comment':'Családi program'},
    )
    assert created.status_code == 201, created.text
    request_id = created.json()['id']
    assert created.json()['requested_days'] == '5.00'
    assert created.json()['employee_name'] == 'Első Dolgozó'

    # A strukturált audit őrizze a döntési nyomot, de ne szivárogtassa ki
    # a pontos szabadságdátumot vagy a munkavállaló szabad szöveges megjegyzését.
    db = TestingSessionLocal()
    leave_audit = (
        db.query(AuditLog)
        .filter(AuditLog.action == 'LEAVE_REQUEST_SUBMITTED', AuditLog.entity_id == str(request_id))
        .order_by(AuditLog.id.desc())
        .first()
    )
    assert leave_audit is not None
    assert '2026-09-07' not in leave_audit.description
    assert '2026-09-11' not in leave_audit.description
    assert leave_audit.after_data['requested_days'] == '5.00'
    assert 'start_date' not in leave_audit.after_data
    assert 'end_date' not in leave_audit.after_data
    assert 'employee_comment' not in leave_audit.after_data
    db.close()

    own1 = client.get('/api/hr/me?year=2026', headers=employee1_headers)
    own2 = client.get('/api/hr/me?year=2026', headers=employee2_headers)
    assert own1.status_code == 200 and len(own1.json()['leave_requests']) == 1
    assert own2.status_code == 200 and own2.json()['leave_requests'] == []
    assert own1.json()['balance']['pending_days'] == '5.00'
    assert own1.json()['balance']['available_to_request_days'] == '20.00'

    # Egy normál munkavállaló nem kaphat HR admin vagy összesítő adatot.
    assert client.get('/api/hr/admin/employees?year=2026', headers=employee1_headers).status_code == 403
    assert client.get('/api/hr/summary?year=2026', headers=employee1_headers).status_code == 403

    approvals = client.get('/api/hr/approvals?status=pending', headers=approver_headers)
    assert approvals.status_code == 200, approvals.text
    assert [row['id'] for row in approvals.json()] == [request_id]
    approved = client.post(
        f'/api/hr/approvals/{request_id}/approve', headers=approver_headers,
        json={'version': created.json()['version'], 'comment':'Jóváhagyva'},
    )
    assert approved.status_code == 200, approved.text
    assert approved.json()['status'] == 'approved'

    own1_after = client.get('/api/hr/me?year=2026', headers=employee1_headers).json()
    assert own1_after['balance']['approved_days'] == '5.00'
    assert own1_after['balance']['pending_days'] == '0.00'
    assert own1_after['balance']['remaining_days'] == '20.00'

    summary = client.get('/api/hr/summary?year=2026', headers=summary_headers)
    assert summary.status_code == 200, summary.text
    by_name = {row['full_name']: row for row in summary.json()}
    assert by_name['Első Dolgozó']['approved_days'] == '5.00'
    assert by_name['Második Dolgozó']['approved_days'] == '0.00'


def test_hr_approver_can_only_decide_assigned_requests_and_calendar_overrides_workdays():
    admin_headers = auth_headers()
    leader1, leader1_headers = _create_hr_user(admin_headers, 'leader1@example.com', 'Vezető Egy', [PERMISSION_HR_ACCESS, PERMISSION_HR_LEAVE_APPROVE])
    leader2, leader2_headers = _create_hr_user(admin_headers, 'leader2@example.com', 'Vezető Kettő', [PERMISSION_HR_ACCESS, PERMISSION_HR_LEAVE_APPROVE])
    employee, employee_headers = _create_hr_user(admin_headers, 'calendar.employee@example.com', 'Naptár Dolgozó', [PERMISSION_HR_ACCESS, PERMISSION_HR_LEAVE_SELF, PERMISSION_HR_LEAVE_REQUEST])
    profile = client.put(f"/api/hr/admin/employees/{employee['id']}/profile", headers=admin_headers, json={'leave_approver_user_id': leader1['id'], 'active': True})
    assert profile.status_code == 200, profile.text
    assert client.put(f"/api/hr/admin/employees/{employee['id']}/entitlement", headers=admin_headers, json={'year':2026,'base_days':25}).status_code == 200

    # Hétfő kivételesen nem munkanap, a következő szombat munkanap.
    assert client.put('/api/hr/admin/calendar', headers=admin_headers, json={'calendar_date':'2026-09-07','is_working_day':False,'label':'teszt pihenőnap'}).status_code == 200
    assert client.put('/api/hr/admin/calendar', headers=admin_headers, json={'calendar_date':'2026-09-12','is_working_day':True,'label':'teszt munkanap'}).status_code == 200
    request = client.post('/api/hr/me/leave-requests', headers=employee_headers, json={'start_date':'2026-09-07','end_date':'2026-09-12'})
    assert request.status_code == 201, request.text
    assert request.json()['requested_days'] == '5.00'  # K-P 4 nap + szombati kivétel 1 nap

    denied = client.post(f"/api/hr/approvals/{request.json()['id']}/approve", headers=leader2_headers, json={'version':1})
    assert denied.status_code == 404
    assert client.get('/api/hr/approvals?status=pending', headers=leader2_headers).json() == []
    assert [r['id'] for r in client.get('/api/hr/approvals?status=pending', headers=leader1_headers).json()] == [request.json()['id']]


def test_approved_leave_cancellation_requires_assigned_approver_and_keeps_balance_until_decision():
    admin_headers = auth_headers()
    leader1, leader1_headers = _create_hr_user(
        admin_headers,
        'cancel-leader1@example.com',
        'Visszavonás Jóváhagyó',
        [PERMISSION_HR_ACCESS, PERMISSION_HR_LEAVE_APPROVE],
    )
    _, leader2_headers = _create_hr_user(
        admin_headers,
        'cancel-leader2@example.com',
        'Másik Jóváhagyó',
        [PERMISSION_HR_ACCESS, PERMISSION_HR_LEAVE_APPROVE],
    )
    employee, employee_headers = _create_hr_user(
        admin_headers,
        'cancel-employee@example.com',
        'Visszavonó Dolgozó',
        [PERMISSION_HR_ACCESS, PERMISSION_HR_LEAVE_SELF, PERMISSION_HR_LEAVE_REQUEST],
    )
    assert client.put(
        f"/api/hr/admin/employees/{employee['id']}/profile",
        headers=admin_headers,
        json={'leave_approver_user_id': leader1['id'], 'active': True},
    ).status_code == 200
    assert client.put(
        f"/api/hr/admin/employees/{employee['id']}/entitlement",
        headers=admin_headers,
        json={'year': 2026, 'base_days': 20},
    ).status_code == 200

    created = client.post(
        '/api/hr/me/leave-requests',
        headers=employee_headers,
        json={'start_date': '2026-11-02', 'end_date': '2026-11-06'},
    )
    assert created.status_code == 201, created.text
    request_id = created.json()['id']
    approved = client.post(
        f'/api/hr/approvals/{request_id}/approve',
        headers=leader1_headers,
        json={'version': created.json()['version']},
    )
    assert approved.status_code == 200, approved.text
    assert approved.json()['status'] == 'approved'

    # A saját jóváhagyott szabadság visszavonásához a hr.leave.self jog elég;
    # új szabadságigény-létrehozási jog nélkül is elindítható a folyamat.
    permissions = client.put(
        f"/api/users/{employee['id']}/permissions",
        headers=admin_headers,
        json={'permission_codes': [PERMISSION_HR_ACCESS, PERMISSION_HR_LEAVE_SELF]},
    )
    assert permissions.status_code == 200, permissions.text
    cancellation = client.post(
        f'/api/hr/me/leave-requests/{request_id}/request-cancellation',
        headers=employee_headers,
        json={'version': approved.json()['version'], 'comment': 'Megváltozott a terv'},
    )
    assert cancellation.status_code == 200, cancellation.text
    cancellation_row = cancellation.json()
    assert cancellation_row['status'] == 'approved'
    assert cancellation_row['cancellation_status'] == 'pending'
    assert cancellation_row['cancellation_approver_user_id'] == leader1['id']
    assert cancellation_row['approval_kind'] == 'cancellation'

    # A kijelölt jóváhagyó addig nem veszítheti el a döntési jogát, amíg
    # hozzá rendelt visszavonási kérelem várakozik.
    protected_approver = client.put(
        f"/api/users/{leader1['id']}/permissions",
        headers=admin_headers,
        json={'permission_codes': [PERMISSION_HR_ACCESS]},
    )
    assert protected_approver.status_code == 409, protected_approver.text

    while_pending = client.get('/api/hr/me?year=2026', headers=employee_headers).json()
    assert while_pending['balance']['approved_days'] == '5.00'
    assert while_pending['balance']['remaining_days'] == '15.00'

    approvals = client.get('/api/hr/approvals?status=pending', headers=leader1_headers)
    assert approvals.status_code == 200, approvals.text
    pending = next(row for row in approvals.json() if row['id'] == request_id)
    assert pending['approval_kind'] == 'cancellation'
    assert pending['cancellation_comment'] == 'Megváltozott a terv'
    denied = client.post(
        f'/api/hr/approvals/{request_id}/cancellation/approve',
        headers=leader2_headers,
        json={'version': cancellation_row['version']},
    )
    assert denied.status_code == 404

    rejected = client.post(
        f'/api/hr/approvals/{request_id}/cancellation/reject',
        headers=leader1_headers,
        json={'version': cancellation_row['version'], 'comment': 'A beosztás miatt maradjon'},
    )
    assert rejected.status_code == 200, rejected.text
    assert rejected.json()['status'] == 'approved'
    assert rejected.json()['cancellation_status'] == 'rejected'

    requested_again = client.post(
        f'/api/hr/me/leave-requests/{request_id}/request-cancellation',
        headers=employee_headers,
        json={'version': rejected.json()['version'], 'comment': 'Új körülmény'},
    )
    assert requested_again.status_code == 200, requested_again.text
    cancelled = client.post(
        f'/api/hr/approvals/{request_id}/cancellation/approve',
        headers=leader1_headers,
        json={'version': requested_again.json()['version'], 'comment': 'Engedélyezve'},
    )
    assert cancelled.status_code == 200, cancelled.text
    assert cancelled.json()['status'] == 'cancelled'
    assert cancelled.json()['cancellation_status'] == 'approved'
    assert cancelled.json()['cancelled_at'] is not None

    after = client.get('/api/hr/me?year=2026', headers=employee_headers).json()
    assert after['balance']['approved_days'] == '0.00'
    assert after['balance']['remaining_days'] == '20.00'
    db = TestingSessionLocal()
    actions = {row.action for row in db.query(AuditLog).filter(AuditLog.entity_id == str(request_id)).all()}
    db.close()
    assert {'LEAVE_CANCELLATION_REQUESTED', 'LEAVE_CANCELLATION_REJECTED', 'LEAVE_CANCELLATION_APPROVED'} <= actions


def test_hr_request_requires_profile_approver_entitlement_and_prevents_overlap_or_cross_year():
    admin_headers = auth_headers()
    leader, _ = _create_hr_user(admin_headers, 'validation.leader@example.com', 'Validáló Vezető', [PERMISSION_HR_ACCESS, PERMISSION_HR_LEAVE_APPROVE])
    employee, employee_headers = _create_hr_user(admin_headers, 'validation.employee@example.com', 'Validáló Dolgozó', [PERMISSION_HR_ACCESS, PERMISSION_HR_LEAVE_SELF, PERMISSION_HR_LEAVE_REQUEST])

    missing_profile = client.get('/api/hr/me?year=2026', headers=employee_headers)
    assert missing_profile.status_code == 409
    assert client.put(f"/api/hr/admin/employees/{employee['id']}/profile", headers=admin_headers, json={'leave_approver_user_id':leader['id'],'active':True}).status_code == 200
    no_entitlement = client.post('/api/hr/me/leave-requests', headers=employee_headers, json={'start_date':'2026-10-05','end_date':'2026-10-05'})
    assert no_entitlement.status_code == 409
    assert client.put(f"/api/hr/admin/employees/{employee['id']}/entitlement", headers=admin_headers, json={'year':2026,'base_days':2}).status_code == 200
    cross_year = client.post('/api/hr/me/leave-requests', headers=employee_headers, json={'start_date':'2026-12-31','end_date':'2027-01-02'})
    assert cross_year.status_code == 400
    first = client.post('/api/hr/me/leave-requests', headers=employee_headers, json={'start_date':'2026-10-05','end_date':'2026-10-05'})
    assert first.status_code == 201, first.text
    overlap = client.post('/api/hr/me/leave-requests', headers=employee_headers, json={'start_date':'2026-10-05','end_date':'2026-10-06'})
    assert overlap.status_code == 409


def _create_crm_user(email: str, permission_codes: list[str]):
    admin_headers = auth_headers()
    db = TestingSessionLocal()
    role_id = db.query(Role).filter(Role.name == ROLE_OFFICE).one().id
    db.close()
    created = client.post(
        "/api/users",
        headers=admin_headers,
        json={
            "email": email,
            "full_name": "CRM Teszt Felhasználó",
            "password": "Ideiglenes!A1",
            "role_id": role_id,
            "is_active": True,
            "permission_codes": permission_codes,
        },
    )
    assert created.status_code == 201, created.text
    headers = auth_headers_for(email, "Ideiglenes!A1")
    changed = client.post(
        "/api/auth/change-password",
        headers=headers,
        json={"current_password": "Ideiglenes!A1", "new_password": "Vegleges!B2"},
    )
    assert changed.status_code == 200, changed.text
    return created.json(), headers


def test_crm_access_is_readable_but_write_requires_crm_write_permission():
    user, headers = _create_crm_user("crm-read@example.com", [PERMISSION_CRM_ACCESS])
    listing = client.get("/api/crm/customers", headers=headers)
    assert listing.status_code == 200, listing.text
    assert listing.json()[0]["name"] == "Teszt ügyfél"

    db = TestingSessionLocal()
    customer_id = db.query(Customer).first().id
    db.close()
    blocked = client.post("/api/crm/opportunities", headers=headers, json={"customer_id": customer_id, "title": "Tiltott lehetőség"})
    assert blocked.status_code == 403
    assert blocked.json()["detail"]["permission"] == PERMISSION_CRM_WRITE

    granted = client.put(
        f"/api/users/{user['id']}/permissions",
        headers=auth_headers(),
        json={"permission_codes": [PERMISSION_CRM_ACCESS, PERMISSION_CRM_WRITE]},
    )
    assert granted.status_code == 200, granted.text
    allowed = client.post("/api/crm/opportunities", headers=headers, json={"customer_id": customer_id, "title": "Engedélyezett lehetőség", "estimated_value": 1250000})
    assert allowed.status_code == 201, allowed.text
    assert allowed.json()["owner_user_id"] == user["id"]


def test_crm_opportunity_activity_and_customer_360_view_use_shared_customer_data():
    user, headers = _create_crm_user("crm-write@example.com", [PERMISSION_CRM_ACCESS, PERMISSION_CRM_WRITE])
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    customer_id = customer.id
    location_id = db.query(Location).filter(Location.customer_id == customer_id).first().id
    db.close()

    contact = client.post(
        f"/api/crm/customers/{customer_id}/contacts",
        headers=headers,
        json={"name": "CRM Kapcsolattartó", "category": "Sales", "email": "sales@example.com", "location_id": location_id},
    )
    assert contact.status_code == 201, contact.text

    opportunity = client.post(
        "/api/crm/opportunities",
        headers=headers,
        json={
            "customer_id": customer_id,
            "contact_id": contact.json()["id"],
            "title": "Új szkenner projekt",
            "stage": "ajánlat",
            "estimated_value": 2400000,
            "currency": "huf",
            "probability": 60,
            "expected_close_date": "2026-09-30",
            "next_step": "Ajánlat visszaigazolása",
        },
    )
    assert opportunity.status_code == 201, opportunity.text
    opp = opportunity.json()
    assert opp["customer_name"] == "Teszt ügyfél"
    assert opp["contact_name"] == "CRM Kapcsolattartó"
    assert opp["currency"] == "HUF"

    activity = client.post(
        "/api/crm/activities",
        headers=headers,
        json={
            "customer_id": customer_id,
            "contact_id": contact.json()["id"],
            "opportunity_id": opp["id"],
            "activity_type": "telefon",
            "subject": "Ajánlat egyeztetés",
            "description": "Az ügyfél műszaki pontosítást kért.",
            "due_at": "2026-08-20T10:00:00",
        },
    )
    assert activity.status_code == 201, activity.text
    assert activity.json()["opportunity_title"] == "Új szkenner projekt"
    assert activity.json()["owner_user_id"] == user["id"]

    detail = client.get(f"/api/crm/customers/{customer_id}", headers=headers)
    assert detail.status_code == 200, detail.text
    payload = detail.json()
    assert payload["customer"]["id"] == customer_id
    assert payload["summary"]["open_opportunity_count"] == 1
    assert float(payload["summary"]["open_pipeline_value"]) == 2400000
    assert payload["opportunities"][0]["id"] == opp["id"]
    assert payload["activities"][0]["subject"] == "Ajánlat egyeztetés"

    dashboard = client.get("/api/crm/dashboard", headers=headers)
    assert dashboard.status_code == 200, dashboard.text
    assert dashboard.json()["open_opportunity_count"] == 1
    assert float(dashboard.json()["open_pipeline_value"]) == 2400000

    db = TestingSessionLocal()
    audit_actions = {row.action for row in db.query(AuditLog).filter(AuditLog.actor_user_id == user["id"]).all()}
    assert "CRM_OPPORTUNITY_CREATED" in audit_actions
    assert "CRM_ACTIVITY_CREATED" in audit_actions
    assert verify_audit_chain(db)["valid"] is True
    db.close()


def test_crm_activity_can_create_and_link_opportunity_with_one_action():
    user, headers = _create_crm_user(
        "crm-activity-conversion@example.com",
        [PERMISSION_CRM_ACCESS, PERMISSION_CRM_WRITE],
    )
    _, read_headers = _create_crm_user(
        "crm-activity-conversion-read@example.com",
        [PERMISSION_CRM_ACCESS],
    )
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    customer_id = customer.id
    location_id = db.query(Location).filter(Location.customer_id == customer_id).first().id
    db.close()

    contact = client.post(
        f"/api/crm/customers/{customer_id}/contacts",
        headers=headers,
        json={"name": "Érdeklődő Kapcsolat", "category": "Sales", "location_id": location_id},
    )
    assert contact.status_code == 201, contact.text
    activity = client.post(
        "/api/crm/activities",
        headers=headers,
        json={
            "customer_id": customer_id,
            "contact_id": contact.json()["id"],
            "activity_type": "telefon",
            "subject": "Új nyomtató beszerzése",
            "description": "Igényfelmérés és ajánlat készítése",
        },
    )
    assert activity.status_code == 201, activity.text
    activity_row = activity.json()
    assert activity_row["opportunity_id"] is None

    denied = client.post(
        f'/api/crm/activities/{activity_row["id"]}/opportunity',
        headers=read_headers,
        json={"version": activity_row["version"]},
    )
    assert denied.status_code == 403, denied.text

    converted = client.post(
        f'/api/crm/activities/{activity_row["id"]}/opportunity',
        headers=headers,
        json={"version": activity_row["version"]},
    )
    assert converted.status_code == 201, converted.text
    opportunity = converted.json()
    assert opportunity["customer_id"] == customer_id
    assert opportunity["contact_id"] == contact.json()["id"]
    assert opportunity["owner_user_id"] == user["id"]
    assert opportunity["title"] == "Új nyomtató beszerzése"
    assert opportunity["stage"] == "új"
    assert opportunity["currency"] == "HUF"
    assert opportunity["source"] == f'CRM aktivitás #{activity_row["id"]}'
    assert opportunity["next_step"] == "Igényfelmérés és ajánlat készítése"

    listed = client.get("/api/crm/activities", headers=headers)
    assert listed.status_code == 200, listed.text
    linked = next(row for row in listed.json() if row["id"] == activity_row["id"])
    assert linked["opportunity_id"] == opportunity["id"]
    assert linked["opportunity_title"] == opportunity["title"]
    assert linked["version"] == activity_row["version"] + 1

    duplicate = client.post(
        f'/api/crm/activities/{activity_row["id"]}/opportunity',
        headers=headers,
        json={"version": linked["version"]},
    )
    assert duplicate.status_code == 409, duplicate.text

    db = TestingSessionLocal()
    actions = {
        row.action
        for row in db.query(AuditLog).filter(
            AuditLog.actor_user_id == user["id"],
            AuditLog.action.in_([
                "CRM_OPPORTUNITY_CREATED_FROM_ACTIVITY",
                "CRM_ACTIVITY_LINKED_TO_OPPORTUNITY",
            ]),
        ).all()
    }
    db.close()
    assert actions == {
        "CRM_OPPORTUNITY_CREATED_FROM_ACTIVITY",
        "CRM_ACTIVITY_LINKED_TO_OPPORTUNITY",
    }


def test_crm_validates_cross_customer_links_and_optimistic_versions():
    _, headers = _create_crm_user("crm-guard@example.com", [PERMISSION_CRM_ACCESS, PERMISSION_CRM_WRITE])
    db = TestingSessionLocal()
    first = db.query(Customer).first()
    second = Customer(name="Másik CRM ügyfél")
    db.add(second)
    db.flush()
    foreign_contact = CustomerContact(customer_id=second.id, name="Másik kapcsolat", category="Sales")
    db.add(foreign_contact)
    db.commit()
    first_id, contact_id = first.id, foreign_contact.id
    db.close()

    invalid = client.post("/api/crm/opportunities", headers=headers, json={"customer_id": first_id, "contact_id": contact_id, "title": "Rossz kapcsolat"})
    assert invalid.status_code == 400

    created = client.post("/api/crm/opportunities", headers=headers, json={"customer_id": first_id, "title": "Verzióteszt"})
    assert created.status_code == 201, created.text
    opp = created.json()
    updated = client.put(f"/api/crm/opportunities/{opp['id']}", headers=headers, json={"version": opp["version"], "stage": "igényfelmérés"})
    assert updated.status_code == 200, updated.text
    assert updated.json()["version"] == opp["version"] + 1

    conflict = client.put(f"/api/crm/opportunities/{opp['id']}", headers=headers, json={"version": opp["version"], "stage": "ajánlat"})
    assert conflict.status_code == 409

    blocked_delete = client.delete(f"/api/customers/{first_id}", headers=auth_headers())
    assert blocked_delete.status_code == 400
    assert "CRM" in blocked_delete.json()["detail"]



def test_crm_contract_management_permission_scope_and_work_order_integration():
    read_user, read_headers = _create_crm_user("crm-contract-read@example.com", [PERMISSION_CRM_ACCESS])
    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    customer_id = customer.id
    location = db.query(Location).filter(Location.customer_id == customer_id).first()
    asset = Asset(internal_id="CRM-CONTRACT-ASSET-1", customer_id=customer_id, current_location_id=location.id, status="aktív", type="Nyomtató", serial_number="SER-C1")
    db.add(asset)
    db.commit()
    location_id, asset_id = location.id, asset.id
    db.close()

    blocked = client.post('/api/crm/contracts', headers=read_headers, json={
        'customer_id': customer_id, 'contract_number':'SZ-2026-001', 'contract_type':'Átalánydíjas szerviz'
    })
    assert blocked.status_code == 403, blocked.text
    assert blocked.json()['detail']['permission'] == PERMISSION_CRM_CONTRACT_MANAGE

    granted = client.put(
        f"/api/users/{read_user['id']}/permissions", headers=auth_headers(),
        json={'permission_codes':[PERMISSION_CRM_ACCESS, PERMISSION_CRM_CONTRACT_MANAGE]},
    )
    assert granted.status_code == 200, granted.text
    created = client.post('/api/crm/contracts', headers=read_headers, json={
        'customer_id': customer_id,
        'contract_number':'SZ-2026-001',
        'contract_type':'Átalánydíjas szerviz',
        'start_date':'2026-01-01', 'end_date':'2026-12-31', 'status':'aktív',
        'billing_model':'havi átalány', 'sla':'4 órás reakcióidő',
        'location_ids':[location_id], 'asset_ids':[asset_id],
    })
    assert created.status_code == 201, created.text
    contract = created.json()
    assert contract['location_ids'] == [location_id]
    assert contract['asset_ids'] == [asset_id]

    options = client.get(f'/api/work-orders/contract-options?customer_id={customer_id}', headers=auth_headers())
    assert options.status_code == 200, options.text
    assert options.json()[0]['contract_number'] == 'SZ-2026-001'

    work_order = client.post('/api/work-orders', headers=auth_headers(), json={
        'customer_id': customer_id, 'location_id': location_id, 'asset_ids':[asset_id],
        'contract_id': contract['id'], 'description':'Szerződéses javítás', 'status':'új',
        'contract_type':'Ezt a CRM felülírja',
    })
    assert work_order.status_code == 201, work_order.text
    order = work_order.json()
    assert order['contract_id'] == contract['id']
    assert order['contract_number'] == 'SZ-2026-001'
    assert order['contract_type'] == 'Átalánydíjas szerviz'

    update = client.put(f"/api/crm/contracts/{contract['id']}", headers=read_headers, json={
        'version': contract['version'], 'contract_type':'Prémium szerviz'
    })
    assert update.status_code == 200, update.text
    refreshed = client.get(f"/api/work-orders/{order['id']}", headers=auth_headers())
    assert refreshed.status_code == 200, refreshed.text
    assert refreshed.json()['contract_type'] == 'Átalánydíjas szerviz'  # történeti snapshot nem íródik át

    db = TestingSessionLocal()
    actions = {row.action for row in db.query(AuditLog).all()}
    assert 'CRM_CONTRACT_CREATED' in actions
    assert 'CRM_CONTRACT_UPDATED' in actions
    assert verify_audit_chain(db)['valid'] is True
    db.close()


def test_crm_contract_rejects_cross_customer_scope_and_work_order_outside_contract_scope():
    _, headers = _create_crm_user("crm-contract-guard@example.com", [PERMISSION_CRM_ACCESS, PERMISSION_CRM_CONTRACT_MANAGE])
    db = TestingSessionLocal()
    first = db.query(Customer).first()
    first_location = db.query(Location).filter(Location.customer_id == first.id).first()
    second = Customer(name='Másik szerződéses ügyfél')
    db.add(second); db.flush()
    second_location = Location(customer_id=second.id, name='Másik helyszín')
    db.add(second_location); db.flush()
    first_asset = Asset(internal_id='CONTRACT-FIRST', customer_id=first.id, current_location_id=first_location.id, status='aktív')
    second_asset = Asset(internal_id='CONTRACT-SECOND', customer_id=second.id, current_location_id=second_location.id, status='aktív')
    db.add_all([first_asset, second_asset]); db.commit()
    ids = first.id, first_location.id, second_location.id, first_asset.id, second_asset.id
    db.close()
    first_id, first_location_id, second_location_id, first_asset_id, second_asset_id = ids

    bad_scope = client.post('/api/crm/contracts', headers=headers, json={
        'customer_id': first_id, 'contract_number':'SZ-BAD-SCOPE', 'contract_type':'Teszt', 'location_ids':[second_location_id]
    })
    assert bad_scope.status_code == 400
    bad_asset = client.post('/api/crm/contracts', headers=headers, json={
        'customer_id': first_id, 'contract_number':'SZ-BAD-ASSET', 'contract_type':'Teszt', 'asset_ids':[second_asset_id]
    })
    assert bad_asset.status_code == 400

    contract = client.post('/api/crm/contracts', headers=headers, json={
        'customer_id': first_id, 'contract_number':'SZ-SCOPED', 'contract_type':'Eszközre szóló',
        'location_ids':[first_location_id], 'asset_ids':[first_asset_id]
    })
    assert contract.status_code == 201, contract.text
    invalid = client.post('/api/work-orders', headers=auth_headers(), json={
        'customer_id': first_id, 'location_id': first_location_id, 'asset_ids':[],
        'contract_id': contract.json()['id'], 'description':'Eszköz nélkül még megengedett ügyfélszintű munkalap'
    })
    assert invalid.status_code == 201, invalid.text

    db = TestingSessionLocal()
    other_same_customer = Asset(internal_id='CONTRACT-OTHER', customer_id=first_id, current_location_id=first_location_id, status='aktív')
    db.add(other_same_customer); db.commit(); other_id = other_same_customer.id; db.close()
    outside_asset = client.post('/api/work-orders', headers=auth_headers(), json={
        'customer_id': first_id, 'location_id': first_location_id, 'asset_ids':[other_id],
        'contract_id': contract.json()['id'], 'description':'Nem szerződéses eszköz'
    })
    assert outside_asset.status_code == 400


def test_phase6_notifications_reports_and_csv_exports():
    headers = auth_headers()
    db = TestingSessionLocal()
    admin = db.query(User).filter(User.email == "admin@example.com").one()
    customer = db.query(Customer).filter(Customer.name == "Teszt ügyfél").one()
    today = date.today()
    db.add(WorkOrder(
        number="ML-PHASE6-1", customer_id=customer.id, status="ütemezve", priority="sürgős",
        technician_id=admin.id, description="Mai teszt munkalap", planned_date=today,
        labor_hours=2, work_type="javítás",
    ))
    db.add(CrmOpportunity(
        customer_id=customer.id, owner_user_id=admin.id, title="Riport lehetőség",
        stage="ajánlat", estimated_value=100000, probability=50,
    ))
    db.add(CrmActivity(
        customer_id=customer.id, owner_user_id=admin.id, activity_type="feladat",
        subject="Lejárt CRM feladat", due_at=datetime.utcnow(),
    ))
    db.add(CrmContract(
        customer_id=customer.id, contract_number="PHASE6-CONTRACT", contract_type="Szerviz",
        status="aktív", start_date=today, end_date=today,
    ))
    profile = EmployeeProfile(user_id=admin.id, employee_number="A-001", organizational_unit="Vezetés", active=True)
    db.add(profile)
    db.flush()
    db.add(LeaveEntitlement(employee_id=profile.id, year=today.year, base_days=25, carried_days=0, adjustment_days=0))
    db.commit()
    db.close()

    notifications = client.get("/api/notifications", headers=headers)
    assert notifications.status_code == 200, notifications.text
    kinds = {row["notification_type"] for row in notifications.json()}
    assert "crm_activity_overdue" in kinds
    assert "contract_expiring" in kinds
    assert "work_order_today" in kinds

    count = client.get("/api/notifications/unread-count", headers=headers)
    assert count.status_code == 200
    assert count.json()["unread_count"] >= 3
    first_id = notifications.json()[0]["id"]
    assert client.post(f"/api/notifications/{first_id}/read", headers=headers).status_code == 200
    assert client.post("/api/notifications/read-all", headers=headers).status_code == 200
    assert client.get("/api/notifications/unread-count", headers=headers).json()["unread_count"] == 0

    service = client.get(f"/api/reports/service?date_from={today}&date_to={today}", headers=headers)
    assert service.status_code == 200, service.text
    assert service.json()["work_order_count"] == 1
    assert service.json()["urgent_count"] == 1

    hr = client.get(f"/api/reports/hr?year={today.year}", headers=headers)
    assert hr.status_code == 200, hr.text
    assert hr.json()["employee_count"] == 1
    assert float(hr.json()["total_days"]) == 25

    crm = client.get(f"/api/reports/crm?date_from={today}&date_to={today}", headers=headers)
    assert crm.status_code == 200, crm.text
    assert crm.json()["opportunity_count"] == 1
    assert float(crm.json()["weighted_pipeline_value"]) == 50000
    assert crm.json()["activities_by_type"]["feladat"] == 1
    assert crm.json()["expiring_contracts"][0]["contract_number"] == "PHASE6-CONTRACT"

    for path in (
        f"/api/reports/service.csv?date_from={today}&date_to={today}",
        f"/api/reports/hr.csv?year={today.year}",
        f"/api/reports/crm.csv?date_from={today}&date_to={today}",
    ):
        response = client.get(path, headers=headers)
        assert response.status_code == 200, response.text
        assert "text/csv" in response.headers["content-type"]
        assert response.content.startswith(b"\xef\xbb\xbf")


def test_phase6_report_permissions_preserve_hr_privacy_and_crm_read_only_access():
    admin_headers = auth_headers()
    _, hr_basic_headers = _create_hr_user(
        admin_headers, "hr-basic-report@example.com", "HR Alap",
        [PERMISSION_HR_ACCESS, PERMISSION_HR_LEAVE_SELF],
    )
    denied = client.get(f"/api/reports/hr?year={date.today().year}", headers=hr_basic_headers)
    assert denied.status_code == 403
    assert denied.json()["detail"]["permission"] == PERMISSION_HR_SUMMARY_VIEW

    _, crm_read_headers = _create_crm_user("crm-report-read@example.com", [PERMISSION_CRM_ACCESS])
    allowed = client.get(f"/api/reports/crm?date_from={date.today()}&date_to={date.today()}", headers=crm_read_headers)
    assert allowed.status_code == 200, allowed.text
    export = client.get("/api/reports/crm-opportunities.csv", headers=crm_read_headers)
    assert export.status_code == 200, export.text
    assert "text/csv" in export.headers["content-type"]

    service_denied = client.get(f"/api/reports/service?date_from={date.today()}&date_to={date.today()}", headers=crm_read_headers)
    assert service_denied.status_code == 403
    assert service_denied.json()["detail"]["permission"] == PERMISSION_SERVICE_ACCESS


def test_synthetic_hr_rlb_source_totals_are_consistent():
    source_path = FIXTURE_DIR / "hr_leave_2026_rlb.json"
    payload = json.loads(source_path.read_text(encoding="utf-8"))
    assert len(payload["people"]) == 1
    for person in payload["people"]:
        annual = sum(row[3] for row in person["absences"] if row[0] == "annual_leave")
        sick = sum(row[3] for row in person["absences"] if row[0] == "sick_leave")
        assert person["full_name"] == "Minta Dolgozó"
        assert (person["base_days"], person["child_days"], person["sick_leave_days"], annual, sick) == (30, 0, 15, 12, 3)


def test_hr_rlb_import_applies_historical_leave_without_requiring_approval():
    from app.services.hr_import import auto_apply_rlb_2026_imports
    from app.services.hr_leave import leave_balance

    db = TestingSessionLocal()
    role = db.query(Role).filter(Role.name == ROLE_OFFICE).one()
    user = User(
        email="minta.dolgozo@example.com",
        full_name="Minta Dolgozó",
        password_hash=get_password_hash("Teszt!123"),
        role_id=role.id,
    )
    db.add(user)
    db.commit()

    stats = auto_apply_rlb_2026_imports(db)
    assert stats["applied"] == 1
    profile = db.query(EmployeeProfile).filter(EmployeeProfile.user_id == user.id).one()
    assert profile.employee_number == "TEST-001"
    assert profile.employment_start_date == date(2024, 7, 11)
    assert profile.job_title == "Teszt munkakör"

    entitlement = db.query(LeaveEntitlement).filter(LeaveEntitlement.employee_id == profile.id, LeaveEntitlement.year == 2026).one()
    assert float(entitlement.base_days) == 30
    assert float(entitlement.child_days) == 0
    assert float(entitlement.sick_leave_days) == 15

    annual_rows = db.query(LeaveRequest).filter(LeaveRequest.employee_id == profile.id, LeaveRequest.leave_type == "annual_leave").all()
    sick_rows = db.query(LeaveRequest).filter(LeaveRequest.employee_id == profile.id, LeaveRequest.leave_type == "sick_leave").all()
    assert sum(float(row.requested_days) for row in annual_rows) == 12
    assert sum(float(row.requested_days) for row in sick_rows) == 3
    assert all(row.status == "approved" and row.approver_user_id is None and row.source == "rlb_pdf_2026" for row in annual_rows + sick_rows)

    balance = leave_balance(db, profile, 2026)
    assert float(balance["total_days"]) == 30
    assert float(balance["approved_days"]) == 12
    assert float(balance["remaining_days"]) == 18
    assert float(balance["sick_leave_used_days"]) == 3
    assert float(balance["sick_leave_remaining_days"]) == 12
    db.close()


def test_contract_excel_fields_service_lines_and_pdf_archive_admin_delete():
    admin_headers = auth_headers()
    _, manager_headers = _create_crm_user("contract-archive@example.com", [PERMISSION_CRM_ACCESS, PERMISSION_CRM_CONTRACT_MANAGE])
    db = TestingSessionLocal(); customer = db.query(Customer).first(); customer_id = customer.id; db.close()
    created = client.post('/api/crm/contracts', headers=manager_headers, json={
        'customer_id': customer_id, 'company_code':'GXR', 'contract_number':'EXCEL-001', 'contract_type':'Vállalkozási szerződés',
        'contracting_party_name':'Teszt ügyfél Kft.', 'contracting_party_address':'1111 Budapest, Teszt u. 1.',
        'subject':'Borítékológép javítása, karbantartása', 'start_date':'2026-01-01', 'term_text':'határozatlan',
        'billing_frequency':'havi', 'payment_terms':'30 nap', 'billing_name':'Teszt ügyfél Kft.', 'service_confirmation':'munkalap',
        'e_invoice_email':'szamla@example.com', 'price_adjustment_terms':'KSH fogyasztói árindex',
        'service_lines':[{
            'rental_fee':'10.000 Ft / hó','operating_fee':'20.000 Ft / hó','click_fee':'1,5 Ft / boríték','included_quantity':'1000 db',
            'operator_company':'Teszt ügyfél Kft.','operator_site':'1111 Budapest','maintenance_cycle':'3 hó','device_group':'borítékológép',
            'device_name':'PB DI380','serial_number':'TEST-SERIAL-113','parts_terms':'tartalmazza','toner_terms':'-','travel_labor_terms':'tartalmazza',
            'repair_terms':'tartalmazza','sla':'8 óra','replacement_device':'igen'
        }]
    })
    assert created.status_code == 201, created.text
    row = created.json(); assert row['company_code'] == 'GXR'; assert row['subject'].startswith('Borítékológép')
    assert row['service_lines'][0]['device_name'] == 'PB DI380'; assert row['service_lines'][0]['maintenance_cycle'] == '3 hó'

    pdf = b'%PDF-1.4\n% contract archive\n%%EOF\n'
    upload = client.post(f"/api/crm/contracts/{row['id']}/archives", headers=manager_headers, data={'note':'Aláírt papír szerződés'}, files={'archive_file':('szerzodes.pdf',pdf,'application/pdf')})
    assert upload.status_code == 201, upload.text
    archive = upload.json(); assert archive['archive_number'].startswith('SZ-A-'); assert archive['original_filename'] == 'szerzodes.pdf'
    listing = client.get(f"/api/crm/contracts/{row['id']}/archives", headers=manager_headers)
    assert listing.status_code == 200 and len(listing.json()) == 1
    preview = client.get(f"/api/crm/contract-archives/{archive['id']}/preview", headers=manager_headers)
    assert preview.status_code == 200 and preview.content == pdf and 'application/pdf' in preview.headers['content-type']
    download = client.get(f"/api/crm/contract-archives/{archive['id']}/download", headers=manager_headers)
    assert download.status_code == 200 and download.content == pdf
    denied = client.request('DELETE', f"/api/crm/contract-archives/{archive['id']}", headers=manager_headers, json={'confirmation':'TÖRÖL'})
    assert denied.status_code == 403
    bad = client.request('DELETE', f"/api/crm/contract-archives/{archive['id']}", headers=admin_headers, json={'confirmation':'torol'})
    assert bad.status_code == 400
    deleted = client.request('DELETE', f"/api/crm/contract-archives/{archive['id']}", headers=admin_headers, json={'confirmation':'TÖRÖL'})
    assert deleted.status_code == 204


def test_hr_summary_calendar_requires_summary_permission_and_returns_leave_without_comments():
    admin_headers = auth_headers()
    user_id, summary_headers = _create_hr_user(admin_headers, 'calendar-summary@example.com', 'Naptár Összesítő', [PERMISSION_HR_ACCESS, PERMISSION_HR_SUMMARY_VIEW])
    _, basic_headers = _create_hr_user(admin_headers, 'calendar-basic@example.com', 'Naptár Alap', [PERMISSION_HR_ACCESS, PERMISSION_HR_LEAVE_SELF])
    db = TestingSessionLocal(); admin = db.query(User).filter(User.email=='admin@example.com').one()
    profile = EmployeeProfile(user_id=admin.id, employee_number='CAL-1', company_code='DRH', active=True); db.add(profile); db.flush()
    db.add(LeaveRequest(employee_id=profile.id, approver_user_id=None, start_date=date(2026,8,17), end_date=date(2026,8,21), requested_days=5, status='approved', leave_type='annual_leave', employee_comment='titkos megjegyzés', approver_comment='titkos vezetői megjegyzés'))
    db.commit(); db.close()
    denied = client.get('/api/hr/summary/calendar?date_from=2026-08-01&date_to=2026-08-31', headers=basic_headers)
    assert denied.status_code == 403
    allowed = client.get('/api/hr/summary/calendar?date_from=2026-08-01&date_to=2026-08-31', headers=summary_headers)
    assert allowed.status_code == 200, allowed.text
    events = allowed.json()['events']; assert len(events) == 1; assert events[0]['full_name'] == 'Admin'; assert events[0]['days'] == 5
    assert 'employee_comment' not in events[0] and 'approver_comment' not in events[0]


def _create_procurement_test_user(email: str, name: str, *, permission_codes=None):
    admin_headers = auth_headers()
    db = TestingSessionLocal()
    role_id = db.query(Role).filter(Role.name == ROLE_OFFICE).one().id
    db.close()
    password = "Ideiglenes!A1"
    response = client.post(
        "/api/users",
        headers=admin_headers,
        json={
            "email": email,
            "full_name": name,
            "password": password,
            "role_id": role_id,
            "is_active": True,
            "permission_codes": permission_codes or [],
        },
    )
    assert response.status_code == 201, response.text
    headers = auth_headers_for(email, password)
    changed = client.post(
        "/api/auth/change-password",
        headers=headers,
        json={"current_password": password, "new_password": "Vegleges!B2"},
    )
    assert changed.status_code == 200, changed.text
    return response.json(), headers


def _create_procurement_request(headers, *, subject="Teszt beszerzés", amount="125000.00"):
    response = client.post(
        "/api/procurement/requests",
        headers=headers,
        json={
            "subject": subject,
            "purpose": "A működéshez szükséges beszerzés",
            "description": "Részletes beszerzési igény tesztelési célból.",
            "total_amount": amount,
            "currency": "HUF",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_procurement_module_is_available_to_every_active_user_and_approval_is_separate_permission():
    requester, requester_headers = _create_procurement_test_user(
        "proc-requester@example.com", "Beszerzés Kérelmező"
    )
    assert requester["modules"] == ["procurement"]
    assert PERMISSION_PROCUREMENT_APPROVE not in requester["permissions"]

    approver, approver_headers = _create_procurement_test_user(
        "proc-approver@example.com", "Beszerzés Jóváhagyó", permission_codes=[PERMISSION_PROCUREMENT_APPROVE]
    )
    assert approver["modules"] == ["procurement"]
    assert PERMISSION_PROCUREMENT_APPROVE in approver["permissions"]

    row = _create_procurement_request(requester_headers)
    assert row["request_number"].startswith("BESZ-")
    assert row["status"] == "pending"
    assert row["requester"]["full_name"] == "Beszerzés Kérelmező"

    own = client.get("/api/procurement/requests", headers=requester_headers)
    assert own.status_code == 200, own.text
    assert [item["id"] for item in own.json()["items"]] == [row["id"]]

    blocked = client.get("/api/procurement/approvals", headers=requester_headers)
    assert blocked.status_code == 403
    assert blocked.json()["detail"]["permission"] == PERMISSION_PROCUREMENT_APPROVE

    approvals = client.get("/api/procurement/approvals?status=pending", headers=approver_headers)
    assert approvals.status_code == 200, approvals.text
    assert any(item["id"] == row["id"] for item in approvals.json()["items"])


def test_procurement_approval_rejection_withdrawal_and_self_approval_guard():
    requester, requester_headers = _create_procurement_test_user(
        "workflow-requester@example.com", "Workflow Kérelmező"
    )
    approver, approver_headers = _create_procurement_test_user(
        "workflow-approver@example.com", "Workflow Jóváhagyó", permission_codes=[PERMISSION_PROCUREMENT_APPROVE]
    )

    approved_row = _create_procurement_request(requester_headers, subject="Jóváhagyandó")
    approved = client.post(
        f"/api/procurement/requests/{approved_row['id']}/approve",
        headers=approver_headers,
        json={"version": approved_row["version"], "comment": "Rendben."},
    )
    assert approved.status_code == 200, approved.text
    assert approved.json()["status"] == "approved"
    assert approved.json()["approver"]["id"] == approver["id"]

    second_decision = client.post(
        f"/api/procurement/requests/{approved_row['id']}/approve",
        headers=approver_headers,
        json={"version": approved_row["version"], "comment": "Dupla döntés"},
    )
    assert second_decision.status_code == 409

    reject_row = _create_procurement_request(requester_headers, subject="Elutasítandó")
    missing_comment = client.post(
        f"/api/procurement/requests/{reject_row['id']}/reject",
        headers=approver_headers,
        json={"version": reject_row["version"]},
    )
    assert missing_comment.status_code == 422
    rejected = client.post(
        f"/api/procurement/requests/{reject_row['id']}/reject",
        headers=approver_headers,
        json={"version": reject_row["version"], "comment": "Nincs rá keret."},
    )
    assert rejected.status_code == 200, rejected.text
    assert rejected.json()["status"] == "rejected"

    withdraw_row = _create_procurement_request(requester_headers, subject="Visszavonandó")
    withdrawn = client.post(
        f"/api/procurement/requests/{withdraw_row['id']}/withdraw",
        headers=requester_headers,
        json={"version": withdraw_row["version"]},
    )
    assert withdrawn.status_code == 200, withdrawn.text
    assert withdrawn.json()["status"] == "withdrawn"

    own_approval_row = _create_procurement_request(approver_headers, subject="Saját jóváhagyás tiltva")
    own_approval = client.post(
        f"/api/procurement/requests/{own_approval_row['id']}/approve",
        headers=approver_headers,
        json={"version": own_approval_row["version"]},
    )
    assert own_approval.status_code == 403
    assert "Saját beszerzési igény" in own_approval.json()["detail"]


def test_procurement_attachments_are_protected_and_limited_to_pending_owner_uploads():
    requester, requester_headers = _create_procurement_test_user(
        "file-requester@example.com", "Fájl Kérelmező"
    )
    approver, approver_headers = _create_procurement_test_user(
        "file-approver@example.com", "Fájl Jóváhagyó", permission_codes=[PERMISSION_PROCUREMENT_APPROVE]
    )
    outsider, outsider_headers = _create_procurement_test_user(
        "file-outsider@example.com", "Fájl Kívülálló"
    )
    row = _create_procurement_request(requester_headers, subject="Csatolmányos")

    uploaded = client.post(
        f"/api/procurement/requests/{row['id']}/attachments",
        headers=requester_headers,
        files={"files": ("ajanlat.pdf", b"%PDF-1.4\n% procurement test\n", "application/pdf")},
    )
    assert uploaded.status_code == 201, uploaded.text
    attachment = uploaded.json()[0]
    assert attachment["original_filename"] == "ajanlat.pdf"
    assert len(attachment["sha256"]) == 64

    owner_download = client.get(f"/api/procurement/attachments/{attachment['id']}/download", headers=requester_headers)
    assert owner_download.status_code == 200
    assert owner_download.content.startswith(b"%PDF-")

    approver_download = client.get(f"/api/procurement/attachments/{attachment['id']}/download", headers=approver_headers)
    assert approver_download.status_code == 200

    outsider_download = client.get(f"/api/procurement/attachments/{attachment['id']}/download", headers=outsider_headers)
    assert outsider_download.status_code == 403

    forbidden_type = client.post(
        f"/api/procurement/requests/{row['id']}/attachments",
        headers=requester_headers,
        files={"files": ("malware.exe", b"MZ", "application/octet-stream")},
    )
    assert forbidden_type.status_code == 415

    detail = client.get(f"/api/procurement/requests/{row['id']}", headers=requester_headers).json()
    approved = client.post(
        f"/api/procurement/requests/{row['id']}/approve",
        headers=approver_headers,
        json={"version": detail["version"]},
    )
    assert approved.status_code == 200, approved.text
    upload_after_decision = client.post(
        f"/api/procurement/requests/{row['id']}/attachments",
        headers=requester_headers,
        files={"files": ("masik.pdf", b"%PDF-1.4\n", "application/pdf")},
    )
    assert upload_after_decision.status_code == 409



def _create_mobile_technician(email: str = "tech@example.com", password: str = "TechPassword123!"):
    db = TestingSessionLocal()
    role = db.query(Role).filter(Role.name == ROLE_TECHNICIAN).one()
    user = User(email=email, full_name="Teszt Technikus", password_hash=get_password_hash(password), role_id=role.id)
    db.add(user)
    db.flush()
    grant_service_access(db, user)
    customer = db.query(Customer).first()
    location = db.query(Location).filter(Location.customer_id == customer.id).first()
    order = WorkOrder(
        number=f"MOB-{user.id:04d}", customer_id=customer.id, location_id=location.id,
        technician_id=user.id, status="ütemezve", priority="normál",
        description="Mobil aláírás teszt", planned_date=date.today(), work_type="javítás",
    )
    db.add(order)
    db.commit()
    user_id, order_id = user.id, order.id
    db.close()
    return user_id, order_id, password


def _mobile_login(email: str, password: str, device_uuid: str = "android-test-device-001"):
    return client.post("/api/mobile/auth/login", json={
        "email": email,
        "password": password,
        "device_uuid": device_uuid,
        "device_name": "Teszt Android",
        "platform": "android",
        "app_version": "0.1-test",
    })


def _signature_png_bytes() -> bytes:
    image = Image.new("RGBA", (800, 240), "white")
    draw = ImageDraw.Draw(image)
    draw.line([(70, 150), (180, 70), (250, 165), (350, 75), (470, 150), (650, 85)], fill="black", width=8)
    output = BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()




def _photo_jpeg_bytes() -> bytes:
    image = Image.new("RGB", (1280, 720), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((100, 100, 600, 500), outline="black", width=12)
    draw.text((140, 140), "METER 12345", fill="black")
    output = BytesIO()
    image.save(output, format="JPEG", quality=90)
    return output.getvalue()


def test_mobile_login_refresh_rotation_and_replay_revokes_session():
    _, _, password = _create_mobile_technician()
    login = _mobile_login("tech@example.com", password)
    assert login.status_code == 200, login.text
    first = login.json()
    assert first["access_token"] and first["refresh_token"] and first["device_id"]
    mobile_headers = {"Authorization": f"Bearer {first['access_token']}"}
    assert client.get("/api/mobile/me", headers=mobile_headers).status_code == 200
    config = client.get("/api/mobile/config", headers=mobile_headers)
    assert config.status_code == 200
    assert config.json()["acceptance_text"]
    assert config.json()["photo_max_size_mb"] >= 1
    assert "meter" in config.json()["photo_categories"]

    refreshed = client.post("/api/mobile/auth/refresh", json={"refresh_token": first["refresh_token"]})
    assert refreshed.status_code == 200, refreshed.text
    second = refreshed.json()
    assert second["refresh_token"] != first["refresh_token"]

    replay = client.post("/api/mobile/auth/refresh", json={"refresh_token": first["refresh_token"]})
    # A replayelt refresh token hitelesitesi hiba: a session visszavonasra kerul.
    assert replay.status_code == 401
    # A bizonyitott refresh-token replay az egesz mobil sessiont visszavonja.
    assert client.get("/api/mobile/me", headers={"Authorization": f"Bearer {second['access_token']}"}).status_code == 401


def test_mobile_work_order_completion_signature_pdf_and_revision_integrity(tmp_path):
    _, order_id, password = _create_mobile_technician()
    login = _mobile_login("tech@example.com", password)
    assert login.status_code == 200, login.text
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

    detail = client.get(f"/api/mobile/work-orders/{order_id}", headers=headers)
    assert detail.status_code == 200, detail.text
    version = detail.json()["version"]

    started = client.post(f"/api/mobile/work-orders/{order_id}/start", headers=headers, json={"version": version})
    assert started.status_code == 200, started.text
    version = started.json()["version"]

    settings = get_settings()
    old_runtime = settings.runtime_dir
    old_signature_dir = settings.mobile_signature_dir
    try:
        settings.runtime_dir = str(tmp_path)
        settings.mobile_signature_dir = ""
        completed = client.post(f"/api/mobile/work-orders/{order_id}/complete", headers=headers, json={
            "version": version,
            "work_done": "A hibás egység ellenőrizve és javítva.",
            "final_status": "Működőképes",
            "labor_hours": "1.25",
            "customer_note": "A berendezés átadva.",
        })
        assert completed.status_code == 200, completed.text
        snapshot = completed.json()
        assert snapshot["snapshot_sha256"] and len(snapshot["snapshot_sha256"]) == 64
        assert snapshot["snapshot"]["work_order"]["work_done"].startswith("A hibás")

        signed = client.post(
            f"/api/mobile/work-orders/{order_id}/signature",
            headers=headers,
            data={"snapshot_id": str(snapshot["id"]), "signer_name": "Kovács János", "signer_role": "üzemeltetési vezető"},
            files={"signature": ("signature.png", _signature_png_bytes(), "image/png")},
        )
        assert signed.status_code == 200, signed.text
        signature = signed.json()
        assert len(signature["signed_pdf_sha256"]) == 64

        pdf = client.get(f"/api/mobile/work-orders/{order_id}/signed-pdf", headers=headers)
        assert pdf.status_code == 200, pdf.text
        assert pdf.content.startswith(b"%PDF-")
        assert hashlib.sha256(pdf.content).hexdigest() == signature["signed_pdf_sha256"]

        from app.services.backup_restore import create_backup_archive
        backup_bytes, _, _ = create_backup_archive(engine, settings)
        with zipfile.ZipFile(BytesIO(backup_bytes)) as archive:
            evidence_names = [name for name in archive.namelist() if name.startswith("work_order_signatures/")]
            assert len(evidence_names) == 2
            assert any(name.endswith("-signature.png") for name in evidence_names)
            assert any(name.endswith("-signed.pdf") for name in evidence_names)

        current = client.get(f"/api/mobile/work-orders/{order_id}", headers=headers)
        assert current.status_code == 200
        assert current.json()["status"] == "lezárva"
        assert current.json()["signature_current"] is True

        db = TestingSessionLocal()
        order = db.get(WorkOrder, order_id)
        order.customer_note = "Utólagos admin módosítás teszt"
        db.commit()
        db.close()
        changed = client.get(f"/api/mobile/work-orders/{order_id}", headers=headers)
        assert changed.status_code == 200
        assert changed.json()["signature_current"] is False
    finally:
        settings.runtime_dir = old_runtime
        settings.mobile_signature_dir = old_signature_dir


def test_mobile_signature_rejects_blank_canvas(tmp_path):
    _, order_id, password = _create_mobile_technician()
    login = _mobile_login("tech@example.com", password)
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    detail = client.get(f"/api/mobile/work-orders/{order_id}", headers=headers).json()
    started = client.post(f"/api/mobile/work-orders/{order_id}/start", headers=headers, json={"version": detail["version"]}).json()
    settings = get_settings(); old_runtime = settings.runtime_dir
    try:
        settings.runtime_dir = str(tmp_path)
        completed = client.post(f"/api/mobile/work-orders/{order_id}/complete", headers=headers, json={"version": started["version"], "work_done": "Teszt javítás"})
        assert completed.status_code == 200
        blank = Image.new("RGBA", (800, 240), "white"); output = BytesIO(); blank.save(output, format="PNG")
        response = client.post(
            f"/api/mobile/work-orders/{order_id}/signature", headers=headers,
            data={"snapshot_id": str(completed.json()["id"]), "signer_name": "Teszt Aláíró"},
            files={"signature": ("blank.png", output.getvalue(), "image/png")},
        )
        assert response.status_code == 400
    finally:
        settings.runtime_dir = old_runtime


def test_mobile_technician_cannot_access_another_technicians_work_order():
    _, first_order_id, first_password = _create_mobile_technician(
        email="mobile-tech-one@example.com", password="TechOnePassword123!"
    )
    _, second_order_id, _ = _create_mobile_technician(
        email="mobile-tech-two@example.com", password="TechTwoPassword123!"
    )
    login = _mobile_login("mobile-tech-one@example.com", first_password, device_uuid="android-tech-one")
    assert login.status_code == 200, login.text
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

    own = client.get(f"/api/mobile/work-orders/{first_order_id}", headers=headers)
    assert own.status_code == 200, own.text
    forbidden = client.get(f"/api/mobile/work-orders/{second_order_id}", headers=headers)
    assert forbidden.status_code == 403

    today = client.get(f"/api/mobile/work-orders/today?day={date.today().isoformat()}", headers=headers)
    assert today.status_code == 200, today.text
    returned_ids = {row["id"] for row in today.json()}
    assert first_order_id in returned_ids
    assert second_order_id not in returned_ids


def test_admin_can_revoke_lost_mobile_device_and_invalidate_access():
    user_id, _, password = _create_mobile_technician(
        email="lost-device-tech@example.com", password="LostDevicePassword123!"
    )
    login = _mobile_login("lost-device-tech@example.com", password, device_uuid="android-lost-device")
    assert login.status_code == 200, login.text
    payload = login.json()
    mobile_headers = {"Authorization": f"Bearer {payload['access_token']}"}
    assert client.get("/api/mobile/me", headers=mobile_headers).status_code == 200

    admin_headers = auth_headers()
    devices = client.get(f"/api/users/{user_id}/mobile-devices", headers=admin_headers)
    assert devices.status_code == 200, devices.text
    assert any(row["id"] == payload["device_id"] for row in devices.json())

    revoked = client.delete(
        f"/api/users/{user_id}/mobile-devices/{payload['device_id']}", headers=admin_headers
    )
    assert revoked.status_code == 204, revoked.text
    assert client.get("/api/mobile/me", headers=mobile_headers).status_code == 401
    refresh = client.post("/api/mobile/auth/refresh", json={"refresh_token": payload["refresh_token"]})
    assert refresh.status_code == 401


def test_v026_production_security_validates_mobile_timezone_and_acceptance_contract():
    settings = get_settings()
    original = {
        "environment": settings.environment,
        "mobile_access_token_expire_minutes": settings.mobile_access_token_expire_minutes,
        "mobile_refresh_token_expire_days": settings.mobile_refresh_token_expire_days,
        "mobile_signature_max_size_kb": settings.mobile_signature_max_size_kb,
        "mobile_photo_max_size_mb": settings.mobile_photo_max_size_mb,
        "mobile_photo_max_pixels": settings.mobile_photo_max_pixels,
        "mobile_photo_max_dimension": settings.mobile_photo_max_dimension,
        "business_timezone": settings.business_timezone,
        "work_order_acceptance_text_version": settings.work_order_acceptance_text_version,
        "work_order_acceptance_text": settings.work_order_acceptance_text,
    }
    try:
        settings.environment = "production"
        settings.mobile_access_token_expire_minutes = 31
        settings.mobile_refresh_token_expire_days = 91
        settings.mobile_signature_max_size_kb = 4097
        settings.mobile_photo_max_size_mb = 26
        settings.mobile_photo_max_pixels = 40_000_001
        settings.mobile_photo_max_dimension = 4097
        settings.business_timezone = "Not/A-Timezone"
        settings.work_order_acceptance_text_version = ""
        settings.work_order_acceptance_text = ""
        errors = production_security_errors(settings)
        assert "MOBILE_ACCESS_TOKEN_EXPIRE_MINUTES must be between 1 and 30 in production" in errors
        assert "MOBILE_REFRESH_TOKEN_EXPIRE_DAYS must be between 1 and 90 in production" in errors
        assert "MOBILE_SIGNATURE_MAX_SIZE_KB must be between 1 and 4096 in production" in errors
        assert "MOBILE_PHOTO_MAX_SIZE_MB must be between 1 and 25 in production" in errors
        assert "MOBILE_PHOTO_MAX_PIXELS must be between 1000000 and 40000000 in production" in errors
        assert "MOBILE_PHOTO_MAX_DIMENSION must be between 640 and 4096 in production" in errors
        assert "BUSINESS_TIMEZONE must be a valid IANA timezone in production" in errors
        assert "WORK_ORDER_ACCEPTANCE_TEXT_VERSION must not be empty in production" in errors
        assert "WORK_ORDER_ACCEPTANCE_TEXT must contain the reviewed customer acceptance statement in production" in errors
    finally:
        for key, value in original.items():
            setattr(settings, key, value)


def test_mobile_photo_is_sanitized_snapshotted_and_locked_after_completion(tmp_path):
    _, order_id, password = _create_mobile_technician(email="photo-tech@example.com", password="PhotoTechPassword123!")
    login = _mobile_login("photo-tech@example.com", password, device_uuid="android-photo-test")
    assert login.status_code == 200, login.text
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    settings = get_settings()
    old_runtime = settings.runtime_dir
    old_photo_dir = settings.mobile_photo_dir
    try:
        settings.runtime_dir = str(tmp_path)
        settings.mobile_photo_dir = ""
        detail = client.get(f"/api/mobile/work-orders/{order_id}", headers=headers).json()
        started = client.post(f"/api/mobile/work-orders/{order_id}/start", headers=headers, json={"version": detail["version"]})
        assert started.status_code == 200, started.text
        upload = client.post(
            f"/api/mobile/work-orders/{order_id}/photos", headers=headers,
            data={"category": "meter", "note": "Átadás előtti számláló"},
            files={"photo": ("meter.jpg", _photo_jpeg_bytes(), "image/jpeg")},
        )
        assert upload.status_code == 201, upload.text
        row = upload.json()
        assert row["category"] == "meter"
        assert len(row["sha256"]) == 64
        assert row["mime_type"] == "image/jpeg"
        content = client.get(f"/api/mobile/work-orders/{order_id}/photos/{row['id']}/content", headers=headers)
        assert content.status_code == 200
        assert content.content.startswith(b"\xff\xd8\xff")
        completed = client.post(
            f"/api/mobile/work-orders/{order_id}/complete", headers=headers,
            json={"version": started.json()["version"], "work_done": "Fotóval dokumentált javítás"},
        )
        assert completed.status_code == 200, completed.text
        photos = completed.json()["snapshot"]["photos"]
        assert len(photos) == 1
        assert photos[0]["sha256"] == row["sha256"]
        blocked = client.delete(f"/api/mobile/work-orders/{order_id}/photos/{row['id']}", headers=headers)
        assert blocked.status_code == 409
        from app.services.backup_restore import create_backup_archive
        backup_bytes, _, _ = create_backup_archive(engine, settings)
        with zipfile.ZipFile(BytesIO(backup_bytes)) as archive:
            photo_names = [name for name in archive.namelist() if name.startswith("work_order_photos/")]
            assert len(photo_names) == 1
        stored_files = list((tmp_path / "work_order_photos").rglob("*.jpg"))
        assert len(stored_files) == 1
        stored_files[0].unlink()
        from app.services.backup_restore import restore_backup_archive
        restore_backup_archive(engine, settings, backup_bytes)
        restored_files = list((tmp_path / "work_order_photos").rglob("*.jpg"))
        assert len(restored_files) == 1
        assert restored_files[0].read_bytes().startswith(b"\xff\xd8\xff")
    finally:
        settings.runtime_dir = old_runtime
        settings.mobile_photo_dir = old_photo_dir



def test_v028_mobile_push_registration_and_work_order_assignment_outbox():
    user_id, _, password = _create_mobile_technician(
        email="push-tech@example.com", password="PushTechPassword123!"
    )
    login = _mobile_login("push-tech@example.com", password, device_uuid="android-push-device")
    assert login.status_code == 200, login.text
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    registered = client.put(
        "/api/mobile/devices/push-token",
        headers=headers,
        json={"token": "fcm-registration-token-abcdefghijklmnopqrstuvwxyz-1234567890", "enabled": True},
    )
    assert registered.status_code == 200, registered.text
    assert registered.json()["push_enabled"] is True
    assert registered.json()["push_token_updated_at"]

    db = TestingSessionLocal()
    customer = db.query(Customer).first()
    location = db.query(Location).filter(Location.customer_id == customer.id).first()
    customer_id, location_id = customer.id, location.id
    db.close()

    created = client.post(
        "/api/work-orders",
        headers=auth_headers(),
        json={
            "customer_id": customer_id,
            "location_id": location_id,
            "technician_id": user_id,
            "description": "Push kiosztás teszt",
            "priority": "normál",
            "status": "ütemezve",
            "planned_date": date.today().isoformat(),
        },
    )
    assert created.status_code == 201, created.text
    order = created.json()

    db = TestingSessionLocal()
    rows = db.query(PushDeliveryOutbox).filter(PushDeliveryOutbox.user_id == user_id).all()
    assert len(rows) == 1
    assert rows[0].status == "pending"
    assert rows[0].title == "Új munkalapot kaptál"
    assert rows[0].data["source_type"] == "work_order"
    assert rows[0].data["source_id"] == str(order["id"])
    db.close()

    changed = client.put(
        f"/api/work-orders/{order['id']}",
        headers=auth_headers(),
        json={"version": order["version"], "planned_date": (date.today()).isoformat(), "planned_start_at": "2026-09-07T10:30:00"},
    )
    assert changed.status_code == 200, changed.text
    db = TestingSessionLocal()
    rows = db.query(PushDeliveryOutbox).filter(PushDeliveryOutbox.user_id == user_id).order_by(PushDeliveryOutbox.id).all()
    assert len(rows) == 2
    assert rows[-1].title == "Módosult a munkalap időpontja"
    db.close()

    disabled = client.delete("/api/mobile/devices/push-token", headers=headers)
    assert disabled.status_code == 204, disabled.text
    db = TestingSessionLocal()
    device = db.query(MobileDevice).filter(MobileDevice.id == login.json()["device_id"]).one()
    assert device.push_token is None
    assert device.push_enabled is False
    assert all(row.status == "cancelled" for row in db.query(PushDeliveryOutbox).filter(PushDeliveryOutbox.user_id == user_id).all())
    db.close()


def test_v028_backup_redacts_fcm_tokens_and_excludes_push_outbox(tmp_path):
    _, _, password = _create_mobile_technician(
        email="push-backup@example.com", password="PushBackupPassword123!"
    )
    login = _mobile_login("push-backup@example.com", password, device_uuid="android-push-backup")
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    assert client.put(
        "/api/mobile/devices/push-token",
        headers=headers,
        json={"token": "fcm-backup-token-abcdefghijklmnopqrstuvwxyz-123456789", "enabled": True},
    ).status_code == 200

    from app.services.notifications import create_notification
    db = TestingSessionLocal()
    device = db.get(MobileDevice, login.json()["device_id"])
    create_notification(
        db,
        user_id=device.user_id,
        module="service",
        notification_type="work_order_assigned",
        title="Backup push teszt",
        message="Nem kerülhet token a mentésbe",
        source_type="work_order",
        source_id=999,
        dedupe_key="test:backup:push:1",
    )
    db.commit()
    assert db.query(PushDeliveryOutbox).count() == 1
    db.close()

    from app.services.backup_restore import create_backup_archive
    settings = get_settings()
    backup_bytes, _, _ = create_backup_archive(engine, settings)
    with zipfile.ZipFile(BytesIO(backup_bytes)) as archive:
        payload = json.loads(archive.read("database.json"))
    assert "push_delivery_outbox" not in payload["tables"]
    mobile_rows = payload["tables"]["mobile_devices"]["rows"]
    row = next(item for item in mobile_rows if item["id"] == login.json()["device_id"])
    assert row["push_token"] is None
    assert row["push_enabled"] is False
    assert row["push_token_updated_at"] is None


def test_v028_production_security_requires_fcm_configuration_when_push_enabled():
    settings = get_settings()
    original = {
        "environment": settings.environment,
        "push_notifications_enabled": settings.push_notifications_enabled,
        "fcm_project_id": settings.fcm_project_id,
        "fcm_service_account_file": settings.fcm_service_account_file,
    }
    try:
        settings.environment = "production"
        settings.push_notifications_enabled = True
        settings.fcm_project_id = ""
        settings.fcm_service_account_file = ""
        errors = production_security_errors(settings)
        assert "FCM_PROJECT_ID is required when push notifications are enabled" in errors
        assert "FCM_SERVICE_ACCOUNT_FILE is required when push notifications are enabled" in errors
    finally:
        for key, value in original.items():
            setattr(settings, key, value)


def test_v028_push_worker_claims_and_marks_delivery_sent(monkeypatch):
    from app import push_worker

    db = TestingSessionLocal()
    user = db.query(User).filter(User.email == "admin@example.com").one()
    device = MobileDevice(
        user_id=user.id,
        device_uuid="worker-device-1",
        device_name="Worker phone",
        platform="android",
        app_version="0.28.0",
        push_token="fcm-worker-token-12345678901234567890",
        push_enabled=True,
        push_token_updated_at=datetime.utcnow(),
        last_seen_at=datetime.utcnow(),
    )
    db.add(device)
    db.flush()
    outbox = PushDeliveryOutbox(
        user_id=user.id,
        device_id=device.id,
        dedupe_key="worker:sent:1",
        title="Teszt push",
        body="Teszt",
        data={"source_type": "work_order", "source_id": "1"},
        status="pending",
        attempt_count=0,
        next_attempt_at=datetime.utcnow(),
        created_at=datetime.utcnow(),
    )
    db.add(outbox)
    db.commit()
    outbox_id = outbox.id
    db.close()

    monkeypatch.setattr(push_worker, "SessionLocal", TestingSessionLocal)
    monkeypatch.setattr(push_worker, "send_fcm_delivery", lambda delivery, device, settings: "projects/test/messages/123")

    claimed = push_worker.claim_delivery()
    assert claimed == outbox_id
    push_worker.process_delivery(claimed)

    db = TestingSessionLocal()
    row = db.get(PushDeliveryOutbox, outbox_id)
    assert row.status == "sent"
    assert row.attempt_count == 1
    assert row.provider_message_id == "projects/test/messages/123"
    assert row.sent_at is not None
    assert row.locked_at is None
    db.close()


def test_v028_push_worker_unregisters_invalid_fcm_token(monkeypatch):
    from app import push_worker
    from app.services.push_notifications import FcmSendError

    db = TestingSessionLocal()
    user = db.query(User).filter(User.email == "admin@example.com").one()
    device = MobileDevice(
        user_id=user.id,
        device_uuid="worker-device-2",
        device_name="Invalid token phone",
        platform="android",
        app_version="0.28.0",
        push_token="fcm-invalid-token-12345678901234567890",
        push_enabled=True,
        push_token_updated_at=datetime.utcnow(),
        last_seen_at=datetime.utcnow(),
    )
    db.add(device)
    db.flush()
    outbox = PushDeliveryOutbox(
        user_id=user.id,
        device_id=device.id,
        dedupe_key="worker:invalid:1",
        title="Teszt push",
        body="Teszt",
        data={},
        status="pending",
        attempt_count=0,
        next_attempt_at=datetime.utcnow(),
        created_at=datetime.utcnow(),
    )
    db.add(outbox)
    db.commit()
    outbox_id = outbox.id
    device_id = device.id
    db.close()

    monkeypatch.setattr(push_worker, "SessionLocal", TestingSessionLocal)

    def invalid_token(*args, **kwargs):
        raise FcmSendError("UNREGISTERED", permanent=True, unregister_token=True)

    monkeypatch.setattr(push_worker, "send_fcm_delivery", invalid_token)

    claimed = push_worker.claim_delivery()
    assert claimed == outbox_id
    push_worker.process_delivery(claimed)

    db = TestingSessionLocal()
    row = db.get(PushDeliveryOutbox, outbox_id)
    device = db.get(MobileDevice, device_id)
    assert row.status == "failed_permanent"
    assert row.last_error == "UNREGISTERED"
    assert device.push_token is None
    assert device.push_enabled is False
    assert device.push_token_updated_at is None
    db.close()


def test_v028_fcm_http_v1_payload_and_unregistered_detection(monkeypatch):
    import httpx
    from app.services import push_notifications as push_service

    settings = get_settings()
    old_enabled = settings.push_notifications_enabled
    old_project = settings.fcm_project_id
    old_channel = settings.push_android_channel_id
    try:
        settings.push_notifications_enabled = True
        settings.fcm_project_id = "firebase-test-project"
        settings.push_android_channel_id = "work_orders"
        monkeypatch.setattr(push_service, "_load_fcm_credentials", lambda current: "oauth-test-token")

        device = MobileDevice(
            id=123,
            user_id=1,
            device_uuid="fcm-payload-device",
            device_name="FCM phone",
            platform="android",
            app_version="0.28.0",
            push_token="fcm-payload-token-12345678901234567890",
            push_enabled=True,
            last_seen_at=datetime.utcnow(),
        )
        delivery = PushDeliveryOutbox(
            id=456,
            user_id=1,
            device_id=123,
            dedupe_key="fcm:payload:1",
            title="Új munkalapot kaptál",
            body="ML-2026-000001 – 2026.09.07. 10:30",
            data={"source_type": "work_order", "source_id": 42, "ignored": None},
            status="processing",
            attempt_count=1,
            next_attempt_at=datetime.utcnow(),
            created_at=datetime.utcnow(),
        )
        captured = {}

        class SuccessResponse:
            is_success = True
            status_code = 200
            text = ""
            def json(self):
                return {"name": "projects/firebase-test-project/messages/abc"}

        def fake_post(url, *, headers, json, timeout):
            captured.update(url=url, headers=headers, json=json, timeout=timeout)
            return SuccessResponse()

        monkeypatch.setattr(httpx, "post", fake_post)
        provider_id = push_service.send_fcm_delivery(delivery, device, settings)
        assert provider_id.endswith("/messages/abc")
        assert captured["url"] == "https://fcm.googleapis.com/v1/projects/firebase-test-project/messages:send"
        assert captured["headers"]["Authorization"] == "Bearer oauth-test-token"
        message = captured["json"]["message"]
        assert message["token"] == device.push_token
        assert message["android"]["notification"]["channel_id"] == "work_orders"
        assert message["data"]["source_id"] == "42"
        assert "ignored" not in message["data"]

        class UnregisteredResponse:
            is_success = False
            status_code = 404
            text = "not found"
            def json(self):
                return {
                    "error": {
                        "message": "Requested entity was not found.",
                        "details": [{"errorCode": "UNREGISTERED"}],
                    }
                }

        monkeypatch.setattr(httpx, "post", lambda *args, **kwargs: UnregisteredResponse())
        try:
            push_service.send_fcm_delivery(delivery, device, settings)
            assert False, "UNREGISTERED FCM response should fail"
        except push_service.FcmSendError as exc:
            assert exc.permanent is True
            assert exc.unregister_token is True
    finally:
        settings.push_notifications_enabled = old_enabled
        settings.fcm_project_id = old_project
        settings.push_android_channel_id = old_channel


def test_v028_push_scope_does_not_expose_hr_crm_or_procurement_events_to_technician_channel():
    from app.services.notifications import create_notification

    db = TestingSessionLocal()
    user = db.query(User).filter(User.email == "admin@example.com").one()
    device = MobileDevice(
        user_id=user.id,
        device_uuid="privacy-scope-device",
        device_name="Privacy scope phone",
        platform="android",
        app_version="0.28.0",
        push_token="fcm-privacy-token-12345678901234567890",
        push_enabled=True,
        push_token_updated_at=datetime.utcnow(),
        last_seen_at=datetime.utcnow(),
    )
    db.add(device)
    db.flush()
    create_notification(
        db,
        user_id=user.id,
        module="hr",
        notification_type="leave_approval_pending",
        title="Szabadság jóváhagyás",
        message="Érzékeny HR esemény",
        source_type="leave_request",
        source_id=1,
        dedupe_key="privacy:hr:1",
    )
    db.commit()
    assert db.query(PushDeliveryOutbox).filter(PushDeliveryOutbox.device_id == device.id).count() == 0
    db.close()


def test_v029_admin_mobile_device_status_and_test_push_queue_without_exposing_token():
    db = TestingSessionLocal()
    user = db.query(User).filter(User.email == "admin@example.com").one()
    device = MobileDevice(
        user_id=user.id,
        device_uuid="v029-admin-device",
        device_name="V029 telefon",
        platform="android",
        app_version="0.29.0",
        push_token="super-secret-fcm-token-12345678901234567890",
        push_enabled=True,
        push_token_updated_at=datetime.utcnow(),
        last_seen_at=datetime.utcnow(),
    )
    db.add(device)
    db.commit()
    user_id = user.id
    device_id = device.id
    db.close()

    headers = auth_headers()
    listing = client.get(f"/api/users/{user_id}/mobile-devices", headers=headers)
    assert listing.status_code == 200, listing.text
    rows = listing.json()
    assert len(rows) == 1
    assert rows[0]["id"] == device_id
    assert rows[0]["push_enabled"] is True
    assert "push_token" not in rows[0]
    assert rows[0]["last_push"] is None

    queued = client.post(f"/api/users/{user_id}/mobile-devices/{device_id}/test-push", headers=headers)
    assert queued.status_code == 202, queued.text
    delivery_id = queued.json()["delivery_id"]

    db = TestingSessionLocal()
    delivery = db.get(PushDeliveryOutbox, delivery_id)
    assert delivery is not None
    assert delivery.device_id == device_id
    assert delivery.user_id == user_id
    assert delivery.status == "pending"
    assert delivery.notification_id is None
    assert delivery.data["notification_type"] == "system_test"
    assert "secret" not in json.dumps(delivery.data).lower()
    db.close()

    refreshed = client.get(f"/api/users/{user_id}/mobile-devices", headers=headers)
    assert refreshed.status_code == 200
    last_push = refreshed.json()[0]["last_push"]
    assert last_push["id"] == delivery_id
    assert last_push["status"] == "pending"


def test_v029_admin_test_push_rejects_revoked_or_unregistered_device():
    db = TestingSessionLocal()
    user = db.query(User).filter(User.email == "admin@example.com").one()
    no_push = MobileDevice(
        user_id=user.id,
        device_uuid="v029-no-push",
        device_name="Nincs push",
        platform="android",
        app_version="0.29.0",
        push_enabled=False,
        last_seen_at=datetime.utcnow(),
    )
    revoked = MobileDevice(
        user_id=user.id,
        device_uuid="v029-revoked",
        device_name="Visszavont",
        platform="android",
        app_version="0.29.0",
        push_token="revoked-fcm-token-12345678901234567890",
        push_enabled=True,
        revoked_at=datetime.utcnow(),
        last_seen_at=datetime.utcnow(),
    )
    db.add_all([no_push, revoked])
    db.commit()
    user_id, no_push_id, revoked_id = user.id, no_push.id, revoked.id
    db.close()

    headers = auth_headers()
    response = client.post(f"/api/users/{user_id}/mobile-devices/{no_push_id}/test-push", headers=headers)
    assert response.status_code == 409
    assert "push" in response.json()["detail"].lower()

    response = client.post(f"/api/users/{user_id}/mobile-devices/{revoked_id}/test-push", headers=headers)
    assert response.status_code == 409
    assert "vissza" in response.json()["detail"].lower()


def _enable_admin_mfa_and_stale_session(headers: dict[str, str]) -> tuple[str, str, int]:
    settings = get_settings()
    settings.mfa_encryption_key = "v030-test-mfa-encryption-key-that-is-long-enough"
    token = headers["Authorization"].split(" ", 1)[1]
    session_id = str(decode_token_claims(token)["sid"])
    db = TestingSessionLocal()
    user = db.query(User).filter(User.email == "admin@example.com").one()
    secret = "JBSWY3DPEHPK3PXP"
    user.mfa_enabled = True
    user.mfa_secret_encrypted = encrypt_mfa_secret(secret, settings)
    user.mfa_last_counter = None
    session = db.query(AuthSession).filter(AuthSession.session_id == session_id).one()
    session.mfa_verified_at = datetime.utcnow() - timedelta(seconds=settings.sensitive_mfa_max_age_seconds + 30)
    db.commit()
    user_id = user.id
    db.close()
    return secret, session_id, user_id


def test_v030_high_risk_permission_requires_mfa_during_login():
    db = TestingSessionLocal()
    role = db.query(Role).filter(Role.name == ROLE_OFFICE).one()
    user = User(
        email="restore-operator@example.com",
        full_name="Restore Operator",
        password_hash=get_password_hash("Restore!Password123"),
        role_id=role.id,
    )
    db.add(user)
    db.flush()
    ensure_permission_catalog(db)
    set_user_permissions(db, user, [PERMISSION_ADMIN_ACCESS, PERMISSION_ADMIN_BACKUP_MANAGE, PERMISSION_ADMIN_BACKUP_RESTORE])
    db.commit()
    db.close()

    response = client.post(
        "/api/auth/login",
        data={"username": "restore-operator@example.com", "password": "Restore!Password123"},
    )
    assert response.status_code == 200
    assert response.json()["mfa_setup_required"] is True
    assert response.json()["access_token"] is None


def test_v030_step_up_refreshes_only_current_session_and_never_audits_code(monkeypatch):
    headers = auth_headers()
    secret, session_id, user_id = _enable_admin_mfa_and_stale_session(headers)
    db = TestingSessionLocal()
    user = db.get(User, user_id)
    other, _token, _csrf = create_auth_session(db, user, get_settings(), mfa_verified=False)
    other_id = other.session_id
    db.commit()
    db.close()

    code = _totp_at(secret, int(datetime.utcnow().timestamp()) // 30)
    response = client.post("/api/auth/mfa/step-up", headers=headers, json={"code": code})
    assert response.status_code == 200, response.text
    assert response.json()["ok"] is True

    db = TestingSessionLocal()
    current = db.query(AuthSession).filter(AuthSession.session_id == session_id).one()
    other = db.query(AuthSession).filter(AuthSession.session_id == other_id).one()
    assert current.mfa_verified_at is not None
    assert other.mfa_verified_at is None
    audit = db.query(AuditLog).filter(AuditLog.action == "MFA_STEP_UP_SUCCESS").order_by(AuditLog.id.desc()).first()
    assert audit is not None
    assert code not in json.dumps({"description": audit.description, "before": audit.before_data, "after": audit.after_data})
    db.close()

    monkeypatch.setattr("app.routers.backups.agent_post", lambda *_args, **_kwargs: {"restore_id": "RST-test"})
    restore = client.post(
        "/api/backups/MST-20260910-120000-000001/restore",
        headers=headers,
        json={"confirmation": "VISSZAALLIT"},
    )
    assert restore.status_code == 202, restore.text
    assert restore.json()["restore_id"] == "RST-test"

    replay = client.post("/api/auth/mfa/step-up", headers=headers, json={"code": code})
    # A web session tovabbra is ervenyes; a mar felhasznalt TOTP kod hibas input.
    assert replay.status_code == 400


def test_v030_sensitive_endpoints_reject_stale_mfa(monkeypatch):
    headers = auth_headers()
    _enable_admin_mfa_and_stale_session(headers)
    monkeypatch.setattr("app.routers.backups.agent_post", lambda *_args, **_kwargs: {"restore_id": "unexpected"})

    restore = client.post(
        "/api/backups/MST-20260910-120000-000001/restore",
        headers=headers,
        json={"confirmation": "VISSZAALLIT"},
    )
    assert restore.status_code == 403
    assert restore.json()["detail"]["code"] == "recent_mfa_required"

    recovery = client.post(
        "/api/backups/recovery/clear",
        headers=headers,
        json={"confirmation": "HELYREALLITAS_KESZ"},
    )
    assert recovery.status_code == 403

    settings_restore = client.post(
        "/api/settings/restore",
        headers=headers,
        data={"confirmation": "VISSZAÁLLÍTÁS"},
        files={"backup_file": ("backup.zip", b"not-a-zip", "application/zip")},
    )
    assert settings_restore.status_code == 403

    db = TestingSessionLocal()
    role_id = db.query(Role).filter(Role.name == ROLE_OFFICE).one().id
    db.close()
    create = client.post(
        "/api/users",
        headers=headers,
        json={
            "email": "blocked-create@example.com", "full_name": "Blocked Create",
            "password": "Blocked!Password123", "role_id": role_id,
            "permission_codes": [PERMISSION_SERVICE_ACCESS],
        },
    )
    assert create.status_code == 403


def test_v030_step_up_has_separate_rate_limit_bucket():
    headers = auth_headers()
    _enable_admin_mfa_and_stale_session(headers)
    settings = get_settings()
    old_attempts = settings.mfa_step_up_rate_limit_attempts
    try:
        settings.mfa_step_up_rate_limit_attempts = 2
        assert client.post("/api/auth/mfa/step-up", headers=headers, json={"code": "000000"}).status_code == 400
        assert client.post("/api/auth/mfa/step-up", headers=headers, json={"code": "000001"}).status_code == 400
        blocked = client.post("/api/auth/mfa/step-up", headers=headers, json={"code": "000002"})
        assert blocked.status_code == 429
        db = TestingSessionLocal()
        keys = [row.bucket_key for row in db.query(SecurityRateLimit).all()]
        db.close()
        assert any("mfa-step-up" in key for key in keys)
        assert all(not key.startswith("login:user:") for key in keys)
    finally:
        settings.mfa_step_up_rate_limit_attempts = old_attempts


def test_v031_production_security_validates_login_limits_and_proxy_host_cidrs():
    settings = get_settings()
    original = {
        "environment": settings.environment,
        "login_rate_limit_attempts": settings.login_rate_limit_attempts,
        "login_ip_rate_limit_attempts": settings.login_ip_rate_limit_attempts,
        "login_rate_limit_window_seconds": settings.login_rate_limit_window_seconds,
        "login_rate_limit_block_seconds": settings.login_rate_limit_block_seconds,
        "trusted_proxy_cidrs": settings.trusted_proxy_cidrs,
    }
    try:
        settings.environment = "production"
        settings.login_rate_limit_attempts = 2
        settings.login_ip_rate_limit_attempts = 201
        settings.login_rate_limit_window_seconds = 59
        settings.login_rate_limit_block_seconds = 299
        settings.trusted_proxy_cidrs = "172.30.250.0/24"
        errors = production_security_errors(settings)
        assert "LOGIN_RATE_LIMIT_ATTEMPTS must be between 3 and 10 in production" in errors
        assert "LOGIN_IP_RATE_LIMIT_ATTEMPTS must be between 10 and 200 in production" in errors
        assert "LOGIN_RATE_LIMIT_WINDOW_SECONDS must be between 60 and 1800 in production" in errors
        assert "LOGIN_RATE_LIMIT_BLOCK_SECONDS must be between 300 and 86400 in production" in errors
        assert "TRUSTED_PROXY_CIDRS must use host CIDRs only in production: 172.30.250.0/24" in errors
    finally:
        for key, value in original.items():
            setattr(settings, key, value)
