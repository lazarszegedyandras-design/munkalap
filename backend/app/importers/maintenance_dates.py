from __future__ import annotations

from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from io import BytesIO
import math
import unicodedata
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy.orm import Session

from ..models import Asset, User
from ..services.audit import add_asset_history

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
EXPECTED_SHEET_NAME = "Karbantartás import"
MAX_ERROR_SAMPLES = 100


@dataclass
class ParsedMaintenanceRow:
    excel_row: int
    import_batch: str | None
    import_row: int | None
    serial_number: str | None
    company_code: str | None
    customer_name: str | None
    asset_name: str | None
    maintenance_date: date
    source_sheet: str | None = None
    source_row: int | None = None
    source_work_order: str | None = None
    technician: str | None = None
    reason: str | None = None
    work_done: str | None = None
    note: str | None = None
    asset_id: int | None = None


@dataclass
class MaintenanceImportIssue:
    excel_row: int | None
    reason: str
    serial_number: str | None = None
    import_row: int | None = None
    maintenance_date: str | None = None


@dataclass
class MaintenanceImportResult:
    mode: str
    filename: str
    sheet_name: str
    total_rows: int = 0
    valid_rows: int = 0
    would_update_count: int = 0
    updated_count: int = 0
    unchanged_count: int = 0
    conflict_count: int = 0
    invalid_count: int = 0
    missing_asset_count: int = 0
    ambiguous_asset_count: int = 0
    affected_assets_count: int = 0
    skipped_count: int = 0
    issues: list[MaintenanceImportIssue] = field(default_factory=list)

    def add_issue(self, issue: MaintenanceImportIssue) -> None:
        if len(self.issues) < MAX_ERROR_SAMPLES:
            self.issues.append(issue)

    def as_dict(self) -> dict[str, Any]:
        payload = asdict(self)
        payload["issues_truncated"] = self.skipped_count > len(self.issues)
        return payload


def _normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    for char in (" ", "-", "/", ".", "(", ")"):
        text = text.replace(char, "_")
    while "__" in text:
        text = text.replace("__", "_")
    return text.strip("_")


HEADER_ALIASES = {
    "import_batch": "import_batch",
    "import_row": "import_row",
    "gyari_szam": "serial_number",
    "sorozatszam": "serial_number",
    "serial_number": "serial_number",
    "cegkod": "company_code",
    "ugyfel": "customer_name",
    "eszkoz": "asset_name",
    "utolso_karbantartas_datuma": "maintenance_date",
    "utolso_karbantartas": "maintenance_date",
    "karbantartas_datuma": "maintenance_date",
    "last_maintenance_date": "maintenance_date",
    "forras_munkalap": "source_sheet",
    "forras_excel_sor": "source_row",
    "forras_sor": "source_row",
    "munkalap_szama": "source_work_order",
    "forras_munkalapszam": "source_work_order",
    "technikus": "technician",
    "kiszallas_oka": "reason",
    "elvegzett_munka": "work_done",
    "megjegyzes": "note",
}


def _clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_int(value: Any, field_name: str) -> int:
    if value is None or value == "":
        raise ValueError(f"A(z) {field_name} mező üres")
    if isinstance(value, bool):
        raise ValueError(f"A(z) {field_name} mező nem szám")
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            raise ValueError(f"A(z) {field_name} mező érvénytelen")
        if float(value) != int(value):
            raise ValueError(f"A(z) {field_name} mező csak egész szám lehet")
        return int(value)
    normalized = str(value).replace(" ", "").replace("\u00a0", "").replace(",", ".").strip()
    number = float(normalized)
    if number != int(number):
        raise ValueError(f"A(z) {field_name} mező csak egész szám lehet")
    return int(number)


def _parse_optional_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    return _parse_int(value, "import_row")


def _parse_date(value: Any, epoch: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        converted = from_excel(value, epoch)
        if isinstance(converted, datetime):
            return converted.date()
        if isinstance(converted, date):
            return converted
    text = str(value or "").strip()
    if not text:
        raise ValueError("Az utolsó karbantartás dátuma üres")
    for fmt in ("%Y-%m-%d", "%Y.%m.%d.", "%Y.%m.%d", "%d.%m.%Y", "%d.%m.%Y."):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError as exc:
        raise ValueError(f"Nem értelmezhető karbantartási dátum: {text}") from exc


def parse_maintenance_workbook(content: bytes, filename: str) -> tuple[str, list[ParsedMaintenanceRow], list[MaintenanceImportIssue], int]:
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Csak .xlsx formátumú Excel fájl tölthető fel")
    if not content:
        raise HTTPException(status_code=400, detail="A feltöltött Excel fájl üres")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Az Excel fájl legfeljebb 25 MB lehet")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="A feltöltött fájl nem olvasható Excel munkafüzetként") from exc

    sheet = None
    expected_normalized = _normalize_header(EXPECTED_SHEET_NAME)
    for candidate in workbook.worksheets:
        if _normalize_header(candidate.title) == expected_normalized:
            sheet = candidate
            break
    if sheet is None:
        sheet = workbook.active

    rows_iter = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows_iter)
    except StopIteration as exc:
        raise HTTPException(status_code=400, detail="Az Excel munkalap üres") from exc

    columns: dict[str, int] = {}
    for index, raw_header in enumerate(raw_headers):
        canonical = HEADER_ALIASES.get(_normalize_header(raw_header))
        if canonical and canonical not in columns:
            columns[canonical] = index
    if "maintenance_date" not in columns:
        raise HTTPException(status_code=400, detail="Hiányzó kötelező oszlop: utolsó karbantartás dátuma")
    if "serial_number" not in columns and not {"import_batch", "import_row"}.issubset(columns):
        raise HTTPException(
            status_code=400,
            detail="Az eszköz azonosításához gyári szám vagy import_batch + import_row oszlop szükséges",
        )

    parsed_rows: list[ParsedMaintenanceRow] = []
    issues: list[MaintenanceImportIssue] = []
    total_rows = 0

    def cell(values: tuple[Any, ...], name: str) -> Any:
        index = columns.get(name)
        return values[index] if index is not None and index < len(values) else None

    for excel_row, values in enumerate(rows_iter, start=2):
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        total_rows += 1
        serial_number = _clean_text(cell(values, "serial_number"))
        import_row = None
        maintenance_date = None
        try:
            import_row = _parse_optional_int(cell(values, "import_row"))
            maintenance_date = _parse_date(cell(values, "maintenance_date"), workbook.epoch)
            if maintenance_date > date.today():
                raise ValueError("Az utolsó karbantartás dátuma nem lehet jövőbeli")
            parsed_rows.append(
                ParsedMaintenanceRow(
                    excel_row=excel_row,
                    import_batch=_clean_text(cell(values, "import_batch")),
                    import_row=import_row,
                    serial_number=serial_number,
                    company_code=_clean_text(cell(values, "company_code")),
                    customer_name=_clean_text(cell(values, "customer_name")),
                    asset_name=_clean_text(cell(values, "asset_name")),
                    maintenance_date=maintenance_date,
                    source_sheet=_clean_text(cell(values, "source_sheet")),
                    source_row=_parse_optional_int(cell(values, "source_row")),
                    source_work_order=_clean_text(cell(values, "source_work_order")),
                    technician=_clean_text(cell(values, "technician")),
                    reason=_clean_text(cell(values, "reason")),
                    work_done=_clean_text(cell(values, "work_done")),
                    note=_clean_text(cell(values, "note")),
                )
            )
        except (ValueError, TypeError) as exc:
            if len(issues) < MAX_ERROR_SAMPLES:
                issues.append(
                    MaintenanceImportIssue(
                        excel_row=excel_row,
                        reason=str(exc),
                        serial_number=serial_number,
                        import_row=import_row,
                        maintenance_date=maintenance_date.isoformat() if maintenance_date else None,
                    )
                )
    workbook.close()
    return sheet.title, parsed_rows, issues, total_rows


def _match_asset(db: Session, row: ParsedMaintenanceRow, cache: dict[tuple[Any, ...], list[Asset]]) -> tuple[Asset | None, str | None]:
    if row.import_batch and row.import_row is not None:
        key = ("import", row.import_batch, row.import_row)
        if key not in cache:
            cache[key] = db.query(Asset).filter(Asset.import_batch == row.import_batch, Asset.import_row == row.import_row).all()
        matches = cache[key]
        if len(matches) == 1:
            return matches[0], None
        if len(matches) > 1:
            return None, "Az import_batch + import_row több eszközt azonosít"

    if row.serial_number:
        normalized_serial = row.serial_number.strip()
        if row.company_code:
            key = ("serial_company", normalized_serial, row.company_code.strip().lower())
            if key not in cache:
                cache[key] = db.query(Asset).filter(
                    Asset.serial_number == normalized_serial,
                    Asset.company_code.ilike(row.company_code.strip()),
                ).all()
            matches = cache[key]
            if len(matches) == 1:
                return matches[0], None
            if len(matches) > 1:
                return None, "A gyári szám és cégkód több eszközt azonosít"
        key = ("serial", normalized_serial)
        if key not in cache:
            cache[key] = db.query(Asset).filter(Asset.serial_number == normalized_serial).all()
        matches = cache[key]
        if len(matches) == 1:
            return matches[0], None
        if len(matches) > 1:
            return None, "A gyári szám több eszköznél is szerepel; cégkód vagy importazonosító szükséges"
    return None, "Nem található az Excel sorhoz tartozó eszköz"


def import_maintenance_workbook(
    db: Session,
    current_user: User,
    content: bytes,
    filename: str,
    *,
    dry_run: bool,
) -> MaintenanceImportResult:
    sheet_name, parsed_rows, parse_issues, total_rows = parse_maintenance_workbook(content, filename)
    result = MaintenanceImportResult(
        mode="preview" if dry_run else "import",
        filename=filename,
        sheet_name=sheet_name,
        total_rows=total_rows,
    )
    result.invalid_count = len(parse_issues)
    result.skipped_count = len(parse_issues)
    for issue in parse_issues:
        result.add_issue(issue)

    asset_cache: dict[tuple[Any, ...], list[Asset]] = {}
    latest_by_asset: dict[int, tuple[ParsedMaintenanceRow, Asset]] = {}

    for row in parsed_rows:
        asset, match_error = _match_asset(db, row, asset_cache)
        if not asset:
            result.skipped_count += 1
            if match_error and "több" in match_error:
                result.ambiguous_asset_count += 1
            else:
                result.missing_asset_count += 1
            result.add_issue(
                MaintenanceImportIssue(
                    excel_row=row.excel_row,
                    reason=match_error or "Nem található eszköz",
                    serial_number=row.serial_number,
                    import_row=row.import_row,
                    maintenance_date=row.maintenance_date.isoformat(),
                )
            )
            continue
        row.asset_id = asset.id
        current = latest_by_asset.get(asset.id)
        if current is None or row.maintenance_date > current[0].maintenance_date:
            if current is not None:
                result.skipped_count += 1
                result.add_issue(
                    MaintenanceImportIssue(
                        excel_row=current[0].excel_row,
                        reason="Ugyanahhoz az eszközhöz újabb dátum is szerepel; a régebbi sor kihagyva",
                        serial_number=current[0].serial_number,
                        import_row=current[0].import_row,
                        maintenance_date=current[0].maintenance_date.isoformat(),
                    )
                )
            latest_by_asset[asset.id] = (row, asset)
        else:
            result.skipped_count += 1
            result.add_issue(
                MaintenanceImportIssue(
                    excel_row=row.excel_row,
                    reason="Ugyanahhoz az eszközhöz azonos vagy újabb dátum már szerepel az importban",
                    serial_number=row.serial_number,
                    import_row=row.import_row,
                    maintenance_date=row.maintenance_date.isoformat(),
                )
            )

    result.valid_rows = len(latest_by_asset)
    changed_asset_ids: set[int] = set()

    for row, asset in latest_by_asset.values():
        existing = asset.last_maintenance_date
        if existing == row.maintenance_date:
            result.unchanged_count += 1
            changed_asset_ids.add(asset.id)
            continue
        if existing and existing > row.maintenance_date:
            result.conflict_count += 1
            result.skipped_count += 1
            result.add_issue(
                MaintenanceImportIssue(
                    excel_row=row.excel_row,
                    reason=f"Az adatbázisban újabb karbantartási dátum szerepel: {existing.isoformat()}",
                    serial_number=row.serial_number,
                    import_row=row.import_row,
                    maintenance_date=row.maintenance_date.isoformat(),
                )
            )
            continue

        changed_asset_ids.add(asset.id)
        if dry_run:
            result.would_update_count += 1
            continue

        asset.last_maintenance_date = row.maintenance_date
        details = [f"Utolsó karbantartás dátuma importálva: {row.maintenance_date.isoformat()}"]
        if row.source_work_order:
            details.append(f"forrás munkalap: {row.source_work_order}")
        if row.technician:
            details.append(f"technikus: {row.technician}")
        if row.source_row is not None:
            details.append(f"forrás Excel sor: {row.source_row}")
        add_asset_history(db, asset.id, current_user, "karbantartási dátum import", "; ".join(details))
        result.updated_count += 1

    result.affected_assets_count = len(changed_asset_ids)
    if dry_run:
        db.rollback()
    else:
        db.flush()
    return result
