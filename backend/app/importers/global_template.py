from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

TEMPLATE_FILENAME = "globalis_import_minta.xlsx"

DARK_BLUE = "1F4E78"
MEDIUM_BLUE = "5B9BD5"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "E7E6E6"
LIGHT_ORANGE = "FCE4D6"
LIGHT_GREEN = "E2F0D9"
WHITE = "FFFFFF"
INPUT_BLUE = "0000FF"
STATIC_GRAY = "666666"
ERROR_RED = "F4CCCC"

ASSET_HEADERS = [
    "Cégkód",
    "Ügyfél neve",
    "Ügyfél címe",
    "Helyszín neve",
    "Helyszín címe",
    "Kapcsolattartó neve",
    "Kapcsolattartó típusa",
    "Kapcsolattartó beosztása",
    "Kapcsolattartó telefon",
    "Kapcsolattartó email",
    "Kapcsolattartó elsődleges",
    "Belső azonosító",
    "Gyári szám",
    "Eszköz típusa",
    "Gyártó",
    "Modell",
    "Kategória",
    "Állapot",
    "Belső használat",
    "Vásárlás dátuma",
    "Garancia lejárata",
    "Utolsó karbantartás",
    "Karbantartási ciklus hónap",
    "Karbantartási ciklus megjegyzés",
    "Következő karbantartás",
    "Eszköz megjegyzés",
]

CONTACT_HEADERS = [
    "Cégkód",
    "Ügyfél neve",
    "Helyszín neve",
    "Helyszín címe",
    "Kapcsolattartó típusa",
    "Név",
    "Beosztás",
    "Telefon",
    "Email",
    "Elsődleges",
    "Megjegyzés",
]

METER_HEADERS = [
    "Cégkód",
    "Ügyfél neve",
    "Helyszín neve",
    "Belső azonosító",
    "Gyári szám",
    "Mérés időpontja",
    "Számlálóállás",
    "FF számlálóállás",
    "Színes számlálóállás",
    "Scan számlálóállás",
    "Munkalapszám",
    "Megjegyzés",
]


def _style_header(sheet, headers):
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 34
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headers)).coordinate}"
    sheet.sheet_view.showGridLines = False


def _prepare_input_rows(sheet, columns, rows=300):
    for row in sheet.iter_rows(min_row=2, max_row=rows + 1, max_col=columns):
        for cell in row:
            cell.font = Font(color=INPUT_BLUE)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.conditional_formatting.add(
        f"A2:{sheet.cell(rows + 1, columns).coordinate}",
        FormulaRule(formula=["COUNTA($A2:$ZZ2)>0"], fill=PatternFill("solid", fgColor="F8FBFE")),
    )


def _set_widths(sheet, widths):
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def build_global_import_template() -> bytes:
    wb = Workbook()
    guide = wb.active
    guide.title = "Útmutató"
    guide.sheet_view.showGridLines = False
    guide.merge_cells("A1:F1")
    guide["A1"] = "Globális Excel-import – mintafájl"
    guide["A1"].fill = PatternFill("solid", fgColor=DARK_BLUE)
    guide["A1"].font = Font(color=WHITE, bold=True, size=16)
    guide["A1"].alignment = Alignment(horizontal="left", vertical="center")
    guide.row_dimensions[1].height = 30

    instructions = [
        ("Cél", "Egyetlen munkafüzetből ügyfelek, helyszínek, eszközök, kapcsolattartók, karbantartási adatok és számlálóállások importálhatók."),
        ("Használat", "1. Töltsd ki az Eszközök lapot. 2. Szükség esetén töltsd ki a Kapcsolattartók és Számlálók lapot. 3. A Beállítások oldalon először futtasd az Import ellenőrzése műveletet."),
        ("Üres cellák", "Az üres cella nem törli a már meglévő adatot. Csak a kitöltött mezők frissülnek."),
        ("Eszközazonosítás", "Elsőként a Belső azonosító, utána a Gyári szám + Cégkód, végül az egyedi Gyári szám alapján történik."),
        ("Új eszköz", "Új eszköznél az Ügyfél neve és az Eszköz típusa kötelező. Belső azonosító vagy Gyári szám közül legalább az egyiket add meg."),
        ("Karbantartási ciklus", "Egész hónapszám: például 1, 3, 6 vagy 12. Nem periodikus szabály a Karbantartási ciklus megjegyzés mezőbe írható."),
        ("Dátumformátum", "Ajánlott: ÉÉÉÉ-HH-NN. A mérés időpontja lehet ÉÉÉÉ-HH-NN ÓÓ:PP."),
        ("Számláló", "Az összes számláló kötelező; az FF, színes és scan számláló opcionális. Azonos időpontra a hiányzó kiegészítő értékek később pótolhatók. Az összes számláló csökkenő értéke nem importálható; a részletes csatornák újraindulhatnak."),
        ("Kapcsolattartók", "Az Eszközök lapon egy elsődleges kapcsolattartó megadható. További kapcsolattartókat a Kapcsolattartók lapon adj meg."),
        ("Minták", "A Példák lapon másolható mintasorok találhatók. A Példák lapot a rendszer nem importálja."),
    ]
    guide["A3"] = "Téma"
    guide["B3"] = "Leírás"
    for cell in guide[3]:
        cell.fill = PatternFill("solid", fgColor=MEDIUM_BLUE)
        cell.font = Font(color=WHITE, bold=True)
    for row_index, (topic, description) in enumerate(instructions, start=4):
        guide.cell(row_index, 1, topic).font = Font(bold=True, color=STATIC_GRAY)
        guide.cell(row_index, 2, description).font = Font(color=STATIC_GRAY)
        guide.cell(row_index, 2).alignment = Alignment(wrap_text=True, vertical="top")
    guide["A15"] = "Fontos"
    guide["B15"] = "Import előtt mindig készíts teljes adatbázismentést, és először az ellenőrzési módot használd."
    guide["A15"].fill = PatternFill("solid", fgColor=LIGHT_ORANGE)
    guide["B15"].fill = PatternFill("solid", fgColor=LIGHT_ORANGE)
    guide["A15"].font = Font(bold=True)
    guide["B15"].font = Font(bold=True)
    guide.column_dimensions["A"].width = 25
    guide.column_dimensions["B"].width = 105
    for row in range(4, 16):
        guide.row_dimensions[row].height = 34

    assets = wb.create_sheet("Eszközök")
    _style_header(assets, ASSET_HEADERS)
    _prepare_input_rows(assets, len(ASSET_HEADERS))
    _set_widths(assets, {
        "A": 12, "B": 28, "C": 34, "D": 24, "E": 34, "F": 24, "G": 22, "H": 22,
        "I": 18, "J": 28, "K": 18, "L": 20, "M": 20, "N": 24, "O": 18, "P": 20,
        "Q": 22, "R": 18, "S": 16, "T": 17, "U": 17, "V": 20, "W": 25, "X": 30,
        "Y": 22, "Z": 36,
    })
    for col in (20, 21, 22, 25):
        for row in range(2, 302):
            assets.cell(row, col).number_format = "yyyy-mm-dd"

    contacts = wb.create_sheet("Kapcsolattartók")
    _style_header(contacts, CONTACT_HEADERS)
    _prepare_input_rows(contacts, len(CONTACT_HEADERS))
    _set_widths(contacts, {"A": 12, "B": 28, "C": 24, "D": 34, "E": 22, "F": 24, "G": 22, "H": 18, "I": 28, "J": 15, "K": 40})

    meters = wb.create_sheet("Számlálók")
    _style_header(meters, METER_HEADERS)
    _prepare_input_rows(meters, len(METER_HEADERS), rows=1000)
    _set_widths(meters, {"A": 12, "B": 28, "C": 24, "D": 20, "E": 20, "F": 22, "G": 18, "H": 18, "I": 20, "J": 18, "K": 18, "L": 40})
    for row in range(2, 1002):
        meters.cell(row, 6).number_format = "yyyy-mm-dd hh:mm"
        for col in range(7, 11):
            meters.cell(row, col).number_format = "#,##0"

    examples = wb.create_sheet("Példák")
    examples.sheet_view.showGridLines = False
    examples["A1"] = "Eszközök lap – mintasor"
    examples["A1"].fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    examples["A1"].font = Font(bold=True)
    for index, header in enumerate(ASSET_HEADERS, start=1):
        examples.cell(2, index, header)
    asset_example = [
        "DRH", "Minta Ügyfél Kft.", "1111 Budapest, Minta utca 1.", "Központ", "1111 Budapest, Minta utca 1.",
        "Minta János", "Szerviz", "üzemeltetési vezető", "+36 30 123 4567", "minta@example.com", "Igen",
        "DRH-NY-001", "SN-123456", "Multifunkciós nyomtató", "Ricoh", "IM C3000", "Nyomtató / MFP", "aktív", "Nem",
        "2024-01-15", "2027-01-15", "2026-06-10", 3, None, None, "Mintasor – másold az Eszközök lapra, majd írd át.",
    ]
    for index, value in enumerate(asset_example, start=1):
        examples.cell(3, index, value)
        examples.cell(3, index).font = Font(color=INPUT_BLUE)
    examples["A5"] = "Kapcsolattartók lap – további kapcsolattartó"
    examples["A5"].fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    examples["A5"].font = Font(bold=True)
    for index, header in enumerate(CONTACT_HEADERS, start=1):
        examples.cell(6, index, header)
    contact_example = ["DRH", "Minta Ügyfél Kft.", "Központ", "1111 Budapest, Minta utca 1.", "Sales", "Minta Anna", "beszerző", "+36 30 765 4321", "anna@example.com", "Igen", None]
    for index, value in enumerate(contact_example, start=1):
        examples.cell(7, index, value)
        examples.cell(7, index).font = Font(color=INPUT_BLUE)
    examples["A9"] = "Számlálók lap – két történeti mérés"
    examples["A9"].fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    examples["A9"].font = Font(bold=True)
    for index, header in enumerate(METER_HEADERS, start=1):
        examples.cell(10, index, header)
    meter_examples = [
        ["DRH", "Minta Ügyfél Kft.", "Központ", "DRH-NY-001", "SN-123456", "2026-06-01 12:00", 120000, 90000, 30000, 2500, None, None],
        ["DRH", "Minta Ügyfél Kft.", "Központ", "DRH-NY-001", "SN-123456", "2026-07-01 12:00", 126500, 94500, 32000, 2800, None, None],
    ]
    for row_index, values in enumerate(meter_examples, start=11):
        for column_index, value in enumerate(values, start=1):
            examples.cell(row_index, column_index, value)
            examples.cell(row_index, column_index).font = Font(color=INPUT_BLUE)
    for column in range(1, max(len(ASSET_HEADERS), len(CONTACT_HEADERS), len(METER_HEADERS)) + 1):
        examples.column_dimensions[examples.cell(1, column).column_letter].width = 20
    examples.freeze_panes = "A2"

    lists = wb.create_sheet("Listák")
    values = {
        "A": ["aktív", "hibás", "javítás alatt", "selejtezett", "raktáron"],
        "B": ["Szerviz", "Sales", "Egyéb"],
        "C": ["Igen", "Nem"],
    }
    for column, items in values.items():
        for index, item in enumerate(items, start=1):
            lists[f"{column}{index}"] = item
    lists.sheet_state = "hidden"

    status_validation = DataValidation(type="list", formula1="=Listák!$A$1:$A$5", allow_blank=True)
    asset_contact_validation = DataValidation(type="list", formula1="=Listák!$B$1:$B$3", allow_blank=True)
    asset_boolean_validation = DataValidation(type="list", formula1="=Listák!$C$1:$C$2", allow_blank=True)
    contact_category_validation = DataValidation(type="list", formula1="=Listák!$B$1:$B$3", allow_blank=True)
    contact_boolean_validation = DataValidation(type="list", formula1="=Listák!$C$1:$C$2", allow_blank=True)
    assets.add_data_validation(status_validation)
    status_validation.add("R2:R301")
    assets.add_data_validation(asset_contact_validation)
    asset_contact_validation.add("G2:G301")
    assets.add_data_validation(asset_boolean_validation)
    asset_boolean_validation.add("K2:K301")
    asset_boolean_validation.add("S2:S301")
    contacts.add_data_validation(contact_category_validation)
    contact_category_validation.add("E2:E301")
    contacts.add_data_validation(contact_boolean_validation)
    contact_boolean_validation.add("J2:J301")

    assets["B1"].comment = Comment("Új eszköznél kötelező.", "OpenAI")
    assets["N1"].comment = Comment("Új eszköznél kötelező.", "OpenAI")
    assets["L1"].comment = Comment("Elsődleges eszközazonosító, ha ki van töltve.", "OpenAI")
    assets["M1"].comment = Comment("Belső azonosító hiányában ezzel és a cégkóddal történik a párosítás.", "OpenAI")
    meters["F1"].comment = Comment("Ajánlott formátum: 2026-07-01 12:00", "OpenAI")
    meters["G1"].comment = Comment("Összes nyomatszámláló. Nem lehet negatív, és a történeti sorrendben nem csökkenhet.", "OpenAI")
    meters["H1"].comment = Comment("Opcionális fekete-fehér nyomatszámláló.", "OpenAI")
    meters["I1"].comment = Comment("Opcionális színes nyomatszámláló.", "OpenAI")
    meters["J1"].comment = Comment("Opcionális scan / beolvasási számláló.", "OpenAI")

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
