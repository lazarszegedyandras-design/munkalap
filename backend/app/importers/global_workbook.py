from __future__ import annotations

from bisect import bisect_left
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, time
from io import BytesIO
import math
import re
import unicodedata
import zlib
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy.orm import Session

from ..models import (
    ASSET_STATUSES,
    CONTACT_CATEGORIES,
    Asset,
    AssetMeterReading,
    Customer,
    CustomerContact,
    Location,
    User,
    WorkOrder,
)
from ..services.audit import add_asset_history
from ..services.maintenance_schedule import maintenance_cycle_label
from .customer_assets import clean_asset_note_after_import, parse_machine_type

MAX_UPLOAD_BYTES = 50 * 1024 * 1024
MAX_ERROR_SAMPLES = 200
GLOBAL_IMPORT_BATCH = "global_excel"

ASSET_SHEET = "Eszközök"
CONTACT_SHEET = "Kapcsolattartók"
METER_SHEET = "Számlálók"


@dataclass
class GlobalImportIssue:
    sheet: str
    excel_row: int | None
    entity: str
    identifier: str | None
    reason: str
    level: str = "error"


@dataclass
class GlobalImportResult:
    mode: str
    filename: str
    asset_rows: int = 0
    contact_rows: int = 0
    meter_rows: int = 0
    customers_created: int = 0
    customers_updated: int = 0
    locations_created: int = 0
    locations_updated: int = 0
    contacts_created: int = 0
    contacts_updated: int = 0
    assets_created: int = 0
    assets_updated: int = 0
    assets_unchanged: int = 0
    maintenance_dates_updated: int = 0
    maintenance_cycles_updated: int = 0
    maintenance_conflicts: int = 0
    meter_readings_imported: int = 0
    meter_readings_would_import: int = 0
    meter_readings_unchanged: int = 0
    meter_conflicts: int = 0
    missing_asset_count: int = 0
    ambiguous_asset_count: int = 0
    invalid_count: int = 0
    skipped_count: int = 0
    affected_assets_count: int = 0
    issues: list[GlobalImportIssue] = field(default_factory=list)

    def add_issue(self, issue: GlobalImportIssue) -> None:
        self.skipped_count += 1
        if issue.level == "error":
            self.invalid_count += 1
        if len(self.issues) < MAX_ERROR_SAMPLES:
            self.issues.append(issue)

    def as_dict(self) -> dict[str, Any]:
        payload = asdict(self)
        payload["total_rows"] = self.asset_rows + self.contact_rows + self.meter_rows
        payload["issues_truncated"] = self.skipped_count > len(self.issues)
        return payload


def _normalize(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).strip().split())
    return text or None


def _key(value: Any) -> str:
    return (_clean(value) or "").casefold()


def _parse_int(value: Any, label: str, *, optional: bool = True) -> int | None:
    if value is None or str(value).strip() == "":
        if optional:
            return None
        raise ValueError(f"A(z) {label} mező kötelező")
    if isinstance(value, bool):
        raise ValueError(f"A(z) {label} mező nem lehet logikai érték")
    try:
        number = float(str(value).replace(" ", "").replace("\u00a0", "").replace(",", "."))
    except ValueError as exc:
        raise ValueError(f"A(z) {label} mező nem szám") from exc
    if math.isnan(number) or math.isinf(number) or number != int(number):
        raise ValueError(f"A(z) {label} mező csak egész szám lehet")
    return int(number)


def _parse_bool(value: Any) -> bool | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        return value
    normalized = _normalize(value)
    if normalized in {"igen", "i", "yes", "true", "1", "x"}:
        return True
    if normalized in {"nem", "n", "no", "false", "0"}:
        return False
    raise ValueError(f"Nem értelmezhető igen/nem érték: {value}")


def _parse_date(value: Any, epoch: Any, label: str) -> date | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        converted = from_excel(value, epoch)
        return converted.date() if isinstance(converted, datetime) else converted
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y.%m.%d.", "%Y.%m.%d", "%d.%m.%Y", "%d.%m.%Y."):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError as exc:
        raise ValueError(f"Nem értelmezhető {label}: {text}") from exc


def _parse_datetime(value: Any, epoch: Any) -> datetime:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, time(hour=12))
    if isinstance(value, (int, float)):
        converted = from_excel(value, epoch)
        if isinstance(converted, datetime):
            return converted.replace(tzinfo=None)
        return datetime.combine(converted, time(hour=12))
    text = str(value or "").strip()
    if not text:
        raise ValueError("A mérés időpontja kötelező")
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        pass
    for fmt in (
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
        "%Y.%m.%d. %H:%M",
        "%Y.%m.%d.",
        "%d.%m.%Y %H:%M",
        "%d.%m.%Y",
    ):
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed if parsed.time() != time.min else parsed.replace(hour=12)
        except ValueError:
            continue
    raise ValueError(f"Nem értelmezhető mérési időpont: {text}")


ASSET_HEADERS = {
    "cegkod": "company_code",
    "ugyfel_neve": "customer_name",
    "ugyfel": "customer_name",
    "ugyfel_cime": "customer_address",
    "helyszin_neve": "location_name",
    "telephely_neve": "location_name",
    "helyszin_cime": "location_address",
    "telephely_cime": "location_address",
    "kapcsolattarto_neve": "contact_name",
    "kapcsolattarto_tipusa": "contact_category",
    "kapcsolattarto_beosztasa": "contact_title",
    "kapcsolattarto_telefon": "contact_phone",
    "kapcsolattarto_email": "contact_email",
    "kapcsolattarto_elsodleges": "contact_primary",
    "belso_azonosito": "internal_id",
    "eszkoz_belso_azonosito": "internal_id",
    "gyari_szam": "serial_number",
    "eszkoz_tipusa": "asset_type",
    "gep_tipusa": "asset_type",
    "gyarto": "manufacturer",
    "modell": "model",
    "kategoria": "category",
    "allapot": "status",
    "belso_hasznalat": "internal_use",
    "vasarlas_datuma": "purchase_date",
    "garancia_lejarata": "warranty_expiry",
    "utolso_karbantartas": "last_maintenance_date",
    "utolso_karbantartas_datuma": "last_maintenance_date",
    "karbantartasi_ciklus_honap": "maintenance_cycle_months",
    "karb_ciklus_honap": "maintenance_cycle_months",
    "karbantartasi_ciklus_megjegyzes": "maintenance_cycle_note",
    "kovetkezo_karbantartas": "next_maintenance_date",
    "eszkoz_megjegyzes": "asset_note",
}

CONTACT_HEADERS = {
    "cegkod": "company_code",
    "ugyfel_neve": "customer_name",
    "ugyfel": "customer_name",
    "helyszin_neve": "location_name",
    "helyszin_cime": "location_address",
    "tipus": "category",
    "kapcsolattarto_tipusa": "category",
    "nev": "name",
    "kapcsolattarto_neve": "name",
    "beosztas": "title",
    "telefon": "phone",
    "email": "email",
    "elsodleges": "is_primary",
    "megjegyzes": "note",
}

METER_HEADERS = {
    "cegkod": "company_code",
    "ugyfel_neve": "customer_name",
    "ugyfel": "customer_name",
    "helyszin_neve": "location_name",
    "belso_azonosito": "internal_id",
    "gyari_szam": "serial_number",
    "meres_idopontja": "recorded_at",
    "meres_datuma": "recorded_at",
    "szamlaloallas": "value",
    "szamlalo_allas": "value",
    "ff_szamlaloallas": "black_white_value",
    "ff_zaro": "black_white_value",
    "fekete_feher_szamlaloallas": "black_white_value",
    "szines_szamlaloallas": "color_value",
    "szines_zaro": "color_value",
    "scan_szamlaloallas": "scan_value",
    "scan_zaro": "scan_value",
    "munkalapszam": "work_order_number",
    "megjegyzes": "note",
}


def _sheet(workbook: Any, expected: str) -> Any | None:
    normalized = _normalize(expected)
    return next((item for item in workbook.worksheets if _normalize(item.title) == normalized), None)


def _columns(sheet: Any, aliases: dict[str, str]) -> tuple[dict[str, int], Any]:
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = next(rows)
    except StopIteration:
        return {}, iter(())
    columns: dict[str, int] = {}
    for index, value in enumerate(headers):
        canonical = aliases.get(_normalize(value))
        if canonical and canonical not in columns:
            columns[canonical] = index
    return columns, rows


def _value(values: tuple[Any, ...], columns: dict[str, int], name: str) -> Any:
    index = columns.get(name)
    return values[index] if index is not None and index < len(values) else None


def _find_customer(db: Session, company_code: str | None, name: str) -> Customer | None:
    query = db.query(Customer)
    if company_code:
        query = query.filter(Customer.company_code.ilike(company_code))
    candidates = query.all()
    wanted = _key(name)
    return next((customer for customer in candidates if _key(customer.name) == wanted), None)


def _upsert_customer(db: Session, row: dict[str, Any], result: GlobalImportResult) -> Customer:
    name = _clean(row.get("customer_name"))
    if not name:
        raise ValueError("Az ügyfél neve kötelező")
    company_code = _clean(row.get("company_code"))
    customer = _find_customer(db, company_code, name)
    created = customer is None
    if customer is None:
        customer = Customer(name=name, company_code=company_code)
        db.add(customer)
        db.flush()
        result.customers_created += 1
    changed = False
    mapping = {
        "address": _clean(row.get("customer_address")),
        "company_code": company_code,
    }
    for field_name, incoming in mapping.items():
        if incoming is not None and getattr(customer, field_name) != incoming:
            setattr(customer, field_name, incoming)
            changed = True
    if changed and not created:
        result.customers_updated += 1
    return customer


def _find_location(db: Session, customer: Customer, name: str, address: str | None) -> Location | None:
    wanted_name, wanted_address = _key(name), _key(address)
    return next(
        (
            location
            for location in db.query(Location).filter(Location.customer_id == customer.id).all()
            if _key(location.name) == wanted_name and _key(location.address) == wanted_address
        ),
        None,
    )


def _upsert_location(db: Session, customer: Customer, row: dict[str, Any], result: GlobalImportResult) -> Location:
    address = _clean(row.get("location_address")) or _clean(row.get("customer_address"))
    name = _clean(row.get("location_name")) or "Alapértelmezett helyszín"
    location = _find_location(db, customer, name, address)
    if location is None:
        location = Location(customer_id=customer.id, name=name, address=address)
        db.add(location)
        db.flush()
        result.locations_created += 1
        return location
    if not location.is_active:
        raise ValueError(f"Az importált helyszín inaktív; előbb aktiváld: {location.name}")
    changed = False
    if address is not None and location.address != address:
        location.address = address
        changed = True
    if changed:
        result.locations_updated += 1
    return location


def _find_contact(
    db: Session,
    customer: Customer,
    location: Location | None,
    name: str,
    email: str | None,
    phone: str | None,
) -> CustomerContact | None:
    candidates = db.query(CustomerContact).filter(CustomerContact.customer_id == customer.id).all()
    for contact in candidates:
        if _key(contact.name) != _key(name):
            continue
        if email and _key(contact.email) == _key(email):
            return contact
        if phone and _key(contact.phone) == _key(phone):
            return contact
        if not email and not phone and contact.location_id == (location.id if location else None):
            return contact
    return None


def _upsert_contact(
    db: Session,
    customer: Customer,
    location: Location | None,
    row: dict[str, Any],
    result: GlobalImportResult,
) -> CustomerContact | None:
    name = _clean(row.get("contact_name") or row.get("name"))
    phone = _clean(row.get("contact_phone") or row.get("phone"))
    email = _clean(row.get("contact_email") or row.get("email"))
    if not name and not phone and not email:
        return None
    name = name or email or phone or "Importált kapcsolattartó"
    category = _clean(row.get("contact_category") or row.get("category")) or "Szerviz"
    normalized_category = next((item for item in CONTACT_CATEGORIES if _key(item) == _key(category)), None)
    if normalized_category is None:
        raise ValueError(f"Ismeretlen kapcsolattartó-típus: {category}")
    is_primary = row.get("contact_primary") if "contact_primary" in row else row.get("is_primary")
    if is_primary is None:
        is_primary = False
    contact = _find_contact(db, customer, location, name, email, phone)
    created = contact is None
    if contact is None:
        contact = CustomerContact(customer_id=customer.id, name=name, category=normalized_category)
        db.add(contact)
        db.flush()
        result.contacts_created += 1
    changed = False
    values = {
        "location_id": location.id if location else None,
        "category": normalized_category,
        "title": _clean(row.get("contact_title") or row.get("title")),
        "phone": phone,
        "email": email,
        "note": _clean(row.get("note")),
        "is_primary": bool(is_primary),
    }
    for field_name, incoming in values.items():
        if incoming is not None and getattr(contact, field_name) != incoming:
            setattr(contact, field_name, incoming)
            changed = True
    if contact.is_primary:
        query = db.query(CustomerContact).filter(
            CustomerContact.customer_id == customer.id,
            CustomerContact.category == contact.category,
            CustomerContact.id != contact.id,
        )
        if contact.location_id is None:
            query = query.filter(CustomerContact.location_id.is_(None))
        else:
            query = query.filter(CustomerContact.location_id == contact.location_id)
        query.update({CustomerContact.is_primary: False}, synchronize_session=False)
    if changed and not created:
        result.contacts_updated += 1
    if not customer.contact_person and name:
        customer.contact_person = name
    if not customer.phone and phone:
        customer.phone = phone
    if not customer.email and email:
        customer.email = email
    return contact


def _asset_candidates(db: Session, row: dict[str, Any]) -> list[Asset]:
    internal_id = _clean(row.get("internal_id"))
    if internal_id:
        matches = db.query(Asset).filter(Asset.internal_id == internal_id).all()
        if matches:
            return matches
    serial = _clean(row.get("serial_number"))
    if not serial:
        return []
    query = db.query(Asset).filter(Asset.serial_number == serial)
    company_code = _clean(row.get("company_code"))
    if company_code:
        query = query.filter(Asset.company_code.ilike(company_code))
    candidates = query.all()
    customer_name = _clean(row.get("customer_name"))
    if customer_name and len(candidates) > 1:
        filtered = [asset for asset in candidates if asset.customer and _key(asset.customer.name) == _key(customer_name)]
        if filtered:
            candidates = filtered
    location_name = _clean(row.get("location_name"))
    if location_name and len(candidates) > 1:
        filtered = [asset for asset in candidates if asset.current_location and _key(asset.current_location.name) == _key(location_name)]
        if filtered:
            candidates = filtered
    return candidates


def _generated_internal_id(row: dict[str, Any]) -> str:
    identity = "|".join(
        _clean(row.get(key)) or ""
        for key in ("company_code", "serial_number", "customer_name", "location_name", "asset_type")
    )
    checksum = zlib.crc32(identity.encode("utf-8")) & 0xFFFFFFFF
    return f"IMP-00000000-{checksum}"


def _set_if_provided(target: Any, field_name: str, incoming: Any) -> bool:
    if incoming is None:
        return False
    if getattr(target, field_name) == incoming:
        return False
    setattr(target, field_name, incoming)
    return True


def _upsert_asset(
    db: Session,
    row: dict[str, Any],
    customer: Customer,
    location: Location,
    current_user: User,
    filename: str,
    excel_row: int,
    result: GlobalImportResult,
    affected: set[int],
) -> Asset:
    matches = _asset_candidates(db, row)
    if len(matches) > 1:
        result.ambiguous_asset_count += 1
        raise ValueError("Az eszközazonosító több adatbázisrekordhoz tartozik")
    asset = matches[0] if matches else None
    created = asset is None
    asset_type = _clean(row.get("asset_type"))
    if asset is None:
        if not asset_type:
            raise ValueError("Új eszköznél az eszköz típusa kötelező")
        if not _clean(row.get("internal_id")) and not _clean(row.get("serial_number")):
            raise ValueError("Új eszköznél belső azonosító vagy gyári szám szükséges")
        internal_id = _clean(row.get("internal_id")) or _generated_internal_id(row)
        existing_id = db.query(Asset).filter(Asset.internal_id == internal_id).first()
        if existing_id:
            raise ValueError(f"A belső azonosító már használatban van: {internal_id}")
        asset = Asset(
            internal_id=internal_id,
            type=asset_type,
            status=_clean(row.get("status")) or "aktív",
            customer_id=customer.id,
            current_location_id=location.id,
            internal_use=bool(row.get("internal_use") or False),
            import_source=filename,
            import_batch=GLOBAL_IMPORT_BATCH,
            import_row=excel_row,
        )
        db.add(asset)
        db.flush()
        result.assets_created += 1
    changed_fields: list[str] = []

    status = _clean(row.get("status"))
    if status:
        normalized_status = next((item for item in ASSET_STATUSES if _key(item) == _key(status)), None)
        if normalized_status is None:
            raise ValueError(f"Ismeretlen eszközállapot: {status}")
        row["status"] = normalized_status

    manufacturer, model, category = parse_machine_type(asset_type)
    values = {
        "serial_number": _clean(row.get("serial_number")),
        "type": asset_type,
        "manufacturer": _clean(row.get("manufacturer")) or (manufacturer if created else None),
        "model": _clean(row.get("model")) or (model if created else None),
        "category": _clean(row.get("category")) or (category if created else None),
        "status": row.get("status"),
        "customer_id": customer.id,
        "current_location_id": location.id,
        "internal_use": row.get("internal_use"),
        "purchase_date": row.get("purchase_date"),
        "warranty_expiry": row.get("warranty_expiry"),
        "maintenance_cycle_note": _clean(row.get("maintenance_cycle_note")),
        "next_maintenance_date": row.get("next_maintenance_date"),
        "note": _clean(row.get("asset_note")),
        "company_code": _clean(row.get("company_code")),
    }
    for field_name, incoming in values.items():
        if _set_if_provided(asset, field_name, incoming):
            changed_fields.append(field_name)

    incoming_last = row.get("last_maintenance_date")
    if incoming_last:
        if asset.last_maintenance_date and incoming_last < asset.last_maintenance_date:
            result.maintenance_conflicts += 1
            result.issues.append(
                GlobalImportIssue(
                    sheet=ASSET_SHEET,
                    excel_row=excel_row,
                    entity="eszköz",
                    identifier=_clean(row.get("serial_number")) or asset.internal_id,
                    reason=(
                        f"Az adatbázisban újabb utolsó karbantartás van ({asset.last_maintenance_date}); "
                        f"a régebbi {incoming_last} dátum nem írta felül"
                    ),
                    level="warning",
                )
            )
        elif asset.last_maintenance_date != incoming_last:
            asset.last_maintenance_date = incoming_last
            result.maintenance_dates_updated += 1
            changed_fields.append("last_maintenance_date")

    incoming_cycle = row.get("maintenance_cycle_months")
    if incoming_cycle is not None and asset.maintenance_cycle_months != incoming_cycle:
        asset.maintenance_cycle_months = incoming_cycle
        result.maintenance_cycles_updated += 1
        changed_fields.append("maintenance_cycle_months")

    cleaned_note = clean_asset_note_after_import(asset.note)
    if cleaned_note != asset.note:
        asset.note = cleaned_note
        changed_fields.append("note")

    affected.add(asset.id)
    if created:
        add_asset_history(db, asset.id, current_user, "globális Excel-import", f"Eszköz létrehozva: {filename}, {ASSET_SHEET} {excel_row}. sor")
    elif changed_fields:
        result.assets_updated += 1
        add_asset_history(
            db,
            asset.id,
            current_user,
            "globális Excel-import",
            f"Eszköz frissítve ({', '.join(sorted(set(changed_fields)))}): {filename}, {ASSET_SHEET} {excel_row}. sor",
        )
    else:
        result.assets_unchanged += 1
    return asset


def _parse_asset_row(values: tuple[Any, ...], columns: dict[str, int], epoch: Any) -> dict[str, Any]:
    row = {name: _clean(_value(values, columns, name)) for name in columns}
    if not row.get("customer_name"):
        raise ValueError("Az ügyfél neve kötelező")
    if not row.get("asset_type"):
        raise ValueError("Az eszköz típusa kötelező")
    if not row.get("internal_id") and not row.get("serial_number"):
        raise ValueError("Belső azonosító vagy gyári szám szükséges")
    status = row.get("status")
    if status:
        normalized_status = next((item for item in ASSET_STATUSES if _key(item) == _key(status)), None)
        if normalized_status is None:
            raise ValueError(f"Ismeretlen eszközállapot: {status}")
        row["status"] = normalized_status
    contact_category = row.get("contact_category")
    if contact_category:
        normalized_category = next((item for item in CONTACT_CATEGORIES if _key(item) == _key(contact_category)), None)
        if normalized_category is None:
            raise ValueError(f"Ismeretlen kapcsolattartó-típus: {contact_category}")
        row["contact_category"] = normalized_category
    for field_name, label in (
        ("purchase_date", "vásárlás dátuma"),
        ("warranty_expiry", "garancia lejárata"),
        ("last_maintenance_date", "utolsó karbantartás"),
        ("next_maintenance_date", "következő karbantartás"),
    ):
        row[field_name] = _parse_date(_value(values, columns, field_name), epoch, label)
    row["maintenance_cycle_months"] = _parse_int(
        _value(values, columns, "maintenance_cycle_months"), "karbantartási ciklus", optional=True
    )
    if row["maintenance_cycle_months"] is not None and row["maintenance_cycle_months"] <= 0:
        raise ValueError("A karbantartási ciklus csak pozitív hónapszám lehet")
    row["internal_use"] = _parse_bool(_value(values, columns, "internal_use"))
    row["contact_primary"] = _parse_bool(_value(values, columns, "contact_primary"))
    return row


def _process_asset_sheet(
    db: Session,
    workbook: Any,
    current_user: User,
    filename: str,
    result: GlobalImportResult,
    affected: set[int],
) -> None:
    sheet = _sheet(workbook, ASSET_SHEET)
    if sheet is None:
        raise HTTPException(status_code=400, detail=f"Hiányzó kötelező munkalap: {ASSET_SHEET}")
    columns, rows = _columns(sheet, ASSET_HEADERS)
    required = {"customer_name", "asset_type"}
    missing = required - set(columns)
    if missing:
        raise HTTPException(status_code=400, detail=f"Az {ASSET_SHEET} munkalap hiányzó oszlopai: {', '.join(sorted(missing))}")
    if "serial_number" not in columns and "internal_id" not in columns:
        raise HTTPException(status_code=400, detail=f"Az {ASSET_SHEET} munkalapon gyári szám vagy belső azonosító oszlop szükséges")

    for excel_row, values in enumerate(rows, start=2):
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        result.asset_rows += 1
        identifier = None
        try:
            row = _parse_asset_row(values, columns, workbook.epoch)
            identifier = row.get("serial_number") or row.get("internal_id")
            candidates = _asset_candidates(db, row)
            if len(candidates) > 1:
                result.ambiguous_asset_count += 1
                raise ValueError("Az eszközazonosító több adatbázisrekordhoz tartozik")
            customer = _upsert_customer(db, row, result)
            location = _upsert_location(db, customer, row, result)
            _upsert_contact(db, customer, location, row, result)
            _upsert_asset(db, row, customer, location, current_user, filename, excel_row, result, affected)
        except (ValueError, TypeError) as exc:
            result.add_issue(GlobalImportIssue(ASSET_SHEET, excel_row, "eszköz", identifier, str(exc)))


def _process_contact_sheet(db: Session, workbook: Any, result: GlobalImportResult) -> None:
    sheet = _sheet(workbook, CONTACT_SHEET)
    if sheet is None:
        return
    columns, rows = _columns(sheet, CONTACT_HEADERS)
    if "customer_name" not in columns or "name" not in columns:
        raise HTTPException(status_code=400, detail=f"A {CONTACT_SHEET} munkalapon ügyfél neve és név oszlop szükséges")
    for excel_row, values in enumerate(rows, start=2):
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        result.contact_rows += 1
        try:
            row = {name: _clean(_value(values, columns, name)) for name in columns}
            if not row.get("customer_name"):
                raise ValueError("Az ügyfél neve kötelező")
            if not row.get("name"):
                raise ValueError("A kapcsolattartó neve kötelező")
            category = row.get("category") or "Szerviz"
            normalized_category = next((item for item in CONTACT_CATEGORIES if _key(item) == _key(category)), None)
            if normalized_category is None:
                raise ValueError(f"Ismeretlen kapcsolattartó-típus: {category}")
            row["category"] = normalized_category
            row["is_primary"] = _parse_bool(_value(values, columns, "is_primary"))
            customer = _upsert_customer(db, row, result)
            location = None
            if row.get("location_name") or row.get("location_address"):
                location = _upsert_location(db, customer, row, result)
            _upsert_contact(db, customer, location, row, result)
        except (ValueError, TypeError) as exc:
            result.add_issue(
                GlobalImportIssue(CONTACT_SHEET, excel_row, "kapcsolattartó", _clean(_value(values, columns, "name")), str(exc))
            )


def _match_meter_asset(db: Session, row: dict[str, Any]) -> tuple[Asset | None, str | None]:
    matches = _asset_candidates(db, row)
    if len(matches) == 1:
        return matches[0], None
    if len(matches) > 1:
        return None, "A megadott adatok több eszközt azonosítanak"
    return None, "A megadott adatokkal nem található eszköz"


def _meter_values(row: dict[str, Any]) -> dict[str, int | None]:
    return {
        "value": row.get("value"),
        "black_white_value": row.get("black_white_value"),
        "color_value": row.get("color_value"),
        "scan_value": row.get("scan_value"),
    }


def _process_meter_sheet(
    db: Session,
    workbook: Any,
    current_user: User,
    result: GlobalImportResult,
    affected: set[int],
    *,
    dry_run: bool,
) -> None:
    sheet = _sheet(workbook, METER_SHEET)
    if sheet is None:
        return
    columns, rows = _columns(sheet, METER_HEADERS)
    if "recorded_at" not in columns or "value" not in columns:
        raise HTTPException(status_code=400, detail=f"A {METER_SHEET} munkalapon mérés időpontja és számlálóállás oszlop szükséges")
    if "serial_number" not in columns and "internal_id" not in columns:
        raise HTTPException(status_code=400, detail=f"A {METER_SHEET} munkalapon gyári szám vagy belső azonosító oszlop szükséges")

    parsed: list[tuple[int, dict[str, Any], Asset]] = []
    for excel_row, values in enumerate(rows, start=2):
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        result.meter_rows += 1
        identifier = _clean(_value(values, columns, "serial_number")) or _clean(_value(values, columns, "internal_id"))
        try:
            row = {name: _clean(_value(values, columns, name)) for name in columns}
            row["recorded_at"] = _parse_datetime(_value(values, columns, "recorded_at"), workbook.epoch)
            row["value"] = _parse_int(_value(values, columns, "value"), "számlálóállás", optional=False)
            for field, label in (("black_white_value", "FF számlálóállás"), ("color_value", "színes számlálóállás"), ("scan_value", "scan számlálóállás")):
                row[field] = _parse_int(_value(values, columns, field), label, optional=True)
            if any(value is not None and value < 0 for value in _meter_values(row).values()):
                raise ValueError("A számlálóállások nem lehetnek negatívak")
            if row.get("black_white_value") is not None and row.get("color_value") is not None:
                if row["black_white_value"] + row["color_value"] != row["value"]:
                    raise ValueError("Az összes számláló nem egyezik az FF + színes számlálók összegével")
            asset, error = _match_meter_asset(db, row)
            if not asset:
                if error and "több" in error:
                    result.ambiguous_asset_count += 1
                else:
                    result.missing_asset_count += 1
                raise ValueError(error or "Eszköz nem található")
            parsed.append((excel_row, row, asset))
        except (ValueError, TypeError) as exc:
            result.add_issue(GlobalImportIssue(METER_SHEET, excel_row, "számláló", identifier, str(exc)))

    by_asset: dict[int, list[tuple[int, dict[str, Any], Asset]]] = {}
    for item in parsed:
        by_asset.setdefault(item[2].id, []).append(item)

    changed_per_asset: dict[int, int] = {}
    fields = ("value", "black_white_value", "color_value", "scan_value")
    labels = {"black_white_value": "FF", "color_value": "színes", "scan_value": "scan"}
    for asset_id, items in by_asset.items():
        existing_rows = db.query(AssetMeterReading).filter(AssetMeterReading.asset_id == asset_id).all()
        existing_by_time = {reading.recorded_at: reading for reading in existing_rows}
        total_timeline = sorted([(r.recorded_at, int(r.value)) for r in existing_rows], key=lambda item: item[0])
        for excel_row, row, asset in sorted(items, key=lambda item: (item[1]["recorded_at"], item[0])):
            existing = existing_by_time.get(row["recorded_at"])
            if existing:
                if int(existing.value) != row["value"]:
                    result.meter_conflicts += 1
                    result.add_issue(GlobalImportIssue(METER_SHEET, excel_row, "számláló", row.get("serial_number") or row.get("internal_id"), f"Erre az időpontra már {existing.value} összes számláló van eltárolva", level="warning"))
                    continue
                conflict = None
                updates: dict[str, int] = {}
                for field in fields[1:]:
                    incoming = row.get(field)
                    current = getattr(existing, field)
                    if incoming is None:
                        continue
                    if current is not None and int(current) != incoming:
                        conflict = f"Erre az időpontra már eltérő {labels[field]} számláló van eltárolva"
                        break
                    if current is None:
                        updates[field] = incoming
                if conflict:
                    result.meter_conflicts += 1
                    result.add_issue(GlobalImportIssue(METER_SHEET, excel_row, "számláló", row.get("serial_number") or row.get("internal_id"), conflict, level="warning"))
                    continue
                if not updates:
                    result.meter_readings_unchanged += 1
                    continue
                result.meter_readings_would_import += 1
                affected.add(asset.id)
                if not dry_run:
                    for field, incoming in updates.items():
                        setattr(existing, field, incoming)
                    result.meter_readings_imported += 1
                    changed_per_asset[asset.id] = changed_per_asset.get(asset.id, 0) + 1
                continue

            position = bisect_left([item[0] for item in total_timeline], row["recorded_at"])
            previous = total_timeline[position - 1] if position > 0 else None
            following = total_timeline[position] if position < len(total_timeline) else None
            metric_conflict = None
            if previous and row["value"] < previous[1]:
                metric_conflict = f"Az összes számláló kisebb a korábbi {previous[1]} értéknél"
            elif following and row["value"] > following[1]:
                metric_conflict = f"Az összes számláló nagyobb a későbbi {following[1]} értéknél"
            if metric_conflict:
                result.meter_conflicts += 1
                result.add_issue(GlobalImportIssue(METER_SHEET, excel_row, "számláló", row.get("serial_number") or row.get("internal_id"), metric_conflict, level="warning"))
                continue
            result.meter_readings_would_import += 1
            affected.add(asset.id)
            total_timeline.insert(position, (row["recorded_at"], row["value"]))
            if dry_run:
                continue
            work_order_id = None
            work_order_number = _clean(row.get("work_order_number"))
            if work_order_number:
                order = db.query(WorkOrder).filter(WorkOrder.number == work_order_number).first()
                work_order_id = order.id if order else None
            created = AssetMeterReading(
                asset_id=asset.id,
                value=row["value"],
                black_white_value=row.get("black_white_value"),
                color_value=row.get("color_value"),
                scan_value=row.get("scan_value"),
                recorded_at=row["recorded_at"],
                work_order_id=work_order_id,
                recorded_by_user_id=current_user.id,
                note=_clean(row.get("note")),
            )
            db.add(created)
            existing_by_time[row["recorded_at"]] = created
            result.meter_readings_imported += 1
            changed_per_asset[asset.id] = changed_per_asset.get(asset.id, 0) + 1

    if not dry_run:
        db.flush()
        for asset_id, count in changed_per_asset.items():
            add_asset_history(db, asset_id, current_user, "globális Excel-import", f"{count} számlálósor importálva vagy kiegészítve")


def import_global_workbook(
    db: Session,
    current_user: User,
    content: bytes,
    filename: str,
    *,
    dry_run: bool,
) -> GlobalImportResult:
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Csak .xlsx formátumú globális importfájl tölthető fel")
    if not content:
        raise HTTPException(status_code=400, detail="A feltöltött Excel fájl üres")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="A globális importfájl legfeljebb 50 MB lehet")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="A feltöltött fájl nem olvasható Excel munkafüzetként") from exc

    result = GlobalImportResult(mode="preview" if dry_run else "import", filename=filename)
    affected: set[int] = set()
    try:
        _process_asset_sheet(db, workbook, current_user, filename, result, affected)
        _process_contact_sheet(db, workbook, result)
        db.flush()
        _process_meter_sheet(db, workbook, current_user, result, affected, dry_run=dry_run)
        result.affected_assets_count = len(affected)
        if dry_run:
            db.rollback()
        else:
            db.flush()
        return result
    except HTTPException:
        db.rollback()
        raise
    except Exception:
        db.rollback()
        raise
    finally:
        workbook.close()
