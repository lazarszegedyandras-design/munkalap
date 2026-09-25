from __future__ import annotations

from dataclasses import asdict, dataclass, field
from io import BytesIO
import math
import re
import unicodedata
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy.orm import Session, joinedload

from ..models import Asset, User
from ..services.asset_display import asset_display_name
from ..services.audit import add_asset_history
from ..services.maintenance_schedule import maintenance_cycle_label

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_ERROR_SAMPLES = 100
EXPECTED_SHEET_NAMES = {"karbantartasi ciklus import", "munka1"}


@dataclass
class ParsedCycleRow:
    excel_row: int
    import_batch: str | None
    import_row: int | None
    serial_number: str | None
    company_code: str | None
    customer_name: str | None
    location_address: str | None
    asset_name: str | None
    cycle_months: int | None
    cycle_note: str | None
    source_sheet: str | None = None
    source_row: int | None = None
    asset_id: int | None = None


@dataclass
class CycleImportIssue:
    excel_row: int | None
    reason: str
    serial_number: str | None = None
    import_row: int | None = None
    cycle: str | None = None


@dataclass
class CycleImportResult:
    mode: str
    filename: str
    sheet_name: str
    total_rows: int = 0
    valid_rows: int = 0
    would_update_count: int = 0
    updated_count: int = 0
    unchanged_count: int = 0
    invalid_count: int = 0
    missing_asset_count: int = 0
    ambiguous_asset_count: int = 0
    affected_assets_count: int = 0
    non_periodic_count: int = 0
    issues: list[CycleImportIssue] = field(default_factory=list)

    def add_issue(self, issue: CycleImportIssue) -> None:
        if len(self.issues) < MAX_ERROR_SAMPLES:
            self.issues.append(issue)

    def as_dict(self) -> dict[str, Any]:
        payload = asdict(self)
        payload["issues_truncated"] = (
            self.invalid_count + self.missing_asset_count + self.ambiguous_asset_count
        ) > len(self.issues)
        return payload


def _normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


HEADER_ALIASES = {
    "import_batch": "import_batch",
    "import_row": "import_row",
    "ceg": "company_code",
    "cegkod": "company_code",
    "ugyfel": "customer_name",
    "ugyfel_neve": "customer_name",
    "telephely_cim": "location_address",
    "cim": "location_address",
    "gep_tipusa": "asset_name",
    "gep_tipus": "asset_name",
    "eszkoz": "asset_name",
    "gyariszam": "serial_number",
    "gyari_szam": "serial_number",
    "serial_number": "serial_number",
    "karb_ciklus": "cycle_raw",
    "karbantartasi_ciklus": "cycle_raw",
    "karbantartasi_ciklus_honap": "cycle_months",
    "ciklus_honap": "cycle_months",
    "maintenance_cycle_months": "cycle_months",
    "ciklus_megjegyzes": "cycle_note",
    "maintenance_cycle_note": "cycle_note",
    "forras_munkalap": "source_sheet",
    "forras_excel_sor": "source_row",
    "forras_sor": "source_row",
}


def _clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_optional_int(value: Any, field_name: str) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        raise ValueError(f"A(z) {field_name} mező nem szám")
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            raise ValueError(f"A(z) {field_name} mező érvénytelen")
        if float(value) != int(value):
            raise ValueError(f"A(z) {field_name} mező csak egész szám lehet")
        return int(value)
    normalized = str(value).replace(" ", "").replace("\u00a0", "").replace(",", ".").strip()
    number = float(normalized)
    if number != int(number):
        raise ValueError(f"A(z) {field_name} mező csak egész szám lehet")
    return int(number)


def _parse_cycle(raw_value: Any, months_value: Any = None, note_value: Any = None) -> tuple[int | None, str | None]:
    explicit_months = _parse_optional_int(months_value, "karbantartási ciklus (hónap)")
    explicit_note = _clean_text(note_value)
    if explicit_months is not None:
        if not 1 <= explicit_months <= 120:
            raise ValueError("A karbantartási ciklus 1 és 120 hónap közötti lehet")
        return explicit_months, explicit_note

    raw = _clean_text(raw_value)
    if raw is None:
        if explicit_note:
            return None, explicit_note
        raise ValueError("A karbantartási ciklus mező üres")
    normalized = unicodedata.normalize("NFKD", raw.lower())
    normalized = "".join(char for char in normalized if not unicodedata.combining(char))
    normalized = re.sub(r"\s+", " ", normalized).strip()

    month_match = re.fullmatch(r"(\d+)\s*(ho|honap|honapos)", normalized)
    if month_match:
        months = int(month_match.group(1))
        if not 1 <= months <= 120:
            raise ValueError("A karbantartási ciklus 1 és 120 hónap közötti lehet")
        return months, explicit_note

    year_match = re.fullmatch(r"(\d+)\s*(ev|eves)", normalized)
    if year_match:
        months = int(year_match.group(1)) * 12
        if not 1 <= months <= 120:
            raise ValueError("A karbantartási ciklus 1 és 120 hónap közötti lehet")
        return months, explicit_note

    if normalized in {"-", "nincs", "nincs ciklus", "nem szerzodeses"}:
        return None, explicit_note or "Nincs rögzített szerződéses ciklus"
    return None, explicit_note or raw


def parse_cycle_workbook(content: bytes, filename: str) -> tuple[str, list[ParsedCycleRow], list[CycleImportIssue], int]:
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Csak .xlsx formátumú Excel fájl tölthető fel")
    if not content:
        raise HTTPException(status_code=400, detail="A feltöltött Excel fájl üres")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Az Excel fájl legfeljebb 25 MB lehet")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="A feltöltött fájl nem olvasható Excel munkafüzetként") from exc

    sheet = workbook.active
    for candidate in workbook.worksheets:
        normalized_title = _normalize_header(candidate.title).replace("_", " ")
        if normalized_title in EXPECTED_SHEET_NAMES:
            sheet = candidate
            break

    rows_iter = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows_iter)
    except StopIteration as exc:
        raise HTTPException(status_code=400, detail="Az Excel munkalap üres") from exc

    columns: dict[str, int] = {}
    for index, raw_header in enumerate(raw_headers):
        canonical = HEADER_ALIASES.get(_normalize_header(raw_header))
        if canonical and canonical not in columns:
            columns[canonical] = index
    if "cycle_raw" not in columns and "cycle_months" not in columns:
        raise HTTPException(status_code=400, detail="Hiányzó kötelező oszlop: Karb. Ciklus / karbantartási ciklus (hónap)")
    if "serial_number" not in columns and not {"import_batch", "import_row"}.issubset(columns):
        if not {"customer_name", "asset_name"}.issubset(columns):
            raise HTTPException(
                status_code=400,
                detail="Az eszköz azonosításához gyári szám, import_batch + import_row, vagy ügyfél + géptípus szükséges",
            )

    parsed_rows: list[ParsedCycleRow] = []
    issues: list[CycleImportIssue] = []
    total_rows = 0
    forwarded: dict[str, Any] = {"company_code": None, "customer_name": None, "location_address": None}

    def cell(values: tuple[Any, ...], name: str) -> Any:
        index = columns.get(name)
        return values[index] if index is not None and index < len(values) else None

    for excel_row, values in enumerate(rows_iter, start=2):
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        for key in forwarded:
            current = cell(values, key)
            if current is not None and str(current).strip() != "":
                forwarded[key] = current
        # A raw workbook second row contains month headings, not asset data.
        if not any(cell(values, name) not in (None, "") for name in ("serial_number", "asset_name", "customer_name", "import_row")):
            continue
        total_rows += 1
        import_row = None
        serial_number = _clean_text(cell(values, "serial_number"))
        cycle_text = _clean_text(cell(values, "cycle_raw")) or _clean_text(cell(values, "cycle_note"))
        try:
            import_row = _parse_optional_int(cell(values, "import_row"), "import_row")
            cycle_months, cycle_note = _parse_cycle(
                cell(values, "cycle_raw"),
                cell(values, "cycle_months"),
                cell(values, "cycle_note"),
            )
            parsed_rows.append(
                ParsedCycleRow(
                    excel_row=excel_row,
                    import_batch=_clean_text(cell(values, "import_batch")),
                    import_row=import_row,
                    serial_number=serial_number,
                    company_code=_clean_text(cell(values, "company_code")) or _clean_text(forwarded["company_code"]),
                    customer_name=_clean_text(cell(values, "customer_name")) or _clean_text(forwarded["customer_name"]),
                    location_address=_clean_text(cell(values, "location_address")) or _clean_text(forwarded["location_address"]),
                    asset_name=_clean_text(cell(values, "asset_name")),
                    cycle_months=cycle_months,
                    cycle_note=cycle_note,
                    source_sheet=_clean_text(cell(values, "source_sheet")) or sheet.title,
                    source_row=_parse_optional_int(cell(values, "source_row"), "forrás Excel sor") or excel_row,
                )
            )
        except (ValueError, TypeError) as exc:
            issues.append(
                CycleImportIssue(
                    excel_row=excel_row,
                    reason=str(exc),
                    serial_number=serial_number,
                    import_row=import_row,
                    cycle=cycle_text,
                )
            )
    workbook.close()
    return sheet.title, parsed_rows, issues[:MAX_ERROR_SAMPLES], total_rows


def _normalized(value: Any) -> str:
    text = str(value or "").strip().upper()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^A-Z0-9]", "", text)


def _match_asset(row: ParsedCycleRow, assets: list[Asset]) -> tuple[Asset | None, str | None]:
    if row.import_batch and row.import_row is not None:
        exact = [asset for asset in assets if asset.import_batch == row.import_batch and asset.import_row == row.import_row]
        if len(exact) == 1:
            return exact[0], None
        if len(exact) > 1:
            return None, "Az import_batch + import_row több eszközt azonosít"

    serial = _normalized(row.serial_number)
    if serial:
        candidates = [asset for asset in assets if _normalized(asset.serial_number) == serial]
        if not candidates:
            # Csak szűk, karaktereltérés nélküli előtag-eltérést fogadunk el
            # (például egy forrásban levágott rövid utótagot). Teljesen eltérő
            # gyári számnál tilos ügyfél/cím alapján másik eszközre találgatni.
            candidates = [
                asset for asset in assets
                if asset.serial_number
                and (
                    _normalized(asset.serial_number).startswith(serial)
                    or serial.startswith(_normalized(asset.serial_number))
                )
                and abs(len(_normalized(asset.serial_number)) - len(serial)) <= 3
            ]
        company = _normalized(row.company_code)
        if company:
            company_candidates = [asset for asset in candidates if _normalized(asset.company_code) == company]
            if company_candidates:
                candidates = company_candidates
        if len(candidates) == 1:
            return candidates[0], None
        if len(candidates) > 1:
            # Resolve duplicate serials by customer, location and type.
            scored = []
            for asset in candidates:
                score = 0
                if _normalized(asset.customer.name if asset.customer else None) == _normalized(row.customer_name):
                    score += 10
                if _normalized(asset.current_location.address if asset.current_location else None) == _normalized(row.location_address):
                    score += 6
                if _normalized(asset.type) == _normalized(row.asset_name):
                    score += 8
                scored.append((score, asset))
            scored.sort(key=lambda item: item[0], reverse=True)
            if scored and (len(scored) == 1 or scored[0][0] > scored[1][0]):
                return scored[0][1], None
            return None, "A gyári szám több eszközhöz tartozik, a környezeti adatok sem döntik el egyértelműen"
        return None, "A gyári szám alapján az eszköz nem található"

    # Fallback csak akkor használható, ha a forrássorban nincs gyári szám.
    scored = []
    for asset in assets:
        score = 0
        if row.company_code and _normalized(asset.company_code) == _normalized(row.company_code):
            score += 10
        asset_customer = asset.customer.name if asset.customer else None
        if row.customer_name and _normalized(asset_customer) == _normalized(row.customer_name):
            score += 10
        asset_address = asset.current_location.address if asset.current_location else (asset.customer.address if asset.customer else None)
        if row.location_address and _normalized(asset_address) == _normalized(row.location_address):
            score += 6
        if row.asset_name and _normalized(asset.type) == _normalized(row.asset_name):
            score += 8
        if serial and _normalized(asset.serial_number).startswith(serial):
            score += 8
        if score >= 24:
            scored.append((score, asset))
    scored.sort(key=lambda item: item[0], reverse=True)
    if scored and (len(scored) == 1 or scored[0][0] > scored[1][0]):
        return scored[0][1], None
    if scored:
        return None, "Az ügyfél, cím és géptípus alapján több eszköz is megfelel"
    return None, "Az eszköz nem található"


def import_maintenance_cycle_workbook(
    db: Session,
    current_user: User,
    content: bytes,
    filename: str,
    *,
    dry_run: bool,
) -> CycleImportResult:
    sheet_name, rows, parse_issues, total_rows = parse_cycle_workbook(content, filename)
    result = CycleImportResult(
        mode="preview" if dry_run else "run",
        filename=filename,
        sheet_name=sheet_name,
        total_rows=total_rows,
        valid_rows=len(rows),
        invalid_count=len(parse_issues),
        issues=list(parse_issues),
    )
    assets = (
        db.query(Asset)
        .options(joinedload(Asset.customer), joinedload(Asset.current_location))
        .all()
    )
    affected_ids: set[int] = set()

    for row in rows:
        asset, match_error = _match_asset(row, assets)
        if not asset:
            if match_error and "több" in match_error:
                result.ambiguous_asset_count += 1
            else:
                result.missing_asset_count += 1
            result.add_issue(
                CycleImportIssue(
                    excel_row=row.excel_row,
                    reason=match_error or "Az eszköz nem található",
                    serial_number=row.serial_number,
                    import_row=row.import_row,
                    cycle=maintenance_cycle_label(row.cycle_months, row.cycle_note),
                )
            )
            continue

        row.asset_id = asset.id
        affected_ids.add(asset.id)
        if row.cycle_months is None:
            result.non_periodic_count += 1
        if asset.maintenance_cycle_months == row.cycle_months and (asset.maintenance_cycle_note or None) == (row.cycle_note or None):
            result.unchanged_count += 1
            continue

        if dry_run:
            result.would_update_count += 1
            continue

        previous_label = maintenance_cycle_label(asset.maintenance_cycle_months, asset.maintenance_cycle_note) or "nincs megadva"
        next_label = maintenance_cycle_label(row.cycle_months, row.cycle_note) or "nincs megadva"
        asset.maintenance_cycle_months = row.cycle_months
        asset.maintenance_cycle_note = row.cycle_note
        add_asset_history(
            db,
            asset.id,
            current_user,
            "karbantartási ciklus",
            f"Szerződéses karbantartási ciklus: {previous_label} → {next_label}. Forrás: {filename}, {sheet_name} {row.excel_row}. sor",
        )
        result.updated_count += 1

    result.affected_assets_count = len(affected_ids)
    if dry_run:
        db.rollback()
    else:
        db.flush()
    return result
