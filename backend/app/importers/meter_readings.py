from __future__ import annotations

from bisect import bisect_left
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, time
from io import BytesIO
import math
import unicodedata
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy.orm import Session

from ..models import Asset, AssetMeterReading, User
from ..services.audit import add_asset_history

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
EXPECTED_SHEET_NAME = "Számláló import"
MAX_ERROR_SAMPLES = 100


@dataclass
class ParsedMeterRow:
    excel_row: int
    import_batch: str | None
    import_row: int | None
    serial_number: str | None
    company_code: str | None
    customer_name: str | None
    asset_name: str | None
    recorded_at: datetime
    value: int
    black_white_value: int | None = None
    color_value: int | None = None
    scan_value: int | None = None
    source_sheet: str | None = None
    source_row: int | None = None
    note: str | None = None
    asset_id: int | None = None


@dataclass
class MeterImportIssue:
    excel_row: int | None
    reason: str
    serial_number: str | None = None
    import_row: int | None = None
    recorded_at: str | None = None
    value: int | None = None


@dataclass
class MeterImportResult:
    mode: str
    filename: str
    sheet_name: str
    total_rows: int = 0
    valid_rows: int = 0
    would_import_count: int = 0
    imported_count: int = 0
    unchanged_count: int = 0
    conflict_count: int = 0
    invalid_count: int = 0
    missing_asset_count: int = 0
    ambiguous_asset_count: int = 0
    affected_assets_count: int = 0
    skipped_count: int = 0
    issues: list[MeterImportIssue] = field(default_factory=list)

    def add_issue(self, issue: MeterImportIssue) -> None:
        if len(self.issues) < MAX_ERROR_SAMPLES:
            self.issues.append(issue)

    def as_dict(self) -> dict[str, Any]:
        payload = asdict(self)
        payload["issues_truncated"] = self.skipped_count > len(self.issues)
        return payload


def _normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    for char in (" ", "-", "/", ".", "(", ")"):
        text = text.replace(char, "_")
    while "__" in text:
        text = text.replace("__", "_")
    return text.strip("_")


HEADER_ALIASES = {
    "import_batch": "import_batch",
    "import_row": "import_row",
    "gyari_szam": "serial_number",
    "sorozatszam": "serial_number",
    "serial_number": "serial_number",
    "cegkod": "company_code",
    "ugyfel": "customer_name",
    "eszkoz": "asset_name",
    "meres_idopontja": "recorded_at",
    "meres_datum": "recorded_at",
    "recorded_at": "recorded_at",
    "szamlaloallas": "value",
    "szamlalo_allas": "value",
    "value": "value",
    "ff_zaro": "black_white_value",
    "ff_szamlaloallas": "black_white_value",
    "szines_zaro": "color_value",
    "szines_szamlaloallas": "color_value",
    "scan_zaro": "scan_value",
    "scan_szamlaloallas": "scan_value",
    "forras_munkalap": "source_sheet",
    "forras_sor": "source_row",
    "megjegyzes": "note",
}


def _clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_int(value: Any, field_name: str) -> int:
    if value is None or value == "":
        raise ValueError(f"A(z) {field_name} mező üres")
    if isinstance(value, bool):
        raise ValueError(f"A(z) {field_name} mező nem szám")
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            raise ValueError(f"A(z) {field_name} mező érvénytelen")
        if float(value) != int(value):
            raise ValueError(f"A(z) {field_name} mező csak egész szám lehet")
        return int(value)
    normalized = str(value).replace(" ", "").replace("\u00a0", "").replace(",", ".").strip()
    number = float(normalized)
    if number != int(number):
        raise ValueError(f"A(z) {field_name} mező csak egész szám lehet")
    return int(number)


def _parse_optional_int(value: Any, field_name: str = "import_row") -> int | None:
    if value is None or value == "":
        return None
    return _parse_int(value, field_name)


def _parse_datetime(value: Any, epoch: Any) -> datetime:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, time(hour=12))
    if isinstance(value, (int, float)):
        converted = from_excel(value, epoch)
        if isinstance(converted, datetime):
            return converted.replace(tzinfo=None)
        if isinstance(converted, date):
            return datetime.combine(converted, time(hour=12))
    text = str(value or "").strip()
    if not text:
        raise ValueError("A mérés időpontja üres")
    normalized = text.replace("Z", "+00:00")
    try:
        parsed = datetime.fromisoformat(normalized)
        return parsed.replace(tzinfo=None)
    except ValueError:
        pass
    for fmt in ("%Y.%m.%d. %H:%M", "%Y.%m.%d.", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%d.%m.%Y %H:%M", "%d.%m.%Y"):
        try:
            parsed = datetime.strptime(text, fmt)
            if parsed.time() == time.min:
                parsed = parsed.replace(hour=12)
            return parsed
        except ValueError:
            continue
    raise ValueError(f"Nem értelmezhető mérési időpont: {text}")


def parse_meter_workbook(content: bytes, filename: str) -> tuple[str, list[ParsedMeterRow], list[MeterImportIssue], int]:
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Csak .xlsx formátumú Excel fájl tölthető fel")
    if not content:
        raise HTTPException(status_code=400, detail="A feltöltött Excel fájl üres")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Az Excel fájl legfeljebb 25 MB lehet")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="A feltöltött fájl nem olvasható Excel munkafüzetként") from exc

    sheet = None
    expected_normalized = _normalize_header(EXPECTED_SHEET_NAME)
    for candidate in workbook.worksheets:
        if _normalize_header(candidate.title) == expected_normalized:
            sheet = candidate
            break
    if sheet is None:
        sheet = workbook.active

    rows_iter = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows_iter)
    except StopIteration as exc:
        raise HTTPException(status_code=400, detail="Az Excel munkalap üres") from exc

    columns: dict[str, int] = {}
    for index, raw_header in enumerate(raw_headers):
        canonical = HEADER_ALIASES.get(_normalize_header(raw_header))
        if canonical and canonical not in columns:
            columns[canonical] = index
    missing = [name for name in ("recorded_at", "value") if name not in columns]
    if missing:
        raise HTTPException(
            status_code=400,
            detail="Hiányzó kötelező oszlop(ok): " + ", ".join(missing) + ". Elvárt: mérés időpontja, számlálóállás.",
        )
    if "serial_number" not in columns and not {"import_batch", "import_row"}.issubset(columns):
        raise HTTPException(
            status_code=400,
            detail="Az eszköz azonosításához gyári szám vagy import_batch + import_row oszlop szükséges",
        )

    parsed_rows: list[ParsedMeterRow] = []
    issues: list[MeterImportIssue] = []
    total_rows = 0

    def cell(values: tuple[Any, ...], name: str) -> Any:
        index = columns.get(name)
        return values[index] if index is not None and index < len(values) else None

    for excel_row, values in enumerate(rows_iter, start=2):
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        total_rows += 1
        serial_number = _clean_text(cell(values, "serial_number"))
        import_row = None
        try:
            import_row = _parse_optional_int(cell(values, "import_row"))
            value = _parse_int(cell(values, "value"), "számlálóállás")
            black_white_value = _parse_optional_int(cell(values, "black_white_value"), "FF számlálóállás")
            color_value = _parse_optional_int(cell(values, "color_value"), "színes számlálóállás")
            scan_value = _parse_optional_int(cell(values, "scan_value"), "scan számlálóállás")
            if any(item is not None and item < 0 for item in (value, black_white_value, color_value, scan_value)):
                raise ValueError("A számlálóállások nem lehetnek negatívak")
            if black_white_value is not None and color_value is not None and black_white_value + color_value != value:
                raise ValueError("Az összes számláló nem egyezik az FF + színes számlálók összegével")
            recorded_at = _parse_datetime(cell(values, "recorded_at"), workbook.epoch)
            parsed_rows.append(
                ParsedMeterRow(
                    excel_row=excel_row,
                    import_batch=_clean_text(cell(values, "import_batch")),
                    import_row=import_row,
                    serial_number=serial_number,
                    company_code=_clean_text(cell(values, "company_code")),
                    customer_name=_clean_text(cell(values, "customer_name")),
                    asset_name=_clean_text(cell(values, "asset_name")),
                    recorded_at=recorded_at,
                    value=value,
                    black_white_value=black_white_value,
                    color_value=color_value,
                    scan_value=scan_value,
                    source_sheet=_clean_text(cell(values, "source_sheet")),
                    source_row=_parse_optional_int(cell(values, "source_row")),
                    note=_clean_text(cell(values, "note")),
                )
            )
        except (ValueError, TypeError) as exc:
            if len(issues) < MAX_ERROR_SAMPLES:
                issues.append(
                    MeterImportIssue(
                        excel_row=excel_row,
                        reason=str(exc),
                        serial_number=serial_number,
                        import_row=import_row,
                    )
                )
    workbook.close()
    return sheet.title, parsed_rows, issues, total_rows


def _match_asset(db: Session, row: ParsedMeterRow, cache: dict[tuple[Any, ...], list[Asset]]) -> tuple[Asset | None, str | None]:
    if row.import_batch and row.import_row is not None:
        key = ("import", row.import_batch, row.import_row)
        if key not in cache:
            cache[key] = db.query(Asset).filter(Asset.import_batch == row.import_batch, Asset.import_row == row.import_row).all()
        matches = cache[key]
        if len(matches) == 1:
            return matches[0], None
        if len(matches) > 1:
            return None, "Az import_batch + import_row több eszközt azonosít"

    if row.serial_number:
        normalized_serial = row.serial_number.strip()
        if row.company_code:
            key = ("serial_company", normalized_serial, row.company_code.strip().lower())
            if key not in cache:
                cache[key] = db.query(Asset).filter(
                    Asset.serial_number == normalized_serial,
                    Asset.company_code.ilike(row.company_code.strip()),
                ).all()
            matches = cache[key]
            if len(matches) == 1:
                return matches[0], None
            if len(matches) > 1:
                return None, "A gyári szám és cégkód több eszközt azonosít"
        key = ("serial", normalized_serial)
        if key not in cache:
            cache[key] = db.query(Asset).filter(Asset.serial_number == normalized_serial).all()
        matches = cache[key]
        if len(matches) == 1:
            return matches[0], None
        if len(matches) > 1:
            return None, "A gyári szám több eszköznél is szerepel; cégkód vagy importazonosító szükséges"
    return None, "Nem található az Excel sorhoz tartozó eszköz"


def import_meter_workbook(
    db: Session,
    current_user: User,
    content: bytes,
    filename: str,
    *,
    dry_run: bool,
) -> MeterImportResult:
    sheet_name, parsed_rows, parse_issues, total_rows = parse_meter_workbook(content, filename)
    result = MeterImportResult(mode="preview" if dry_run else "import", filename=filename, sheet_name=sheet_name, total_rows=total_rows)
    result.invalid_count = len(parse_issues)
    result.skipped_count = len(parse_issues)
    for issue in parse_issues:
        result.add_issue(issue)

    asset_cache: dict[tuple[Any, ...], list[Asset]] = {}
    rows_by_asset: dict[int, list[ParsedMeterRow]] = defaultdict(list)
    assets_by_id: dict[int, Asset] = {}

    for row in parsed_rows:
        asset, match_error = _match_asset(db, row, asset_cache)
        if not asset:
            if match_error and "több" in match_error:
                result.ambiguous_asset_count += 1
            else:
                result.missing_asset_count += 1
            result.skipped_count += 1
            result.add_issue(
                MeterImportIssue(
                    excel_row=row.excel_row,
                    reason=match_error or "Eszköz nem található",
                    serial_number=row.serial_number,
                    import_row=row.import_row,
                    recorded_at=row.recorded_at.isoformat(),
                    value=row.value,
                )
            )
            continue
        row.asset_id = asset.id
        assets_by_id[asset.id] = asset
        rows_by_asset[asset.id].append(row)

    existing_by_asset: dict[int, list[AssetMeterReading]] = defaultdict(list)
    if rows_by_asset:
        for reading in db.query(AssetMeterReading).filter(AssetMeterReading.asset_id.in_(rows_by_asset.keys())).all():
            existing_by_asset[reading.asset_id].append(reading)

    to_create: list[ParsedMeterRow] = []
    to_update: list[tuple[AssetMeterReading, dict[str, int]]] = []
    imported_per_asset: dict[int, int] = defaultdict(int)
    affected_assets: set[int] = set()
    fields = ("value", "black_white_value", "color_value", "scan_value")

    for asset_id, asset_rows in rows_by_asset.items():
        existing_rows = existing_by_asset.get(asset_id, [])
        existing_by_time = {row.recorded_at: row for row in existing_rows}
        total_timeline = sorted([(r.recorded_at, int(r.value)) for r in existing_rows], key=lambda item: item[0])
        for row in sorted(asset_rows, key=lambda item: (item.recorded_at, item.excel_row)):
            existing = existing_by_time.get(row.recorded_at)
            if existing:
                if int(existing.value) != row.value:
                    result.conflict_count += 1
                    result.skipped_count += 1
                    result.add_issue(MeterImportIssue(excel_row=row.excel_row, reason=f"Erre az időpontra már {existing.value} összes számlálóállás van eltárolva", serial_number=row.serial_number, import_row=row.import_row, recorded_at=row.recorded_at.isoformat(), value=row.value))
                    continue
                updates = {}
                conflict = False
                for field in fields[1:]:
                    incoming = getattr(row, field)
                    current = getattr(existing, field)
                    if incoming is None:
                        continue
                    if current is not None and int(current) != incoming:
                        conflict = True
                        break
                    if current is None:
                        updates[field] = incoming
                if conflict:
                    result.conflict_count += 1
                    result.skipped_count += 1
                    result.add_issue(MeterImportIssue(excel_row=row.excel_row, reason="Erre az időpontra már eltérő részletes számlálóadat van eltárolva", serial_number=row.serial_number, import_row=row.import_row, recorded_at=row.recorded_at.isoformat(), value=row.value))
                    continue
                if not updates:
                    result.unchanged_count += 1
                    continue
                to_update.append((existing, updates))
                affected_assets.add(asset_id)
                result.valid_rows += 1
                continue

            position = bisect_left([item[0] for item in total_timeline], row.recorded_at)
            previous = total_timeline[position - 1] if position > 0 else None
            following = total_timeline[position] if position < len(total_timeline) else None
            sequence_error = None
            if previous and row.value < previous[1]:
                sequence_error = f"Az összes számláló kisebb a korábbi {previous[1]} értéknél"
            elif following and row.value > following[1]:
                sequence_error = f"Az összes számláló nagyobb a későbbi {following[1]} értéknél"
            if sequence_error:
                result.conflict_count += 1
                result.skipped_count += 1
                result.add_issue(MeterImportIssue(excel_row=row.excel_row, reason=sequence_error, serial_number=row.serial_number, import_row=row.import_row, recorded_at=row.recorded_at.isoformat(), value=row.value))
                continue
            total_timeline.insert(position, (row.recorded_at, row.value))
            to_create.append(row)
            affected_assets.add(asset_id)
            result.valid_rows += 1

    result.would_import_count = len(to_create) + len(to_update)
    result.affected_assets_count = len(affected_assets)
    if dry_run:
        return result

    try:
        for existing, updates in to_update:
            for field, value in updates.items():
                setattr(existing, field, value)
            imported_per_asset[existing.asset_id] += 1
        for row in to_create:
            db.add(
                AssetMeterReading(
                    asset_id=row.asset_id,
                    value=row.value,
                    black_white_value=row.black_white_value,
                    color_value=row.color_value,
                    scan_value=row.scan_value,
                    recorded_at=row.recorded_at,
                    recorded_by_user_id=current_user.id,
                    note=row.note,
                )
            )
            imported_per_asset[row.asset_id] += 1
        db.flush()
        for asset_id, count in imported_per_asset.items():
            add_asset_history(
                db,
                asset_id,
                current_user,
                "számláló Excel-import",
                f"{count} historikus számlálóállás importálva Excel fájlból",
            )
        db.commit()
    except Exception:
        db.rollback()
        raise

    result.imported_count = len(to_create) + len(to_update)
    return result
