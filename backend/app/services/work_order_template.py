from __future__ import annotations

from copy import copy
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.worksheet import Worksheet

from ..config import Settings
from ..models import WorkOrder
from .asset_display import asset_display_name, visible_internal_id
from .company_profiles import resolve_work_order_company_profile

STATIC_TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "templates" / "work_order_template.xlsx"


def work_order_template_path(settings: Settings) -> Path:
    runtime_path = Path(settings.runtime_dir) / "work_order_template.xlsx"
    return runtime_path if runtime_path.exists() else STATIC_TEMPLATE_PATH
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _date(value: Any) -> str:
    if not value:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y.%m.%d. %H:%M")
    if isinstance(value, date):
        return value.strftime("%Y.%m.%d.")
    return str(value)


def _money(value: Any) -> str:
    if value is None or value == "":
        return ""
    try:
        amount = Decimal(str(value))
        return f"{amount:,.2f}".replace(",", " ").replace(".", ",")
    except Exception:
        return str(value)


def _number(value: Any) -> str:
    if value is None or value == "":
        return ""
    try:
        amount = Decimal(str(value))
        if amount == amount.to_integral_value():
            return str(int(amount))
        return str(amount).replace(".", ",")
    except Exception:
        return str(value)


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _safe_filename(value: str) -> str:
    allowed = "abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789-_."
    cleaned = "".join(ch if ch in allowed else "_" for ch in value)
    return cleaned.strip("_") or "munkalap"


def work_order_template_filename(order: WorkOrder) -> str:
    return f"{_safe_filename(order.number)}.xlsx"


def _merged_top_left(ws: Worksheet, cell: str) -> str:
    for merged_range in ws.merged_cells.ranges:
        if cell in merged_range:
            return ws.cell(row=merged_range.min_row, column=merged_range.min_col).coordinate
    return cell


def _set(ws: Worksheet, cell: str, value: Any, wrap: bool = True) -> None:
    target = ws[_merged_top_left(ws, cell)]
    target.value = value
    if wrap:
        current_alignment = copy(target.alignment) if target.alignment else Alignment()
        target.alignment = Alignment(
            horizontal=current_alignment.horizontal,
            vertical=current_alignment.vertical or "top",
            text_rotation=current_alignment.text_rotation,
            wrap_text=True,
            shrink_to_fit=current_alignment.shrink_to_fit,
            indent=current_alignment.indent,
        )


def _line_join(lines: list[str]) -> str:
    return "\n".join(line for line in lines if line)


def _asset_lines(order: WorkOrder) -> tuple[str, str, str, str]:
    names: list[str] = []
    serials: list[str] = []
    identifiers: list[str] = []
    ap_numbers: list[str] = []
    for link in order.asset_links:
        asset = link.asset
        if not asset:
            continue
        display_name = " / ".join(part for part in [asset.manufacturer, asset.model or asset.type] if part)
        names.append(display_name or asset_display_name(asset))
        serials.append(asset.serial_number or "")
        identifiers.append(visible_internal_id(asset.internal_id) or "")
        ap_numbers.append("")
    return _line_join(names), _line_join(serials), _line_join(identifiers), _line_join(ap_numbers)


def _latest_meter(asset):
    return max(asset.meter_readings, key=lambda reading: (reading.recorded_at, reading.id), default=None)


def _meter_lines(order: WorkOrder) -> tuple[str, str]:
    labels: list[str] = []
    values: list[str] = []
    for link in order.asset_links:
        reading = _latest_meter(link.asset) if link.asset else None
        if not reading:
            labels.append("Nincs adat")
            values.append("")
            continue
        labels.append(f"{_date(reading.recorded_at)} · Összes / FF / Színes / Scan")
        metrics = [reading.value, reading.black_white_value, reading.color_value, reading.scan_value]
        values.append(" / ".join(_number(metric) if metric is not None else "" for metric in metrics))
    return _line_join(labels), _line_join(values)


def _material_lines(order: WorkOrder) -> tuple[str, str, str, str, str, str, str]:
    sku: list[str] = []
    name: list[str] = []
    quantity: list[str] = []
    unit: list[str] = []
    price: list[str] = []
    net: list[str] = []
    gross: list[str] = []
    for line in order.material_lines:
        material = line.material
        qty = Decimal(str(line.quantity or 0))
        unit_price = Decimal(str(line.unit_price or 0))
        net_value = qty * unit_price
        sku.append(material.sku if material else "")
        name.append(material.name if material else "")
        quantity.append(_number(qty))
        unit.append(material.unit if material else "")
        price.append(_money(unit_price) if line.unit_price is not None else "")
        net.append(_money(net_value) if line.unit_price is not None else "")
        # A program jelenleg nem tárol ÁFA-kulcsot, ezért a bruttó értéket nem számoljuk ki találgatással.
        gross.append("")
    return _line_join(sku), _line_join(name), _line_join(quantity), _line_join(unit), _line_join(price), _line_join(net), _line_join(gross)


def _apply_dynamic_row_heights(ws: Worksheet, order: WorkOrder) -> None:
    asset_rows = max(1, len(order.asset_links))
    ws.row_dimensions[23].height = max(ws.row_dimensions[23].height or 15, 16 * asset_rows)
    ws.row_dimensions[25].height = max(ws.row_dimensions[25].height or 15, 18 * asset_rows)
    ws.row_dimensions[29].height = max(ws.row_dimensions[29].height or 15, 60)
    # Az elvégzett munka a korábbi megjegyzés- és anyagblokk helyét használja.
    ws.row_dimensions[33].height = max(ws.row_dimensions[33].height or 15, 180)
    for row in (29, 33, 46, 47, 48, 49):
        cell_value = ws.cell(row=row, column=1).value
        if isinstance(cell_value, str):
            lines = max(1, cell_value.count("\n") + 1, len(cell_value) // 80 + 1)
            ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 15, 15 * lines)


def _hide_material_section(ws: Worksheet) -> None:
    # A felhasznált anyagok továbbra is az adatbázisban maradnak, de a nyomtatott
    # Excel munkalapon nem jelennek meg.
    for cell in ("A40", "A41", "C41", "G41", "J41", "K41", "N41", "R41", "A43", "C43", "G43", "J43", "K43", "N43", "R43"):
        _set(ws, cell, "")
    for row in range(40, 44):
        ws.row_dimensions[row].hidden = True


def _append_completion_section(ws: Worksheet, order: WorkOrder) -> None:
    start = 45
    thin = Side(style="thin", color="666666")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    title_fill = PatternFill("solid", fgColor="E7E6E6")
    header_font = Font(bold=True)

    ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=20)
    cell = ws.cell(row=start, column=1)
    cell.value = "Zárási adatok"
    cell.font = header_font
    cell.fill = title_fill
    cell.border = border

    rows = [
        ("Végállapot", order.final_status or ""),
        ("Belső megjegyzés", order.internal_note or ""),
        ("Munka típusa", order.work_type or ""),
        ("Lezárás időpontja", _date(order.closed_at)),
    ]
    row = start + 1
    for label, value in rows:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=20)
        label_cell = ws.cell(row=row, column=1)
        value_cell = ws.cell(row=row, column=5)
        label_cell.value = label
        label_cell.font = header_font
        value_cell.value = value
        value_cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col in range(1, 21):
            ws.cell(row=row, column=col).border = border
        row += 1

    # Nagyobb kézi kitöltési tér a zárási adatok és az aláírások között.
    spacer_start = row
    signature_row = row + 6
    for spacer_row in range(spacer_start, signature_row):
        ws.row_dimensions[spacer_row].height = max(ws.row_dimensions[spacer_row].height or 15, 18)
    ws.merge_cells(start_row=signature_row, start_column=1, end_row=signature_row, end_column=10)
    ws.merge_cells(start_row=signature_row, start_column=11, end_row=signature_row, end_column=20)
    technician_name = ", ".join(name for name in (order.technician.full_name if order.technician else None, order.secondary_technician.full_name if order.secondary_technician else None) if name)
    left = ws.cell(row=signature_row, column=1)
    right = ws.cell(row=signature_row, column=11)
    left.value = _line_join(["Technikus aláírása", technician_name])
    right.value = "Ügyfél aláírása"
    for target in (left, right):
        target.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        target.border = Border(top=thin)
    ws.row_dimensions[signature_row].height = 46
    ws.print_area = f"A1:T{signature_row}"


def _select_customer_contact(order: WorkOrder):
    if getattr(order, "contact", None):
        return order.contact
    customer = order.customer
    contacts = list(getattr(customer, "contacts", []) or []) if customer else []
    if not contacts:
        return None
    location_id = order.location_id
    for category in ["Szerviz", "Sales", "Egyéb"]:
        for contact in contacts:
            if contact.category == category and contact.is_primary and contact.location_id == location_id:
                return contact
        for contact in contacts:
            if contact.category == category and contact.is_primary and contact.location_id is None:
                return contact
        for contact in contacts:
            if contact.category == category and contact.location_id == location_id:
                return contact
        for contact in contacts:
            if contact.category == category and contact.location_id is None:
                return contact
    return contacts[0]


def build_work_order_workbook(order: WorkOrder, settings: Settings) -> BytesIO:
    template_path = work_order_template_path(settings)
    if not template_path.exists():
        raise FileNotFoundError(f"Hiányzó munkalap sablon: {template_path}")
    workbook = load_workbook(template_path)
    ws = workbook.active
    ws.title = "Munkalap"

    customer = order.customer
    location = order.location
    contact = _select_customer_contact(order)
    contact_name = contact.name if contact else (customer.contact_person if customer else "")
    contact_phone = contact.phone if contact and contact.phone else (customer.phone if customer else "")
    contact_email = contact.email if contact and contact.email else (customer.email if customer else "")
    customer_address = (location.address if location and location.address else customer.address if customer else "")

    _set(ws, "A1", "Munkalap")
    supplier = resolve_work_order_company_profile(order, settings)
    if supplier:
        supplier_lines = [
            f"Adószám: {supplier.tax_number}" if supplier.tax_number else "Adószám:",
            f"Cégjegyzékszám: {supplier.company_registration_number}" if supplier.company_registration_number else "Cégjegyzékszám:",
            f"Telefonszám: {supplier.phone_for_worksheet}" if supplier.phone_for_worksheet else "Telefonszám:",
            f"E-mail: {supplier.email_for_worksheet}" if supplier.email_for_worksheet else "E-mail:",
        ]
        _set(ws, "A4", supplier.display_name)
        _set(ws, "A5", _line_join(supplier_lines))
        _set(ws, "A10", "")
        _set(ws, "A12", f"Telefon: {supplier.phone_for_worksheet}" if supplier.phone_for_worksheet else "Telefon:")
        _set(ws, "A14", f"E-Mail: {supplier.email_for_worksheet}" if supplier.email_for_worksheet else "E-Mail:")
    else:
        _set(ws, "A4", "")
        _set(ws, "A5", _line_join(["Adószám:", "Cégjegyzékszám:", "Telefonszám:", "E-mail:"]))
        _set(ws, "A10", "")
        _set(ws, "A12", "Telefon:")
        _set(ws, "A14", "E-Mail:")

    _set(ws, "H4", customer.name if customer else "")
    _set(ws, "H6", f"Cím: {customer_address}" if customer_address else "Cím:")
    _set(ws, "H7", f"Telefonszám: {contact_phone}" if contact_phone else "Telefonszám:")
    _set(ws, "H8", f"E-mail: {contact_email}" if contact_email else "E-mail:")
    _set(ws, "H10", f"Kapcsolattartó: {contact_name}" if contact_name else "Kapcsolattartó:")
    _set(ws, "H12", f"Telefon / Mobil: {contact_phone}" if contact_phone else "Telefon / Mobil:")
    _set(ws, "H14", f"E-Mail: {contact_email}" if contact_email else "E-Mail:")

    _set(ws, "A17", order.number)
    _set(ws, "D17", _date(order.created_at))
    _set(ws, "E17", _date(order.planned_date))
    _set(ws, "G17", order.work_type or "")
    _set(ws, "M17", order.status)
    _set(ws, "Q17", "")
    _set(ws, "T17", order.priority)

    names, serials, identifiers, ap_numbers = _asset_lines(order)
    _set(ws, "A23", names)
    _set(ws, "I23", serials)
    _set(ws, "O23", identifiers)
    _set(ws, "S23", ap_numbers)
    meter_labels, meter_values = _meter_lines(order)
    _set(ws, "B21", "Utolsó mérés / számlálók")
    _set(ws, "F21", "Összes / FF / Színes / Scan")
    _set(ws, "B25", meter_labels)
    _set(ws, "F25", meter_values)

    _set(ws, "A29", order.description)
    _set(ws, "A31", "Elvégzett munka:")
    _set(ws, "A33", order.work_done or "")
    _set(ws, "A38", order.technician.full_name if order.technician else "")
    # A nyomtatott munkalapon a technikus kézzel tölti ki a tényleges időpontokat.
    _set(ws, "L38", "")
    _set(ws, "P38", "")

    _hide_material_section(ws)
    _append_completion_section(ws, order)
    _apply_dynamic_row_heights(ws, order)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
