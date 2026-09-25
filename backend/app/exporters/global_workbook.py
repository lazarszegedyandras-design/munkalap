from __future__ import annotations

from copy import copy
from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session, joinedload, selectinload

from ..importers.global_template import (
    ASSET_HEADERS,
    CONTACT_HEADERS,
    METER_HEADERS,
    build_global_import_template,
)
from ..models import Asset, AssetMeterReading, Customer, CustomerContact

EXPORT_FILENAME_PREFIX = "globalis_adat_export"
IMPORTED_GREEN = "008000"
EXPORT_FILL = "E2F0D9"


def _yes_no(value: bool | None) -> str | None:
    if value is None:
        return None
    return "Igen" if value else "Nem"


def _copy_row_style(sheet, source_row: int, target_row: int, columns: int) -> None:
    if target_row == source_row:
        return
    for column in range(1, columns + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.alignment = copy(source.alignment)
    if source_row in sheet.row_dimensions:
        source_dimension = sheet.row_dimensions[source_row]
        target_dimension = sheet.row_dimensions[target_row]
        target_dimension.height = source_dimension.height


def _write_rows(sheet, headers: list[str], rows: list[list[object]], *, date_columns: set[int] | None = None, datetime_columns: set[int] | None = None) -> None:
    date_columns = date_columns or set()
    datetime_columns = datetime_columns or set()
    max_existing = max(sheet.max_row, len(rows) + 1)
    if max_existing > 1:
        for row in sheet.iter_rows(min_row=2, max_row=max_existing, max_col=len(headers)):
            for cell in row:
                cell.value = None

    for row_number, values in enumerate(rows, start=2):
        _copy_row_style(sheet, 2, row_number, len(headers))
        for column_number, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column_number, value)
            cell.font = Font(name=cell.font.name or "Calibri", size=cell.font.sz or 11, color=IMPORTED_GREEN)
            if column_number in date_columns:
                cell.number_format = "yyyy-mm-dd"
            elif column_number in datetime_columns:
                cell.number_format = "yyyy-mm-dd hh:mm"
    last_row = max(1, len(rows) + 1)
    sheet.auto_filter.ref = f"A1:{sheet.cell(last_row, len(headers)).coordinate}"


def _contact_rank(contact: CustomerContact, asset: Asset) -> tuple[int, int, int]:
    same_location = bool(asset.current_location_id and contact.location_id == asset.current_location_id)
    general = contact.location_id is None
    category_rank = {"Szerviz": 0, "Sales": 1, "Egyéb": 2}.get(contact.category, 3)
    location_rank = 0 if same_location else (1 if general else 2)
    primary_rank = 0 if contact.is_primary else 1
    return (location_rank, primary_rank, category_rank)


def _asset_contact(asset: Asset) -> CustomerContact | None:
    if not asset.customer or not asset.customer.contacts:
        return None
    return sorted(asset.customer.contacts, key=lambda item: _contact_rank(item, asset))[0]


def _asset_rows(assets: list[Asset]) -> list[list[object]]:
    rows: list[list[object]] = []
    for asset in assets:
        customer = asset.customer
        location = asset.current_location
        contact = _asset_contact(asset)
        contact_name = contact.name if contact else (customer.contact_person if customer else None)
        contact_phone = contact.phone if contact else (customer.phone if customer else None)
        contact_email = contact.email if contact else (customer.email if customer else None)
        rows.append([
            asset.company_code or (customer.company_code if customer else None),
            customer.name if customer else "Belső használat",
            customer.address if customer else None,
            location.name if location else ("Belső használat" if asset.internal_use else None),
            location.address if location else None,
            contact_name,
            contact.category if contact else ("Szerviz" if contact_name or contact_phone or contact_email else None),
            contact.title if contact else None,
            contact_phone,
            contact_email,
            _yes_no(contact.is_primary) if contact else (_yes_no(True) if contact_name or contact_phone or contact_email else None),
            asset.internal_id,
            asset.serial_number,
            asset.type,
            asset.manufacturer,
            asset.model,
            asset.category,
            asset.status,
            _yes_no(asset.internal_use),
            asset.purchase_date,
            asset.warranty_expiry,
            asset.last_maintenance_date,
            asset.maintenance_cycle_months,
            asset.maintenance_cycle_note,
            asset.next_maintenance_date,
            asset.note,
        ])
    return rows


def _contact_rows(contacts: list[CustomerContact]) -> list[list[object]]:
    rows: list[list[object]] = []
    for contact in contacts:
        customer = contact.customer
        location = contact.location
        rows.append([
            customer.company_code if customer else None,
            customer.name if customer else None,
            location.name if location else None,
            location.address if location else None,
            contact.category,
            contact.name,
            contact.title,
            contact.phone,
            contact.email,
            _yes_no(contact.is_primary),
            contact.note,
        ])
    return rows


def _meter_rows(readings: list[AssetMeterReading]) -> list[list[object]]:
    rows: list[list[object]] = []
    for reading in readings:
        asset = reading.asset
        customer = asset.customer if asset else None
        location = asset.current_location if asset else None
        rows.append([
            (asset.company_code or (customer.company_code if customer else None)) if asset else None,
            customer.name if customer else None,
            location.name if location else None,
            asset.internal_id if asset else None,
            asset.serial_number if asset else None,
            reading.recorded_at,
            int(reading.value),
            int(reading.black_white_value) if reading.black_white_value is not None else None,
            int(reading.color_value) if reading.color_value is not None else None,
            int(reading.scan_value) if reading.scan_value is not None else None,
            reading.work_order.number if reading.work_order else None,
            reading.note,
        ])
    return rows


def build_global_database_export(db: Session) -> bytes:
    assets = (
        db.query(Asset)
        .options(
            joinedload(Asset.customer).selectinload(Customer.contacts),
            joinedload(Asset.current_location),
        )
        .order_by(Asset.company_code.asc().nullslast(), Asset.internal_id.asc())
        .all()
    )
    contacts = (
        db.query(CustomerContact)
        .options(joinedload(CustomerContact.customer), joinedload(CustomerContact.location))
        .order_by(CustomerContact.customer_id.asc(), CustomerContact.location_id.asc().nullslast(), CustomerContact.category.asc(), CustomerContact.name.asc())
        .all()
    )
    readings = (
        db.query(AssetMeterReading)
        .options(
            joinedload(AssetMeterReading.asset).joinedload(Asset.customer),
            joinedload(AssetMeterReading.asset).joinedload(Asset.current_location),
            joinedload(AssetMeterReading.work_order),
        )
        .order_by(AssetMeterReading.asset_id.asc(), AssetMeterReading.recorded_at.asc())
        .all()
    )

    workbook = load_workbook(BytesIO(build_global_import_template()))
    guide = workbook["Útmutató"]
    guide["A1"] = "Globális Excel-export – visszaimportálható adatállomány"
    guide["A17"] = "Export időpontja"
    guide["B17"] = datetime.now()
    guide["B17"].number_format = "yyyy-mm-dd hh:mm:ss"
    guide["A18"] = "Exportált eszközök"
    guide["B18"] = len(assets)
    guide["A19"] = "Exportált kapcsolattartók"
    guide["B19"] = len(contacts)
    guide["A20"] = "Exportált számlálóállások"
    guide["B20"] = len(readings)
    guide["A21"] = "Visszaimportálás"
    guide["B21"] = "A fájl közvetlenül feltölthető a Globális Excel-import funkcióba. Import előtt készíts teljes rendszermentést."
    for row in range(17, 22):
        guide.cell(row, 1).font = Font(bold=True, color="666666")
        guide.cell(row, 2).font = Font(color="008000")
        guide.cell(row, 1).fill = PatternFill("solid", fgColor=EXPORT_FILL)
        guide.cell(row, 2).fill = PatternFill("solid", fgColor=EXPORT_FILL)

    _write_rows(workbook["Eszközök"], ASSET_HEADERS, _asset_rows(assets), date_columns={20, 21, 22, 25})
    _write_rows(workbook["Kapcsolattartók"], CONTACT_HEADERS, _contact_rows(contacts))
    _write_rows(workbook["Számlálók"], METER_HEADERS, _meter_rows(readings), datetime_columns={6})

    # Az exportban nincs szükség a mintasorokra; eltávolításuk csökkenti a téves szerkesztés kockázatát.
    if "Példák" in workbook.sheetnames:
        workbook.remove(workbook["Példák"])

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
